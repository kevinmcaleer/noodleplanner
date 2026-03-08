"""Tests for the Excel importer module."""

import io
import pytest
from datetime import datetime
from openpyxl import Workbook
import xlrd

from noodle_core.excel_importer import (
    normalize_date,
    detect_hierarchy,
    calculate_duration_from_dates,
    analyze_workbook,
    convert_excel_to_markdown,
)


def _make_xlsx_bytes(sheets_data):
    """Helper: create an xlsx file in memory and return bytes.

    sheets_data: dict of {sheet_name: [row_list, ...]}
    Each row_list is a list of cell values; first row is headers.
    """
    wb = Workbook()
    first = True
    for name, rows in sheets_data.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- TestNormalizeDate ----------

class TestNormalizeDate:
    def test_datetime_object(self):
        dt = datetime(2025, 3, 15)
        assert normalize_date(dt) == "2025-03-15"

    def test_iso_string(self):
        assert normalize_date("2025-03-15") == "2025-03-15"

    def test_dd_mm_yyyy_slash(self):
        assert normalize_date("15/03/2025") == "2025-03-15"

    def test_mm_dd_yyyy_slash(self):
        assert normalize_date("03/15/2025") == "2025-03-15"

    def test_dd_mon_yyyy(self):
        assert normalize_date("15-Mar-2025") == "2025-03-15"

    def test_dd_month_yyyy(self):
        assert normalize_date("15-March-2025") == "2025-03-15"

    def test_excel_serial_number(self):
        # 44000 is ~June 2020
        result = normalize_date(44000)
        assert result is not None
        parsed = datetime.strptime(result, "%Y-%m-%d")
        assert parsed.year == 2020

    def test_float_serial(self):
        result = normalize_date(44000.5)
        assert result is not None

    def test_none_returns_none(self):
        assert normalize_date(None) is None

    def test_empty_string_returns_none(self):
        assert normalize_date("") is None
        assert normalize_date("   ") is None

    def test_invalid_string_returns_none(self):
        assert normalize_date("not-a-date") is None

    def test_date_before_1900_returns_none(self):
        dt = datetime(1899, 12, 31)
        assert normalize_date(dt) is None

    def test_date_after_2100_returns_none(self):
        dt = datetime(2101, 1, 1)
        assert normalize_date(dt) is None

    def test_serial_number_out_of_range(self):
        assert normalize_date(0) is None
        assert normalize_date(-5) is None
        assert normalize_date(300000) is None

    def test_yyyy_mm_dd_slash(self):
        assert normalize_date("2025/03/15") == "2025-03-15"

    def test_dd_mm_yyyy_dot(self):
        assert normalize_date("15.03.2025") == "2025-03-15"


# ---------- TestDetectHierarchy ----------

class TestDetectHierarchy:
    def test_flat_list(self):
        names = ["Task A", "Task B", "Task C"]
        result = detect_hierarchy(names)
        assert result == [(0, "Task A"), (0, "Task B"), (0, "Task C")]

    def test_indented_tasks(self):
        names = ["Phase 1", "  Task A", "  Task B", "Phase 2", "  Task C"]
        result = detect_hierarchy(names)
        assert result == [
            (0, "Phase 1"),
            (1, "Task A"),
            (1, "Task B"),
            (0, "Phase 2"),
            (1, "Task C"),
        ]

    def test_mixed_levels(self):
        names = ["Phase", "  Summary", "    Sub-task"]
        result = detect_hierarchy(names)
        assert result == [(0, "Phase"), (1, "Summary"), (2, "Sub-task")]

    def test_empty_and_none(self):
        names = [None, "", "Task"]
        result = detect_hierarchy(names)
        assert result == [(0, ""), (0, ""), (0, "Task")]


# ---------- TestCalculateDuration ----------

class TestCalculateDuration:
    def test_same_day(self):
        assert calculate_duration_from_dates("2025-01-06", "2025-01-06") == 1

    def test_weekdays_only(self):
        # Mon-Fri = 5 working days
        assert calculate_duration_from_dates("2025-01-06", "2025-01-10") == 5

    def test_across_weekend(self):
        # Mon to next Mon = 6 working days (Mon-Fri + Mon)
        assert calculate_duration_from_dates("2025-01-06", "2025-01-13") == 6

    def test_end_before_start(self):
        assert calculate_duration_from_dates("2025-01-10", "2025-01-06") is None

    def test_none_inputs(self):
        assert calculate_duration_from_dates(None, "2025-01-06") is None
        assert calculate_duration_from_dates("2025-01-06", None) is None

    def test_invalid_date(self):
        assert calculate_duration_from_dates("bad", "2025-01-06") is None


# ---------- TestAnalyzeWorkbook ----------

class TestAnalyzeWorkbook:
    def test_single_sheet(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Start", "Duration"],
                ["Task 1", "2025-01-06", 5],
                ["Task 2", "2025-01-13", 3],
            ]
        })
        result = analyze_workbook(data, "test.xlsx")
        assert len(result["sheets"]) == 1
        sheet = result["sheets"][0]
        assert sheet["name"] == "Tasks"
        assert sheet["row_count"] == 2
        assert sheet["columns"] == ["Task Name", "Start", "Duration"]
        assert len(sheet["sample_rows"]) == 2

    def test_multiple_sheets(self):
        data = _make_xlsx_bytes({
            "Tasks": [["Name"], ["A"]],
            "Milestones": [["Milestone"], ["M1"]],
        })
        result = analyze_workbook(data, "test.xlsx")
        assert len(result["sheets"]) == 2
        assert result["sheets"][0]["name"] == "Tasks"
        assert result["sheets"][1]["name"] == "Milestones"

    def test_empty_sheet(self):
        data = _make_xlsx_bytes({"Empty": []})
        result = analyze_workbook(data, "test.xlsx")
        assert result["sheets"][0]["row_count"] == 0
        assert result["sheets"][0]["columns"] == []

    def test_corrupt_file(self):
        with pytest.raises(ValueError, match="Failed to read"):
            analyze_workbook(b"not an excel file", "bad.xlsx")

    def test_xls_corrupt_file(self):
        with pytest.raises(ValueError, match="Failed to read .xls"):
            analyze_workbook(b"not a real xls file", "old.xls")

    def test_wrong_extension(self):
        with pytest.raises(ValueError, match="Unsupported"):
            analyze_workbook(b"dummy", "file.csv")

    def test_sample_rows_limited_to_5(self):
        rows = [["Name"]] + [[f"Task {i}"] for i in range(20)]
        data = _make_xlsx_bytes({"Tasks": rows})
        result = analyze_workbook(data, "test.xlsx")
        assert len(result["sheets"][0]["sample_rows"]) == 5

    def test_datetime_in_sample_formatted(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Name", "Date"],
                ["T1", datetime(2025, 3, 15)],
            ]
        })
        result = analyze_workbook(data, "test.xlsx")
        assert result["sheets"][0]["sample_rows"][0][1] == "2025-03-15"


# ---------- TestConvertExcelToMarkdown ----------

class TestConvertExcelToMarkdown:
    def test_flat_tasks(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Duration (days)"],
                ["Task 1", 3],
                ["Task 2", 5],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "duration": "Duration (days)",
        })
        assert result["task_count"] == 2
        assert "Task 1" in result["markdown"]
        assert "Task 2" in result["markdown"]
        assert "3d" in result["markdown"]

    def test_hierarchical_tasks(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Duration (days)"],
                ["Phase 1", 0],
                ["  Task A", 3],
                ["  Task B", 2],
                ["Phase 2", 0],
                ["    Deep Task", 1],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "duration": "Duration (days)",
        })
        assert result["phase_count"] >= 2
        md = result["markdown"]
        assert "Phase 1" in md
        assert "Phase 2" in md

    def test_with_resources(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Resources"],
                ["Phase 1", ""],
                ["  Task 1", "John Doe"],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "resources": "Resources",
        })
        assert "@john" in result["markdown"].lower()

    def test_with_dates(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Start", "Finish"],
                ["Phase 1", "", ""],
                ["  Task 1", datetime(2025, 1, 6), datetime(2025, 1, 10)],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "start_date": "Start",
            "end_date": "Finish",
        })
        # Duration should be calculated from dates (5 working days)
        assert "5d" in result["markdown"]

    def test_with_comments(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Comment"],
                ["Task 1", "Important note"],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "comment": "Comment",
        })
        assert "! Important note" in result["markdown"]

    def test_with_percent_complete(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "% Complete"],
                ["Task 1", 75],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "percent_complete": "% Complete",
        })
        assert "75%" in result["markdown"]

    def test_empty_rows_skipped(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name"],
                ["Task 1"],
                [None],
                [""],
                ["Task 2"],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
        })
        assert result["task_count"] == 2

    def test_date_validation_warnings(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Start", "Finish"],
                ["Task 1", datetime(2025, 1, 10), datetime(2025, 1, 6)],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "start_date": "Start",
            "end_date": "Finish",
        })
        assert len(result["warnings"]) > 0
        assert "after end date" in result["warnings"][0]

    def test_missing_task_name_mapping_raises(self):
        data = _make_xlsx_bytes({"Tasks": [["Name"], ["T1"]]})
        with pytest.raises(ValueError, match="task_name"):
            convert_excel_to_markdown(data, "test.xlsx", "Tasks", {})

    def test_sheet_not_found_raises(self):
        data = _make_xlsx_bytes({"Tasks": [["Name"], ["T1"]]})
        with pytest.raises(ValueError, match="not found"):
            convert_excel_to_markdown(data, "test.xlsx", "NonExistent", {
                "task_name": "Name",
            })

    def test_empty_sheet_raises(self):
        data = _make_xlsx_bytes({"Tasks": []})
        with pytest.raises(ValueError, match="empty"):
            convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
                "task_name": "Name",
            })

    def test_no_tasks_found_raises(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name"],
                [None],
                [""],
            ]
        })
        with pytest.raises(ValueError, match="No tasks found"):
            convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
                "task_name": "Task Name",
            })

    def test_front_matter_has_resources(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Resources"],
                ["Task 1", "Alice Smith"],
                ["Task 2", "Bob Jones"],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "resources": "Resources",
        })
        md = result["markdown"]
        assert "---" in md
        assert "Resources:" in md
        assert "@alice" in md.lower()

    def test_roundtrip_noodleplanner_export(self):
        """Test that a NoodlePlanner-style export can be re-imported."""
        data = _make_xlsx_bytes({
            "Tasks": [
                ["ID", "Task Name", "Start", "Finish", "Duration (days)",
                 "Resources", "% Complete", "RAG", "Comment"],
                [1, "Planning Phase", "2025-01-06", "2025-01-17", 10,
                 "", "", "", ""],
                [2, "  Gather Requirements", "2025-01-06", "2025-01-10", 5,
                 "John Doe", 100, "Green", "Done"],
                [3, "  Design Solution", "2025-01-13", "2025-01-17", 5,
                 "Jane Smith", 50, "Amber", "In progress"],
                [4, "Implementation Phase", "2025-01-20", "2025-01-31", 10,
                 "", "", "", ""],
                [5, "  Build Feature", "2025-01-20", "2025-01-31", 10,
                 "John Doe, Jane Smith", 0, "", ""],
            ]
        })
        result = convert_excel_to_markdown(data, "project.xlsx", "Tasks", {
            "task_name": "Task Name",
            "start_date": "Start",
            "end_date": "Finish",
            "duration": "Duration (days)",
            "resources": "Resources",
            "percent_complete": "% Complete",
            "comment": "Comment",
        })
        md = result["markdown"]
        assert result["phase_count"] >= 2
        assert "Planning Phase" in md
        assert "Gather Requirements" in md
        assert "100%" in md
        assert "! Done" in md
        assert result["warnings"] == []

    def test_multiple_resources_mapped(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Resources"],
                ["Task 1", "Alice Smith, Bob Jones"],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "resources": "Resources",
        })
        md = result["markdown"]
        assert "@alice" in md.lower()
        assert "@bob" in md.lower()

    def test_unparseable_date_produces_warning(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Start"],
                ["Task 1", "not-a-date"],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "start_date": "Start",
        })
        assert len(result["warnings"]) > 0
        assert "Could not parse" in result["warnings"][0]

    def test_negative_duration_warning(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Duration"],
                ["Task 1", -5],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "duration": "Duration",
        })
        assert any("Negative duration" in w for w in result["warnings"])

    def test_unparseable_duration_warning(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "Duration"],
                ["Task 1", "abc"],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "duration": "Duration",
        })
        assert any("Could not parse duration" in w for w in result["warnings"])

    def test_percent_out_of_range_warning(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "% Complete"],
                ["Task 1", 150],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "percent_complete": "% Complete",
        })
        assert any("outside 0-100 range" in w for w in result["warnings"])
        # Value should be clamped to 100
        assert "100%" in result["markdown"]

    def test_percent_as_decimal_converted(self):
        """Test that 0.75 is treated as 75%."""
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "% Complete"],
                ["Task 1", 0.75],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "percent_complete": "% Complete",
        })
        assert "75%" in result["markdown"]

    def test_unparseable_percent_warning(self):
        data = _make_xlsx_bytes({
            "Tasks": [
                ["Task Name", "% Complete"],
                ["Task 1", "abc"],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", {
            "task_name": "Task Name",
            "percent_complete": "% Complete",
        })
        assert any("Could not parse percent" in w for w in result["warnings"])


# ---------- TestNormalizeDateFallback ----------

class TestNormalizeDateFallback:
    """Tests for the dateutil fallback in normalize_date."""

    def test_month_name_with_spaces(self):
        result = normalize_date("March 15, 2025")
        assert result == "2025-03-15"

    def test_abbreviated_month_with_comma(self):
        result = normalize_date("Jan 6, 2025")
        assert result == "2025-01-06"

    def test_iso_with_time(self):
        result = normalize_date("2025-03-15T10:30:00")
        assert result == "2025-03-15"

    def test_truly_invalid_still_none(self):
        assert normalize_date("hello world") is None


# ---------- TestXlsSupport ----------

def _make_xls_bytes(sheets_data):
    """Helper: create an xls file in memory and return bytes.

    sheets_data: dict of {sheet_name: [row_list, ...]}
    Each row_list is a list of cell values; first row is headers.
    """
    import xlwt
    wb = xlwt.Workbook()
    for name, rows in sheets_data.items():
        ws = wb.add_sheet(name)
        for row_idx, row in enumerate(rows):
            for col_idx, value in enumerate(row):
                ws.write(row_idx, col_idx, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestXlsSupport:
    """Tests for .xls file support via xlrd."""

    def test_analyze_xls_workbook(self):
        data = _make_xls_bytes({
            "Tasks": [
                ["Task Name", "Duration"],
                ["Task 1", 5],
                ["Task 2", 3],
            ]
        })
        result = analyze_workbook(data, "test.xls")
        assert len(result["sheets"]) == 1
        sheet = result["sheets"][0]
        assert sheet["name"] == "Tasks"
        assert sheet["row_count"] == 2
        assert "Task Name" in sheet["columns"]

    def test_convert_xls_to_markdown(self):
        data = _make_xls_bytes({
            "Tasks": [
                ["Task Name", "Duration"],
                ["Phase 1", 0],
                ["  Task A", 3],
            ]
        })
        result = convert_excel_to_markdown(data, "test.xls", "Tasks", {
            "task_name": "Task Name",
            "duration": "Duration",
        })
        assert "Task A" in result["markdown"]
        assert "3d" in result["markdown"]

    def test_xls_multiple_sheets(self):
        data = _make_xls_bytes({
            "Tasks": [["Name"], ["A"]],
            "Milestones": [["Milestone"], ["M1"]],
        })
        result = analyze_workbook(data, "test.xls")
        assert len(result["sheets"]) == 2

    def test_analyze_xls_ignores_blank_data_rows(self):
        data = _make_xls_bytes({
            "Tasks": [
                ["Task Name", "Duration"],
                ["Task 1", 5],
                ["", ""],
                ["Task 2", 3],
                ["", ""],
            ]
        })
        result = analyze_workbook(data, "test.xls")
        sheet = result["sheets"][0]
        assert sheet["row_count"] == 2
