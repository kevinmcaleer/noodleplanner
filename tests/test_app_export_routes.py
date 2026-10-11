"""Routes and service paths the API tests never hit.

- ``POST /api/export-report-pptx`` (weekly report deck) and the
  ``PlanService.export_report_pptx`` method under it
- ``POST /api/benefits/export-excel`` (two-sheet benefits workbook)
- the ``settings:`` block of ``plan_service.parse_front_matter``
- ``ErrorSanitizationMiddleware`` on an unhandled exception, in both modes
- ``APIKeyAuthMiddleware`` with a key configured
"""

import io
from unittest.mock import patch

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from noodle_web import security
from noodle_web.app import app, plan_service
from noodle_web.plan_service import parse_front_matter
from noodle_web.security import APIKeyAuthMiddleware, ErrorSanitizationMiddleware


@pytest.fixture
def client():
    return TestClient(app)


def _slide_text(data: bytes) -> str:
    """All text on every slide, including table cells (milestones are a table)."""
    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# /api/export-report-pptx
# ---------------------------------------------------------------------------

REPORT = {
    "project_name": "Weekly Report Plan",
    "manager": "Kev",
    "sponsor": "Sponsor",
    "budget": "£10k",
    "date": "2026-03-06",
    "status": "amber",
    "milestones": [{"name": "Go live", "date": "2026-04-01", "rag": "green"}],
    "up_next": [{"name": "Build", "start": "2026-03-09", "finish": "2026-03-13", "rag": "amber"}],
    "highlight": {"date": "2026-03-05", "author": "Kev", "content": "Shipped the thing"},
    "risks_issues": [{"type": "risk", "title": "Vendor slips", "description": "d", "mitigation": "m", "score": 12}],
    "timeline_tasks": [
        {"name": "Build", "start": "2026-03-01", "finish": "2026-03-20", "percent": 40,
         "is_summary": True, "duration_days": 15},
    ],
}


class TestReportPptxRoute:
    def test_returns_a_deck_with_report_and_highlight_slides(self, client):
        resp = client.post("/api/export-report-pptx", json=REPORT)
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
        disposition = resp.headers["content-disposition"]
        assert "Weekly Report Plan - Report - " in disposition
        assert disposition.endswith('.pptx"') or ".pptx" in disposition

        prs = Presentation(io.BytesIO(resp.content))
        assert len(prs.slides) == 2  # report + highlight
        text = _slide_text(resp.content)
        assert "Weekly Report Plan" in text
        assert "Go live" in text
        assert "Shipped the thing" in text
        assert "Latest Highlight" in text

    def test_without_highlight_only_one_slide(self, client):
        body = dict(REPORT, highlight=None)
        resp = client.post("/api/export-report-pptx", json=body)
        assert resp.status_code == 200
        assert len(Presentation(io.BytesIO(resp.content)).slides) == 1

    def test_defaults_when_body_is_minimal(self, client):
        resp = client.post("/api/export-report-pptx", json={})
        assert resp.status_code == 200
        assert "Project - Report - " in resp.headers["content-disposition"]

    def test_validation_rejects_bad_types(self, client):
        resp = client.post("/api/export-report-pptx", json={"milestones": "not a list"})
        assert resp.status_code == 422

    def test_exporter_failure_is_a_sanitised_500(self, client):
        with patch("noodle_web.app.plan_service.export_report_pptx", side_effect=ValueError("boom")):
            with patch("noodle_web.app.is_production", return_value=True):
                resp = client.post("/api/export-report-pptx", json=REPORT)
        assert resp.status_code == 500
        assert resp.json()["detail"] == "Failed to export report"


class TestPlanServiceExportReport:
    def test_filename_and_bytes(self):
        result = plan_service.export_report_pptx(dict(REPORT))
        assert result.filename.startswith("Weekly Report Plan - Report - ")
        assert result.filename.endswith(".pptx")
        assert result.media_type.endswith("presentationml.presentation")
        assert result.content[:2] == b"PK"
        assert "Go live" in _slide_text(result.content)


# ---------------------------------------------------------------------------
# /api/benefits/export-excel
# ---------------------------------------------------------------------------

def _benefit(id_, type_, **kw):
    item = {
        "id": id_, "type": type_, "title": f"{type_} {id_}", "description": "desc",
        "objectiveType": "Strategic", "targetValue": "10", "currentValue": "4",
        "targetDate": "2026-12-31", "measurementMethod": "Survey", "linkedTo": [],
        "contributionPercent": 50, "status": "on-track", "lastUpdated": "2026-03-01",
        "score": 7,
    }
    item.update(kw)
    return item


class TestBenefitsExcelRoute:
    def test_two_sheets_with_tracking_limited_to_benefits(self, client):
        items = [
            _benefit(1, "objective"),
            _benefit(2, "benefit", linkedTo=[1]),
            _benefit(3, "disbenefit", linkedTo=[1, 2]),
            _benefit(4, "enabler"),
            _benefit(5, "change"),
        ]
        resp = client.post("/api/benefits/export-excel", json={"items": items, "project_name": "Ben"})
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert 'Ben-benefits.xlsx' in resp.headers["content-disposition"]

        wb = load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames == ["Benefits Map", "Tracking"]

        ws_map = wb["Benefits Map"]
        rows = list(ws_map.iter_rows(values_only=True))
        assert rows[0][:3] == ("ID", "Type", "Title")
        assert len(rows) == 1 + 5
        assert rows[1][1] == "Objective"
        assert rows[3][9] == "1, 2"          # linkedTo joined
        assert rows[2][10] == 50             # contribution %

        ws_track = wb["Tracking"]
        track = list(ws_track.iter_rows(values_only=True))
        assert track[0][0] == "ID"
        assert [r[1] for r in track[1:]] == ["Benefit", "Disbenefit"]
        assert len(track) == 3

    def test_default_project_name(self, client):
        resp = client.post("/api/benefits/export-excel", json={"items": []})
        assert resp.status_code == 200
        assert "Benefits-benefits.xlsx" in resp.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(resp.content))
        assert list(wb["Tracking"].iter_rows(values_only=True))[0][0] == "ID"

    def test_rejects_unknown_type_and_bad_percent(self, client):
        bad_type = client.post("/api/benefits/export-excel", json={"items": [_benefit(1, "wish")]})
        assert bad_type.status_code == 422
        bad_pct = client.post("/api/benefits/export-excel", json={"items": [_benefit(1, "benefit", contributionPercent=120)]})
        assert bad_pct.status_code == 422


# ---------------------------------------------------------------------------
# parse_front_matter settings block
# ---------------------------------------------------------------------------

class TestParseFrontMatterSettings:
    def test_settings_block_is_nested_and_booleans_converted(self):
        plan = (
            "---\n"
            "title: Settings Plan\n"
            "settings:\n"
            "  weekends: TRUE\n"
            "  theme: dark\n"
            "  autosave: false\n"
            "owner: Kev\n"
            "---\n"
            "Phase\n  Task 1d\n"
        )
        fm = parse_front_matter(plan)
        assert fm["title"] == "Settings Plan"
        assert fm["settings"] == {"weekends": True, "theme": "dark", "autosave": False}
        # The first un-indented line ends the block and is parsed normally.
        assert fm["owner"] == "Kev"
        assert "weekends" not in fm

    def test_empty_settings_block_is_omitted(self):
        fm = parse_front_matter("---\nsettings:\nowner: Kev\n---\n")
        assert "settings" not in fm
        assert fm["owner"] == "Kev"

    def test_no_front_matter(self):
        assert parse_front_matter("Phase\n  Task 1d\n") == {}

    def test_keys_are_lowercased(self):
        assert parse_front_matter("---\nProject Manager: Kev\n---\n")["project manager"] == "Kev"


# ---------------------------------------------------------------------------
# Middleware behaviour on failures
# ---------------------------------------------------------------------------

def _crashing_app(*middlewares):
    mini = FastAPI()

    @mini.get("/boom")
    def boom():
        raise RuntimeError("secret internals")

    @mini.get("/ok")
    def ok():
        return {"ok": True}

    @mini.get("/health")
    def health():
        return {"status": "ok"}

    for mw in middlewares:
        mini.add_middleware(mw)
    return mini


class TestErrorSanitizationMiddleware:
    def test_production_returns_generic_500_json(self):
        client = TestClient(_crashing_app(ErrorSanitizationMiddleware), raise_server_exceptions=False)
        with patch.object(security, "ENVIRONMENT", "production"):
            resp = client.get("/boom")
        assert resp.status_code == 500
        assert resp.json() == {"detail": "An internal error occurred."}
        assert "secret internals" not in resp.text

    def test_development_re_raises(self):
        client = TestClient(_crashing_app(ErrorSanitizationMiddleware), raise_server_exceptions=True)
        with patch.object(security, "ENVIRONMENT", "development"):
            with pytest.raises(RuntimeError, match="secret internals"):
                client.get("/boom")

    def test_successful_requests_pass_through(self):
        client = TestClient(_crashing_app(ErrorSanitizationMiddleware))
        assert client.get("/ok").json() == {"ok": True}


class TestAPIKeyAuthMiddleware:
    def test_requires_bearer_when_key_configured(self):
        client = TestClient(_crashing_app(APIKeyAuthMiddleware))
        with patch.object(security, "API_KEY", "s3cret"):
            assert client.get("/ok").status_code == 401
            assert client.get("/ok", headers={"Authorization": "Bearer wrong"}).status_code == 401
            assert client.get("/ok", headers={"Authorization": "Basic s3cret"}).status_code == 401
            assert client.get("/ok", headers={"Authorization": "Bearer s3cret"}).status_code == 200
            # Health checks stay public.
            assert client.get("/health").status_code == 200

    def test_open_when_no_key_configured(self):
        client = TestClient(_crashing_app(APIKeyAuthMiddleware))
        with patch.object(security, "API_KEY", ""):
            assert client.get("/ok").status_code == 200
