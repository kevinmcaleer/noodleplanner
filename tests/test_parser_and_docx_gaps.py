"""Front-matter parser branches and the comms-plan Word export.

``FrontMatterParser.parse_named_non_working_days`` had no test for bare dates
or unparseable entries, ``parse_dependencies`` none for the multi-line
``key: value`` continuation form or for a non-list line ending the section,
and ``export_comms_to_docx`` was never called at all.
"""

import io

from docx import Document

from noodle_core import FrontMatterParser
from noodle_core.format_converter import export_comms_to_docx


# ---------------------------------------------------------------------------
# parse_named_non_working_days
# ---------------------------------------------------------------------------

class TestNamedNonWorkingDays:
    def test_bare_dates_ranges_and_junk(self):
        plan = (
            "---\n"
            "title: NWD\n"
            "non-working-days:\n"
            "- 2026-12-25\n"
            "- Christmas shutdown: 2026-12-24:2027-01-02\n"
            "- 2026-05-04\n"
            "- some day in June\n"
            "- 2026-5-4\n"
            "# a comment line does not end the list\n"
            "- 2026-08-31\n"
            "owner: Kev\n"
            "- 2026-09-01\n"
            "---\n"
            "Phase\n  Task 1d\n"
        )
        entries = FrontMatterParser(plan).parse_named_non_working_days()
        assert entries == [
            {"name": "", "start": "2026-12-25", "finish": ""},
            {"name": "Christmas shutdown", "start": "2026-12-24", "finish": "2027-01-02"},
            {"name": "", "start": "2026-05-04", "finish": ""},
            {"name": "", "start": "2026-08-31", "finish": ""},
        ]

    def test_holidays_alias_and_absent(self):
        plan = "---\nholidays:\n- Bank holiday: 2026-05-25\n---\n"
        assert FrontMatterParser(plan).parse_named_non_working_days() == [
            {"name": "Bank holiday", "start": "2026-05-25", "finish": ""},
        ]
        assert FrontMatterParser("---\ntitle: x\n---\n").parse_named_non_working_days() == []
        assert FrontMatterParser("Phase\n  Task 1d\n").parse_named_non_working_days() == []


# ---------------------------------------------------------------------------
# parse_dependencies
# ---------------------------------------------------------------------------

class TestDependenciesContinuationForm:
    def test_multi_line_entries_with_continuation_keys(self):
        plan = (
            "---\n"
            "title: Deps\n"
            "dependencies:\n"
            "  - from: Project A\n"
            "    task: Milestone 1\n"
            "    to_task: Design Start\n"
            "    type: SS\n"
            "    lag: 3\n"
            "  - from: Project B\n"
            "    TASK: Handover\n"
            "    to_task: Build Start\n"
            "    lag: not-a-number\n"
            "    colour: ignored\n"
            "  - from: Project C\n"
            "    task: T\n"
            "    to_task: U\n"
            "---\n"
            "Phase\n  Task 1d\n"
        )
        deps = FrontMatterParser(plan).parse_dependencies()
        assert deps == [
            {"from": "Project A", "task": "Milestone 1", "to_task": "Design Start", "type": "SS", "lag": 3},
            {"from": "Project B", "task": "Handover", "to_task": "Build Start", "type": "FS", "lag": 0},
            {"from": "Project C", "task": "T", "to_task": "U", "type": "FS", "lag": 0},
        ]

    def test_inline_lag_that_is_not_an_integer_defaults_to_zero(self):
        plan = "---\ndependencies:\n- lag: 2.5\n  from: X\n---\n"
        deps = FrontMatterParser(plan).parse_dependencies()
        assert deps[0]["lag"] == 0
        assert deps[0]["from"] == "X"

    def test_section_ends_at_a_plain_line(self):
        plan = (
            "---\n"
            "dependencies:\n"
            "- from: A\n"
            "  task: T1\n"
            "Just some text\n"
            "- from: Not A Dependency\n"
            "---\n"
        )
        deps = FrontMatterParser(plan).parse_dependencies()
        assert [d["from"] for d in deps] == ["A"]
        assert deps[0]["task"] == "T1"

    def test_a_following_key_value_line_does_not_end_the_section(self):
        # A ``key: value`` line after the list is treated as a continuation
        # of the current entry when it names a known key, and ignored
        # otherwise; either way it does not terminate the section.
        plan = (
            "---\n"
            "dependencies:\n"
            "- from: A\n"
            "owner: Kev\n"
            "- from: B\n"
            "---\n"
        )
        deps = FrontMatterParser(plan).parse_dependencies()
        assert [d["from"] for d in deps] == ["A", "B"]


# ---------------------------------------------------------------------------
# export_comms_to_docx
# ---------------------------------------------------------------------------

def _doc(data: bytes) -> Document:
    return Document(io.BytesIO(data))


class TestCommsDocx:
    def test_empty_plan_writes_a_placeholder(self):
        doc = _doc(export_comms_to_docx([], project_name="Empty"))
        paragraphs = [p.text for p in doc.paragraphs]
        assert "Empty — Communications Plan" in paragraphs
        assert "No communications plan items defined." in paragraphs
        assert doc.tables == []

    def test_landscape_table_with_one_row_per_item(self):
        items = [
            {"id": 1, "activity": "Steering update", "audience": "Sponsor", "content": "Status",
             "frequency": "Monthly", "channel": "Meeting", "owner": "Kev", "status": "active"},
            {"activity": "Team standup", "audience": "Team"},
        ]
        doc = _doc(export_comms_to_docx(items, project_name="Comms"))

        section = doc.sections[0]
        assert section.page_width > section.page_height  # landscape

        assert len(doc.tables) == 1
        table = doc.tables[0]
        header = [c.text for c in table.rows[0].cells]
        assert header == ["#", "Activity", "Audience", "Content", "Frequency", "Channel", "Owner", "Status"]
        assert len(table.rows) == 3
        first = [c.text for c in table.rows[1].cells]
        assert first == ["1", "Steering update", "Sponsor", "Status", "Monthly", "Meeting", "Kev", "active"]
        second = [c.text for c in table.rows[2].cells]
        # Missing id falls back to the 1-based position; missing fields are blank.
        assert second == ["2", "Team standup", "Team", "", "", "", "", ""]

    def test_default_project_name(self):
        doc = _doc(export_comms_to_docx([{"activity": "x"}]))
        assert doc.paragraphs[0].text == "Project — Communications Plan"
