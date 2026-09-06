"""Cross-runtime parity checks for the browser and Python Excel exporters."""

import json
import subprocess
from datetime import date, datetime
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from noodle_core import extract_budget, parse_budget_markdown
from noodle_web.app import app


ROOT = Path(__file__).resolve().parent.parent
PLAN = ROOT / "tests" / "fixtures" / "roundtrip" / "kitchen-sink.md"
NODE_EXPORT = ROOT / "tests" / "helpers" / "build_browser_excel.mjs"
NODE_IMPORT = ROOT / "tests" / "helpers" / "import_browser_excel.mjs"


def _normalise(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float):
        return round(value, 6)
    return value


def _colour(colour):
    if colour is None:
        return None
    if colour.type == "rgb":
        return colour.rgb[-6:].upper() if colour.rgb else None
    if colour.type == "indexed":
        return ("indexed", colour.indexed)
    if colour.type == "theme":
        return ("theme", colour.theme, colour.tint)
    return None


def _style(cell):
    return {
        "fill": (cell.fill.fill_type, _colour(cell.fill.fgColor)),
        "font": (cell.font.bold, _colour(cell.font.color), cell.font.sz),
        "alignment": (
            cell.alignment.horizontal,
            cell.alignment.vertical,
            cell.alignment.text_rotation,
        ),
        "number_format": cell.number_format,
    }


def _column_width(sheet, column_letter):
    index = column_index_from_string(column_letter)
    for dimension in sheet.column_dimensions.values():
        if dimension.min <= index <= dimension.max:
            return dimension.width
    return sheet.sheet_format.defaultColWidth


def test_browser_workbook_matches_python_export_cell_for_cell(tmp_path):
    plan_text = PLAN.read_text(encoding="utf-8")
    client = TestClient(app)
    parse_response = client.post(
        "/api/parse",
        json={"plan_text": plan_text, "project_name": "Kitchen Sink"},
    )
    assert parse_response.status_code == 200
    parse_result = parse_response.json()
    assert parse_result["success"] is True

    budget_text = extract_budget(plan_text)
    budget_items = parse_budget_markdown(budget_text)
    payload_path = tmp_path / "browser-export.json"
    browser_path = tmp_path / "browser.xlsx"
    python_path = tmp_path / "python.xlsx"
    payload_path.write_text(
        json.dumps(
            {
                "parseResult": parse_result,
                "budgetItems": budget_items,
                "projectName": "Kitchen Sink",
            }
        ),
        encoding="utf-8",
    )
    subprocess.run(
        ["node", str(NODE_EXPORT), str(payload_path), str(browser_path)],
        cwd=ROOT,
        check=True,
    )

    python_response = client.post(
        "/render",
        json={
            "plan_text": plan_text,
            "project_name": "Kitchen Sink",
            "export_excel": True,
        },
    )
    assert python_response.status_code == 200
    python_path.write_bytes(python_response.content)

    browser = load_workbook(browser_path)
    python = load_workbook(python_path)
    assert browser.sheetnames == python.sheetnames

    volatile_cells = {("Summary", "B8")}
    for sheet_name in python.sheetnames:
        browser_sheet = browser[sheet_name]
        python_sheet = python[sheet_name]
        assert browser_sheet.max_row == python_sheet.max_row, sheet_name
        assert browser_sheet.max_column == python_sheet.max_column, sheet_name

        for row in python_sheet.iter_rows():
            for python_cell in row:
                address = (sheet_name, python_cell.coordinate)
                if address in volatile_cells:
                    continue
                browser_cell = browser_sheet[python_cell.coordinate]
                assert _normalise(browser_cell.value) == _normalise(
                    python_cell.value
                ), address
                if python_cell.value is not None or browser_cell.value is not None:
                    assert _style(browser_cell) == _style(python_cell), address

        for column_letter, dimension in python_sheet.column_dimensions.items():
            assert _column_width(browser_sheet, column_letter) == dimension.width, (
                sheet_name,
                column_letter,
            )


def _browser_import(kind, workbook_path):
    result = subprocess.run(
        ["node", str(NODE_IMPORT), kind, str(workbook_path)],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(result.stdout)


def test_targeted_browser_imports_match_python_routes(tmp_path):
    client = TestClient(app)
    raid_export = client.post(
        "/api/raid/export-excel",
        json={
            "project_name": "Parity",
            "items": [
                {
                    "id": 7,
                    "type": "issue",
                    "title": "Blocked",
                    "description": "Waiting",
                    "raised_by": "Jane",
                    "owner": "Quinn",
                    "mitigation_actions": "Escalate",
                    "impact": 4,
                    "likelihood": 5,
                    "score": 20,
                    "status": "transferred",
                }
            ],
        },
    )
    assert raid_export.status_code == 200
    raid_path = tmp_path / "raid.xlsx"
    raid_path.write_bytes(raid_export.content)
    raid_python = client.post(
        "/api/raid/import-excel",
        files={"file": ("raid.xlsx", raid_export.content)},
    ).json()
    assert _browser_import("raid", raid_path) == raid_python

    budget_export = client.post(
        "/api/budget/export-excel",
        json={
            "project_name": "Parity",
            "items": [
                {
                    "id": 3,
                    "description": "Laptop",
                    "estimate": 12.5,
                    "forecast": 15,
                    "type": "Capex",
                    "invoice": "I-3",
                    "po": "P-3",
                    "supplier": "Shop",
                    "total": 9.25,
                    "date_ordered": "2026-01-04",
                    "date_received": "2026-01-10",
                    "category": "Hardware",
                }
            ],
        },
    )
    assert budget_export.status_code == 200
    budget_path = tmp_path / "budget.xlsx"
    budget_path.write_bytes(budget_export.content)
    budget_python = client.post(
        "/api/budget/import-excel",
        files={"file": ("budget.xlsx", budget_export.content)},
    ).json()
    assert _browser_import("budget", budget_path) == budget_python
