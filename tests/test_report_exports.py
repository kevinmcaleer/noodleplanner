"""Excel and PowerPoint export of the Tasks by Assignment and Slippage reports (#776).

The browser computes the reports (static/plan-reports.js) and posts the rows;
these tests post rows shaped as views-reports.js sends them and open the files
that come back.

Run with: uv run pytest tests/test_report_exports.py -q
"""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from noodle_web.app import app

ASSIGNMENTS = {
    "groups": [
        {"name": "Unassigned", "role": "", "unassigned": True,
         "totals": {"tasks": 1, "open": 1, "complete": 0, "workDays": 2, "nextDue": "2026-03-10", "overdue": 0},
         "tasks": [{"name": "Docs", "phase": "Build", "start": "2026-03-09", "finish": "2026-03-10",
                    "duration": 2, "percent": 0, "status": "open", "shared_with": []}]},
        {"name": "Alex Chen", "role": "Developer", "unassigned": False,
         "totals": {"tasks": 2, "open": 1, "complete": 1, "workDays": 3.5, "nextDue": None, "overdue": 1},
         "tasks": [{"name": "Review", "phase": "Design", "start": "2026-03-05", "finish": "2026-03-05",
                    "duration": 1, "percent": 100, "status": "complete", "shared_with": ["Sam Lee"]},
                   {"name": "API", "phase": "Build", "start": "2026-03-06", "finish": "2026-03-12",
                    "duration": 5, "percent": 20, "status": "overdue", "shared_with": []}]},
    ],
}


def _row(name, variance, critical=False):
    return {"name": name, "phase": "Build", "critical": critical, "baseline_start": "2026-03-06",
            "baseline_finish": "2026-03-12", "start": "2026-03-06", "finish": "2026-03-17",
            "start_variance": 0, "finish_variance": variance,
            "status": "slipped" if variance > 0 else "on-track"}


SLIPPAGE = {
    "baseline": "Sign-off (2026-03-01)",
    "project": {"baseline_finish": "2026-03-12", "finish": "2026-03-17", "variance": 3, "status": "slipped"},
    "counts": {"slipped": 1, "onTrack": 1, "pulledForward": 0, "added": 1, "removed": 1},
    "critical": [_row("API", 3, critical=True)],
    "phases": [dict(_row("Build", 3), level=1, slipped_tasks=1, worst_slip=3)],
    "tasks": [_row("API", 3, critical=True), _row("Docs", 0)],
    "added": [{"name": "Training", "phase": "Build"}],
    "removed": [{"name": "Old task"}],
}


@pytest.fixture
def client():
    return TestClient(app)


def export(client, report, fmt, data):
    response = client.post("/api/reports/export", json={
        "report": report, "format": fmt, "project_name": "Launch Plan", "data": data})
    assert response.status_code == 200, response.text
    return response


def test_assignments_excel_has_a_summary_and_a_sheet_per_person(client):
    response = export(client, "assignments", "xlsx", ASSIGNMENTS)
    assert 'filename="Launch_Plan_tasks_by_assignment.xlsx"' in response.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Summary", "Unassigned", "Alex Chen"]
    summary = [[c.value for c in row] for row in wb["Summary"].iter_rows()]
    assert summary[0][:3] == ["Person", "Role", "Tasks"]
    assert summary[1][0] == "Unassigned"
    assert summary[2][:3] == ["Alex Chen", "Developer", 2]
    alex = [[c.value for c in row] for row in wb["Alex Chen"].iter_rows()]
    assert alex[1][0] == "Review" and alex[1][7] == "Sam Lee"
    assert alex[2][6] == "Overdue"


def test_assignments_powerpoint_has_an_overview_and_a_slide_per_person(client):
    deck = Presentation(io.BytesIO(export(client, "assignments", "pptx", ASSIGNMENTS).content))
    titles = [s.shapes[0].text_frame.text for s in deck.slides]
    assert titles[0] == "Launch Plan — Tasks by assignment"
    assert titles[1:] == ["Unassigned", "Alex Chen"]


def test_slippage_excel_sheets(client):
    response = export(client, "slippage", "xlsx", SLIPPAGE)
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Summary", "Critical path", "Phases", "Tasks", "Scope changes"]
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Variance (working days)"] == 3
    assert summary["Compared with"] == "Sign-off (2026-03-01)"
    tasks = [[c.value for c in row] for row in wb["Tasks"].iter_rows()]
    assert tasks[1][0] == "API" and tasks[1][7] == "+3d" and tasks[1][9] == "Yes"
    scope = [[c.value for c in row] for row in wb["Scope changes"].iter_rows(min_row=2)]
    assert scope == [["Training", "Build", "Added"], ["Old task", None, "Removed"]]


def test_slippage_powerpoint_leads_with_the_headline(client):
    deck = Presentation(io.BytesIO(export(client, "slippage", "pptx", SLIPPAGE).content))
    first = " ".join(shape.text_frame.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    assert "Launch Plan — Slippage" in first
    assert "+3d working days" in first
    titles = [s.shapes[0].text_frame.text for s in deck.slides]
    assert "Slippage on the critical path" in titles
    assert "Changed since the baseline" in titles


def test_long_tables_paginate(client):
    many = dict(SLIPPAGE, tasks=[_row(f"Task {i}", i % 3) for i in range(30)])
    deck = Presentation(io.BytesIO(export(client, "slippage", "pptx", many).content))
    titles = [s.shapes[0].text_frame.text for s in deck.slides]
    assert [t for t in titles if t.startswith("Every task")] == ["Every task (1/3)", "Every task (2/3)", "Every task (3/3)"]


def test_bad_requests_are_rejected(client):
    assert client.post("/api/reports/export", json={"report": "nope", "format": "xlsx", "data": {}}).status_code == 422
    assert client.post("/api/reports/export", json={"report": "slippage", "format": "csv", "data": {}}).status_code == 422
    broken = client.post("/api/reports/export", json={"report": "assignments", "format": "xlsx",
                                                       "data": {"groups": [{"name": "x"}]}})
    assert broken.status_code == 422
