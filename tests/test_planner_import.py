"""Tests for Microsoft Planner import functionality."""

import io
import pytest
from datetime import datetime
from openpyxl import Workbook

from noodle_core.excel_importer import (
    parse_planner_duration,
    parse_planner_dependency,
    detect_planner_worksheet,
    parse_planner_header,
    convert_planner_to_markdown,
    convert_excel_to_markdown,
    analyze_workbook,
    _read_workbook,
    _is_task_header_row,
    _split_resource_names,
)


def _make_planner_xlsx_bytes(
    header_fields=None,
    task_columns=None,
    task_rows=None,
    sheet_name="Project tasks",
):
    """Helper: create a Planner-style xlsx file in memory.

    Creates a workbook with a header section followed by a task data table,
    matching the Microsoft Planner export format.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Default header fields
    if header_fields is None:
        header_fields = [
            ("Task Number", "My Project"),
            ("Plan Owner", "John Smith"),
            ("Project Start Date", datetime(2025, 1, 6)),
            ("Project Finish Date", datetime(2025, 6, 30)),
            ("% Complete", 0.45),
            ("Exported on", datetime(2025, 2, 15)),
        ]

    # Write header fields
    for label, value in header_fields:
        ws.append([label, value])

    # Blank row separator
    ws.append([])

    # Default task columns
    if task_columns is None:
        task_columns = [
            "ID", "Task Name", "Duration", "Start", "Finish",
            "Predecessors", "Resource Names", "% Complete", "Outline Level",
        ]

    ws.append(task_columns)

    # Default task rows
    if task_rows is None:
        task_rows = [
            [1, "Planning", "10 days", datetime(2025, 1, 6), datetime(2025, 1, 17), "", "Alice", 1.0, 1],
            [2, "Requirements", "5 days", datetime(2025, 1, 6), datetime(2025, 1, 10), "", "Alice", 1.0, 2],
            [3, "Design", "5 days", datetime(2025, 1, 13), datetime(2025, 1, 17), "2FS", "Bob", 0.5, 2],
            [4, "Implementation", "15 days", datetime(2025, 1, 20), datetime(2025, 2, 7), "", "", 0.0, 1],
            [5, "Build Feature", "10 days", datetime(2025, 1, 20), datetime(2025, 1, 31), "3FS", "Alice, Bob", 0.0, 2],
            [6, "Testing", "5 days", datetime(2025, 2, 3), datetime(2025, 2, 7), "5FS", "Charlie", 0.0, 2],
        ]

    for row in task_rows:
        ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- TestParsePlannerDuration ----------


class TestParsePlannerDuration:
    def test_days_singular(self):
        assert parse_planner_duration("1 day") == (1, "d")

    def test_days_plural(self):
        assert parse_planner_duration("5 days") == (5, "d")

    def test_weeks_singular(self):
        assert parse_planner_duration("1 week") == (1, "w")

    def test_weeks_plural(self):
        assert parse_planner_duration("2 weeks") == (2, "w")

    def test_months_singular(self):
        assert parse_planner_duration("1 month") == (1, "m")

    def test_months_plural(self):
        assert parse_planner_duration("3 months") == (3, "m")

    def test_abbreviated_weeks(self):
        assert parse_planner_duration("2 wks") == (2, "w")

    def test_abbreviated_months(self):
        assert parse_planner_duration("3 mons") == (3, "m")

    def test_numeric_only(self):
        assert parse_planner_duration("5") == (5, "d")

    def test_numeric_float(self):
        assert parse_planner_duration("5.0") == (5, "d")

    def test_with_question_mark(self):
        assert parse_planner_duration("3 days?") == (3, "d")

    def test_none_returns_none(self):
        assert parse_planner_duration(None) is None

    def test_empty_returns_none(self):
        assert parse_planner_duration("") is None

    def test_whitespace_only_returns_none(self):
        assert parse_planner_duration("   ") is None

    def test_invalid_text_returns_none(self):
        assert parse_planner_duration("not a duration") is None

    def test_float_days(self):
        assert parse_planner_duration("1.5 days") == (1, "d")

    def test_zero_days(self):
        assert parse_planner_duration("0 days") == (0, "d")

    def test_negative_returns_none(self):
        assert parse_planner_duration("-5") is None


# ---------- TestParsePlannerDependency ----------


class TestParsePlannerDependency:
    def test_single_fs_dependency(self):
        task_map = {2: "Design"}
        deps, warnings = parse_planner_dependency("2FS", task_map)
        assert deps == ["Design"]
        assert warnings == []

    def test_implicit_fs(self):
        task_map = {2: "Design"}
        deps, warnings = parse_planner_dependency("2", task_map)
        assert deps == ["Design"]
        assert warnings == []

    def test_multiple_dependencies(self):
        task_map = {2: "Design", 5: "Build"}
        deps, warnings = parse_planner_dependency("2FS,5FS", task_map)
        assert deps == ["Design", "Build"]
        assert warnings == []

    def test_ss_dependency_warns(self):
        task_map = {3: "Testing"}
        deps, warnings = parse_planner_dependency("3SS", task_map)
        assert deps == ["Testing"]
        assert any("not supported" in w for w in warnings)

    def test_ff_dependency_warns(self):
        task_map = {3: "Testing"}
        deps, warnings = parse_planner_dependency("3FF", task_map)
        assert deps == ["Testing"]
        assert any("not supported" in w for w in warnings)

    def test_sf_dependency_warns(self):
        task_map = {3: "Testing"}
        deps, warnings = parse_planner_dependency("3SF", task_map)
        assert deps == ["Testing"]
        assert any("not supported" in w for w in warnings)

    def test_unknown_task_number(self):
        task_map = {2: "Design"}
        deps, warnings = parse_planner_dependency("99FS", task_map)
        assert deps == []
        assert any("unknown task number" in w for w in warnings)

    def test_none_returns_empty(self):
        deps, warnings = parse_planner_dependency(None, {})
        assert deps == []
        assert warnings == []

    def test_empty_returns_empty(self):
        deps, warnings = parse_planner_dependency("", {})
        assert deps == []
        assert warnings == []

    def test_invalid_format(self):
        deps, warnings = parse_planner_dependency("abc", {})
        assert deps == []
        assert any("Could not parse" in w for w in warnings)

    def test_multiple_with_spaces(self):
        task_map = {2: "Design", 5: "Build"}
        deps, warnings = parse_planner_dependency("2FS, 5FS", task_map)
        assert deps == ["Design", "Build"]

    def test_lag_positive_days(self):
        task_map = {2: "Design"}
        deps, warnings = parse_planner_dependency("2FS+3d", task_map)
        assert deps == ["Design"]
        assert any("Lag/lead" in w for w in warnings)

    def test_lag_negative(self):
        task_map = {2: "Design"}
        deps, warnings = parse_planner_dependency("2FS-2", task_map)
        assert deps == ["Design"]
        assert any("Lag/lead" in w for w in warnings)

    def test_lag_with_weeks(self):
        task_map = {2: "Design"}
        deps, warnings = parse_planner_dependency("2FS+1w", task_map)
        assert deps == ["Design"]
        assert any("Lag/lead" in w for w in warnings)

    def test_implicit_fs_with_lag(self):
        task_map = {2: "Design"}
        deps, warnings = parse_planner_dependency("2+3d", task_map)
        assert deps == ["Design"]
        assert any("Lag/lead" in w for w in warnings)

    def test_multiple_with_lag(self):
        task_map = {2: "Design", 5: "Build"}
        deps, warnings = parse_planner_dependency("2FS+1d,5FS-2d", task_map)
        assert deps == ["Design", "Build"]
        lag_warnings = [w for w in warnings if "Lag/lead" in w]
        assert len(lag_warnings) == 2


# ---------- TestDetectPlannerWorksheet ----------


class TestDetectPlannerWorksheet:
    def test_detects_project_tasks_sheet(self):
        data = _make_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            result = detect_planner_worksheet(wb)
            assert result == "Project tasks"
        finally:
            wb.close()

    def test_no_planner_sheet(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tasks"
        ws.append(["Name"])
        result = detect_planner_worksheet(wb)
        assert result is None
        wb.close()

    def test_case_insensitive_detection(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "project tasks"
        ws.append(["Task Number", "My Project"])
        result = detect_planner_worksheet(wb)
        assert result == "project tasks"
        wb.close()


# ---------- TestParsePlannerHeader ----------


class TestParsePlannerHeader:
    def test_parses_project_name(self):
        data = _make_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Project tasks"]
            header, task_row = parse_planner_header(ws)
            assert header["project_name"] == "My Project"
        finally:
            wb.close()

    def test_parses_plan_owner(self):
        data = _make_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Project tasks"]
            header, task_row = parse_planner_header(ws)
            assert header["plan_owner"] == "John Smith"
        finally:
            wb.close()

    def test_parses_start_date(self):
        data = _make_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Project tasks"]
            header, task_row = parse_planner_header(ws)
            assert header["start_date"] == "2025-01-06"
        finally:
            wb.close()

    def test_finds_task_header_row(self):
        data = _make_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Project tasks"]
            header, task_row = parse_planner_header(ws)
            assert task_row is not None
            rows = list(ws.iter_rows(values_only=True))
            header_cells = [str(c).strip().lower() for c in rows[task_row] if c is not None]
            assert "task name" in header_cells
        finally:
            wb.close()


# ---------- TestAnalyzeWorkbookPlannerDetection ----------


class TestAnalyzeWorkbookPlannerDetection:
    def test_detects_planner_format(self):
        data = _make_planner_xlsx_bytes()
        result = analyze_workbook(data, "test.xlsx")
        assert result.get("is_planner") is True

    def test_planner_columns_use_task_headers(self):
        """analyze_workbook should return task column headers, not metadata rows."""
        data = _make_planner_xlsx_bytes()
        result = analyze_workbook(data, "test.xlsx")
        sheet = result["sheets"][0]
        columns_lower = [c.lower() for c in sheet["columns"]]
        # Should have task data columns, not metadata labels
        assert "task name" in columns_lower
        assert "duration" in columns_lower
        # Should NOT have metadata labels as columns
        assert "task number" not in columns_lower
        assert "plan owner" not in columns_lower
        # Sample rows should be actual task data, not metadata values
        assert len(sheet["sample_rows"]) > 0
        all_cells = [cell for row in sheet["sample_rows"] for cell in row]
        assert any("Planning" in str(c) or "Requirements" in str(c) for c in all_cells)

    def test_non_planner_not_flagged(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tasks"
        ws.append(["Task Name", "Duration"])
        ws.append(["Task 1", 5])
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()

        result = analyze_workbook(data, "test.xlsx")
        assert "is_planner" not in result


# ---------- TestConvertPlannerToMarkdown ----------


class TestConvertPlannerToMarkdown:
    def test_basic_conversion(self):
        data = _make_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        assert result["task_count"] > 0
        assert result["phase_count"] > 0
        assert "Planning" in result["markdown"]
        assert "Requirements" in result["markdown"]

    def test_front_matter_includes_project_metadata(self):
        data = _make_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "title: My Project" in md
        assert "manager: John Smith" in md
        assert "start: 2025-01-06" in md

    def test_durations_parsed_correctly(self):
        data = _make_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "10d" in md
        assert "5d" in md

    def test_dependencies_converted(self):
        data = _make_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "[depends Requirements]" in md
        assert "[depends Design]" in md
        assert "[depends Build Feature]" in md

    def test_resources_mapped(self):
        data = _make_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "@alice" in md.lower()
        assert "@bob" in md.lower()
        assert "@charlie" in md.lower()

    def test_percent_complete_converted(self):
        data = _make_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "100%" in md
        assert "50%" in md

    def test_hierarchy_from_outline_level(self):
        data = _make_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        # Phase (level 1 = indent 0) should not be indented
        lines = md.split("\n")
        planning_lines = [l for l in lines if "Planning" in l and "Implementation" not in l]
        assert any(not l.startswith("  ") for l in planning_lines)
        # Sub-tasks (level 2 = indent 1) should be indented
        req_lines = [l for l in lines if "Requirements" in l]
        assert any(l.startswith("  ") for l in req_lines)

    def test_week_duration_format(self):
        task_rows = [
            [1, "Phase 1", "2 weeks", datetime(2025, 1, 6), datetime(2025, 1, 17), "", "", 0.0, 1],
            [2, "Task A", "1 week", datetime(2025, 1, 6), datetime(2025, 1, 10), "", "", 0.0, 2],
        ]
        data = _make_planner_xlsx_bytes(task_rows=task_rows)
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "2w" in md
        assert "1w" in md

    def test_month_duration_format(self):
        task_rows = [
            [1, "Phase 1", "3 months", datetime(2025, 1, 6), datetime(2025, 3, 31), "", "", 0.0, 1],
            [2, "Task A", "1 month", datetime(2025, 1, 6), datetime(2025, 2, 3), "", "", 0.0, 2],
        ]
        data = _make_planner_xlsx_bytes(task_rows=task_rows)
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "3m" in md
        assert "1m" in md

    def test_no_project_tasks_sheet_raises(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Other Sheet"
        ws.append(["Name"])
        ws.append(["Task 1"])
        buf = io.BytesIO()
        wb.save(buf)

        with pytest.raises(ValueError, match="No Planner tasks worksheet"):
            convert_planner_to_markdown(buf.getvalue(), "test.xlsx")

    def test_no_tasks_raises(self):
        data = _make_planner_xlsx_bytes(task_rows=[])
        with pytest.raises(ValueError, match="No tasks found"):
            convert_planner_to_markdown(data, "test.xlsx")

    def test_ss_dependency_produces_warning(self):
        task_rows = [
            [1, "Phase 1", "5 days", datetime(2025, 1, 6), datetime(2025, 1, 10), "", "", 0.0, 1],
            [2, "Task A", "3 days", datetime(2025, 1, 6), datetime(2025, 1, 8), "1SS", "", 0.0, 2],
        ]
        data = _make_planner_xlsx_bytes(task_rows=task_rows)
        result = convert_planner_to_markdown(data, "test.xlsx")
        assert any("not supported" in w for w in result["warnings"])

    def test_empty_dependencies_no_depends_tag(self):
        task_rows = [
            [1, "Phase 1", "5 days", datetime(2025, 1, 6), datetime(2025, 1, 10), "", "", 0.0, 1],
            [2, "Task A", "3 days", datetime(2025, 1, 6), datetime(2025, 1, 8), "", "", 0.0, 2],
        ]
        data = _make_planner_xlsx_bytes(task_rows=task_rows)
        result = convert_planner_to_markdown(data, "test.xlsx")
        assert "[depends" not in result["markdown"]

    def test_flat_tasks_without_outline_level(self):
        """Tasks without outline level should be wrapped in a default phase."""
        task_columns = ["ID", "Task Name", "Duration", "Start", "Finish", "Predecessors"]
        task_rows = [
            [1, "Task A", "3 days", datetime(2025, 1, 6), datetime(2025, 1, 8), ""],
            [2, "Task B", "5 days", datetime(2025, 1, 9), datetime(2025, 1, 15), "1FS"],
        ]
        data = _make_planner_xlsx_bytes(task_columns=task_columns, task_rows=task_rows)
        result = convert_planner_to_markdown(data, "test.xlsx")
        assert result["task_count"] == 2
        assert result["phase_count"] == 1

    def test_resource_section_in_front_matter(self):
        data = _make_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "Resources:" in md

    def test_multiple_resources_on_task(self):
        task_rows = [
            [1, "Phase", "5 days", datetime(2025, 1, 6), datetime(2025, 1, 10), "", "", 0.0, 1],
            [2, "Task A", "3 days", datetime(2025, 1, 6), datetime(2025, 1, 8), "", "Alice, Bob", 0.0, 2],
        ]
        data = _make_planner_xlsx_bytes(task_rows=task_rows)
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"].lower()
        assert "@alice" in md
        assert "@bob" in md


# ---------- New-format Planner export (2025+) ----------


def _make_new_planner_xlsx_bytes(
    header_fields=None,
    task_columns=None,
    task_rows=None,
    sheet_name="Tasks",
):
    """Helper: create a new-format Planner export with 8 metadata header rows.

    The newer Planner export uses different field names and has more metadata
    rows than the original format, and the sheet may not be called
    'Project tasks'.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    if header_fields is None:
        header_fields = [
            ("Project Name", "My New Project"),
            ("Plan Owner", "Jane Doe"),
            ("Project Start", datetime(2025, 3, 1)),
            ("Project Finish", datetime(2025, 9, 30)),
            ("Duration", "150 days"),
            ("% Complete", 0.25),
            ("Exported on", datetime(2025, 3, 15)),
            ("Status", "In Progress"),
        ]

    for label, value in header_fields:
        ws.append([label, value])

    # Blank row separator
    ws.append([])

    if task_columns is None:
        task_columns = [
            "Task Number", "Outline Number", "Name", "Assigned to",
            "Bucket", "Labels", "Start", "Finish", "Duration",
            "% Complete", "Priority", "Comment", "Depends on",
            "Dependents (after)", "Effort", "Effort completed",
            "Effort remaining", "Milestone", "Notes", "Completed",
            "Checklist Items", "Sprint", "Goal",
        ]

    ws.append(task_columns)

    if task_rows is None:
        task_rows = [
            [1, "1", "Planning", "Alice", "Phase 1", "", datetime(2025, 3, 1),
             datetime(2025, 3, 14), "10 days", 1.0, "Medium", "", "", "", "",
             "", "", False, "", False, "", "", ""],
            [2, "1.1", "Requirements", "Alice", "Phase 1", "", datetime(2025, 3, 1),
             datetime(2025, 3, 7), "5 days", 1.0, "High", "", "", "", "",
             "", "", False, "", False, "", "", ""],
            [3, "1.2", "Design", "Bob", "Phase 1", "", datetime(2025, 3, 10),
             datetime(2025, 3, 14), "5 days", 0.5, "High", "", "2", "", "",
             "", "", False, "", False, "", "", ""],
            [4, "2", "Implementation", "", "", "", datetime(2025, 3, 17),
             datetime(2025, 4, 4), "15 days", 0.0, "Medium", "", "", "", "",
             "", "", False, "", False, "", "", ""],
            [5, "2.1", "Build Feature", "Alice, Bob", "Phase 2", "", datetime(2025, 3, 17),
             datetime(2025, 3, 28), "10 days", 0.0, "High", "", "3", "", "",
             "", "", False, "", False, "", "", ""],
            [6, "2.2", "Testing", "Charlie", "Phase 2", "", datetime(2025, 3, 31),
             datetime(2025, 4, 4), "5 days", 0.0, "Medium", "", "5", "", "",
             "", "", False, "", False, "", "", ""],
        ]

    for row in task_rows:
        ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestIsTaskHeaderRow:
    def test_metadata_row_not_task_header(self):
        assert _is_task_header_row(["Project Name", "My Project"]) is False

    def test_metadata_row_with_none(self):
        assert _is_task_header_row(["Plan Owner", "John", None, None]) is False

    def test_task_header_row_detected(self):
        row = ["Task Number", "Outline Number", "Name", "Assigned to",
               "Bucket", "Labels", "Start", "Finish", "Duration"]
        assert _is_task_header_row(row) is True

    def test_empty_row(self):
        assert _is_task_header_row([None, None]) is False


class TestNewPlannerDetection:
    def test_detects_by_content_not_sheet_name(self):
        """A sheet named 'Tasks' with Planner metadata should be detected."""
        data = _make_new_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            result = detect_planner_worksheet(wb)
            assert result == "Tasks"
        finally:
            wb.close()

    def test_analyze_workbook_detects_new_format(self):
        data = _make_new_planner_xlsx_bytes()
        result = analyze_workbook(data, "test.xlsx")
        assert result.get("is_planner") is True

    def test_analyze_workbook_shows_task_columns(self):
        """analyze_workbook should skip 8 metadata rows and show task columns."""
        data = _make_new_planner_xlsx_bytes()
        result = analyze_workbook(data, "test.xlsx")
        sheet = result["sheets"][0]
        columns_lower = [c.lower() for c in sheet["columns"]]
        assert "name" in columns_lower
        assert "duration" in columns_lower
        assert "assigned to" in columns_lower
        assert "depends on" in columns_lower
        # Should NOT have metadata labels as columns
        assert "project name" not in columns_lower
        assert "plan owner" not in columns_lower

    def test_analyze_workbook_sample_rows_are_tasks(self):
        data = _make_new_planner_xlsx_bytes()
        result = analyze_workbook(data, "test.xlsx")
        sheet = result["sheets"][0]
        assert len(sheet["sample_rows"]) > 0
        all_cells = [cell for row in sheet["sample_rows"] for cell in row]
        assert any("Planning" in str(c) or "Requirements" in str(c) for c in all_cells)


class TestNewPlannerHeaderParsing:
    def test_parses_project_name(self):
        data = _make_new_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Tasks"]
            header, task_row = parse_planner_header(ws)
            assert header["project_name"] == "My New Project"
        finally:
            wb.close()

    def test_parses_plan_owner(self):
        data = _make_new_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Tasks"]
            header, task_row = parse_planner_header(ws)
            assert header["plan_owner"] == "Jane Doe"
        finally:
            wb.close()

    def test_parses_start_date(self):
        data = _make_new_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Tasks"]
            header, task_row = parse_planner_header(ws)
            assert header["start_date"] == "2025-03-01"
        finally:
            wb.close()

    def test_finds_task_header_row(self):
        data = _make_new_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Tasks"]
            header, task_row = parse_planner_header(ws)
            assert task_row is not None
            rows = list(ws.iter_rows(values_only=True))
            header_cells = [str(c).strip().lower() for c in rows[task_row] if c is not None]
            assert "name" in header_cells
            assert "assigned to" in header_cells
        finally:
            wb.close()

    def test_task_header_row_is_row_9(self):
        """With 8 metadata rows + 1 blank row, task headers should be at index 9."""
        data = _make_new_planner_xlsx_bytes()
        wb = _read_workbook(data, "test.xlsx")
        try:
            ws = wb["Tasks"]
            _header, task_row = parse_planner_header(ws)
            # 8 metadata rows + 1 blank = index 9
            assert task_row == 9
        finally:
            wb.close()


class TestNewPlannerConversion:
    def test_basic_conversion(self):
        data = _make_new_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        assert result["task_count"] > 0
        assert result["phase_count"] > 0
        assert "Planning" in result["markdown"]
        assert "Requirements" in result["markdown"]

    def test_front_matter(self):
        data = _make_new_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "title: My New Project" in md
        assert "manager: Jane Doe" in md
        assert "start: 2025-03-01" in md

    def test_resources_mapped(self):
        data = _make_new_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"].lower()
        assert "@alice" in md
        assert "@bob" in md
        assert "@charlie" in md

    def test_dependencies_converted(self):
        data = _make_new_planner_xlsx_bytes()
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"]
        assert "[depends Requirements]" in md
        assert "[depends Design]" in md
        assert "[depends Build Feature]" in md


class TestGenericConverterWithPlannerData:
    """Test that convert_excel_to_markdown (the generic path used by the
    import wizard) correctly skips Planner metadata header rows."""

    def test_convert_excel_skips_planner_headers_old_format(self):
        data = _make_planner_xlsx_bytes()
        mapping = {"task_name": "Task Name", "duration": "Duration",
                   "start_date": "Start", "end_date": "Finish"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Project tasks", mapping)
        assert result["task_count"] > 0
        assert "Planning" in result["markdown"]

    def test_convert_excel_skips_planner_headers_new_format(self):
        data = _make_new_planner_xlsx_bytes()
        mapping = {"task_name": "Name", "duration": "Duration",
                   "start_date": "Start", "end_date": "Finish",
                   "resources": "Assigned to"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", mapping)
        assert result["task_count"] > 0
        assert "Planning" in result["markdown"]

    def test_convert_excel_finds_name_column(self):
        """The error 'Task name column Name not found' should not occur."""
        data = _make_new_planner_xlsx_bytes()
        mapping = {"task_name": "Name"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", mapping)
        assert result["task_count"] > 0

    def test_convert_excel_parses_text_durations(self):
        """Text durations like '7 days' should be parsed, not produce warnings."""
        data = _make_new_planner_xlsx_bytes()
        mapping = {"task_name": "Name", "duration": "Duration"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", mapping)
        md = result["markdown"]
        assert "10d" in md
        assert "5d" in md
        duration_warnings = [w for w in result["warnings"] if "duration" in w.lower()]
        assert len(duration_warnings) == 0

    def test_convert_excel_parses_week_durations(self):
        """Week durations like '2 weeks' should produce '2w'."""
        task_rows = [
            [1, "1", "Phase", "", "", "", datetime(2025, 3, 1),
             datetime(2025, 3, 14), "2 weeks", 0.0, "", "", "", "", "",
             "", "", False, "", False, "", "", ""],
        ]
        data = _make_new_planner_xlsx_bytes(task_rows=task_rows)
        mapping = {"task_name": "Name", "duration": "Duration"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", mapping)
        assert "2w" in result["markdown"]

    def test_convert_excel_parses_dependencies_new_format(self):
        """Dependencies in 'Depends on' column should produce [depends ...] tags."""
        data = _make_new_planner_xlsx_bytes()
        mapping = {"task_name": "Name", "duration": "Duration",
                   "depends_on": "Depends on"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", mapping)
        md = result["markdown"]
        assert "[depends Requirements]" in md
        assert "[depends Design]" in md
        assert "[depends Build Feature]" in md

    def test_convert_excel_parses_dependencies_old_format(self):
        """Dependencies in 'Predecessors' column with FS suffix should work."""
        data = _make_planner_xlsx_bytes()
        mapping = {"task_name": "Task Name", "duration": "Duration",
                   "depends_on": "Predecessors"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Project tasks", mapping)
        md = result["markdown"]
        assert "[depends Requirements]" in md
        assert "[depends Design]" in md

    def test_convert_excel_no_depends_without_mapping(self):
        """Without depends_on in mapping, no [depends] tags should appear."""
        data = _make_new_planner_xlsx_bytes()
        mapping = {"task_name": "Name", "duration": "Duration"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Tasks", mapping)
        assert "[depends" not in result["markdown"]


# ---------- TestSplitResourceNames ----------


class TestSplitResourceNames:
    """Tests for _split_resource_names which splits on comma, semicolon, and slash."""

    def test_comma_separated(self):
        assert _split_resource_names("Alice, Bob") == ["Alice", "Bob"]

    def test_semicolon_separated(self):
        assert _split_resource_names("Alice; Bob") == ["Alice", "Bob"]

    def test_slash_separated(self):
        assert _split_resource_names("Jack/Sandeep") == ["Jack", "Sandeep"]

    def test_mixed_delimiters(self):
        assert _split_resource_names("Alice, Bob/Charlie; Dan") == [
            "Alice", "Bob", "Charlie", "Dan",
        ]

    def test_trims_whitespace(self):
        assert _split_resource_names("  Alice , Bob  / Charlie ") == [
            "Alice", "Bob", "Charlie",
        ]

    def test_empty_string(self):
        assert _split_resource_names("") == []

    def test_none(self):
        assert _split_resource_names(None) == []

    def test_single_name(self):
        assert _split_resource_names("Alice") == ["Alice"]

    def test_skips_empty_segments(self):
        assert _split_resource_names("Alice,,Bob") == ["Alice", "Bob"]


# ---------- TestSlashSeparatedResourcesEndToEnd ----------


class TestSlashSeparatedResources:
    """End-to-end tests for slash-separated resources in Planner imports (issue #650)."""

    def test_planner_slash_resources_old_format(self):
        """Slash-separated resources like 'Jack/Sandeep' should become @Jack and @Sandeep."""
        task_rows = [
            [1, "Phase", "5 days", datetime(2025, 1, 6), datetime(2025, 1, 10), "", "", 0.0, 1],
            [2, "Task A", "3 days", datetime(2025, 1, 6), datetime(2025, 1, 8), "", "Jack/sandeep", 0.0, 2],
        ]
        data = _make_planner_xlsx_bytes(task_rows=task_rows)
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"].lower()
        assert "@jack" in md
        assert "@sandeep" in md
        # Should NOT have a combined resource
        assert "jack/sandeep" not in md

    def test_planner_slash_resources_new_format(self):
        """Slash-separated resources in new Planner format."""
        task_rows = [
            [1, "1", "Phase", "Jack/Sandeep", "Phase 1", "", datetime(2025, 3, 1),
             datetime(2025, 3, 14), "10 days", 0.0, "Medium", "", "", "", "",
             "", "", False, "", False, "", "", ""],
            [2, "1.1", "Task A", "Alice/Bob", "Phase 1", "", datetime(2025, 3, 1),
             datetime(2025, 3, 7), "5 days", 0.0, "High", "", "", "", "",
             "", "", False, "", False, "", "", ""],
        ]
        data = _make_new_planner_xlsx_bytes(task_rows=task_rows)
        result = convert_planner_to_markdown(data, "test.xlsx")
        md = result["markdown"].lower()
        assert "@jack" in md
        assert "@sandeep" in md
        assert "@alice" in md
        assert "@bob" in md

    def test_generic_converter_slash_resources(self):
        """Slash-separated resources via the generic convert_excel_to_markdown path."""
        task_rows = [
            [1, "Phase", "5 days", datetime(2025, 1, 6), datetime(2025, 1, 10), "", "Jack/Sandeep", 0.0, 1],
            [2, "Task A", "3 days", datetime(2025, 1, 6), datetime(2025, 1, 8), "", "Alice", 0.0, 2],
        ]
        data = _make_planner_xlsx_bytes(task_rows=task_rows)
        mapping = {"task_name": "Task Name", "duration": "Duration",
                   "resources": "Resource Names"}
        result = convert_excel_to_markdown(data, "test.xlsx", "Project tasks", mapping)
        md = result["markdown"].lower()
        assert "@jack" in md
        assert "@sandeep" in md
