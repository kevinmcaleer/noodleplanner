"""Test scenarios for plan scheduling, duration calculation, and exports.

20 test scenarios covering scheduling, timesheets, RAG status, completion,
and export formats. Each test defines a markdown plan, runs it through the
planning engine, and verifies the calculated outputs.

GitHub issue: #107
"""

import csv
import io
import os
import tempfile
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from noodle_core import (
    add_working_days,
    calculate_rag_status,
    get_next_working_day,
    schedule_tasks,
    text_to_markdown_table,
    natural_language_to_yaml,
    export_to_excel,
    export_to_csv,
)


# ---------------------------------------------------------------------------
# Helper: parse plan text and return scheduled tasks
# ---------------------------------------------------------------------------

def parse_and_schedule(plan_text, project_name="Test Project"):
    """Parse a natural-language plan and return the scheduled task list."""
    data = natural_language_to_yaml(plan_text, project_name)
    phases_raw = data[project_name]
    if isinstance(phases_raw, list):
        phases = phases_raw
    elif isinstance(phases_raw, dict):
        phases = [phases_raw]
    else:
        phases = []
    return schedule_tasks(phases)


def find_task(tasks, name):
    """Find a task by name (case-insensitive substring match)."""
    for t in tasks:
        if name.lower() in t.get("name", "").lower():
            return t
    raise ValueError(f"Task '{name}' not found in {[t.get('name') for t in tasks]}")


# ===========================================================================
# Scheduling & Duration scenarios (1-10)
# ===========================================================================


class TestSimpleSequentialTasks:
    """Scenario 1: Three tasks with * prefix start one after the other."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  * Task B @Bob 2d
  * Task C @Alice 4d"""

    def test_task_b_starts_after_task_a(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        b = find_task(tasks, "Task B")
        # Task B should start on or after Task A's finish
        assert b["start"] >= a["finish"]

    def test_task_c_starts_after_task_b(self):
        tasks = parse_and_schedule(self.PLAN)
        b = find_task(tasks, "Task B")
        c = find_task(tasks, "Task C")
        assert c["start"] >= b["finish"]

    def test_correct_number_of_tasks(self):
        tasks = parse_and_schedule(self.PLAN)
        leaf_tasks = [t for t in tasks if not t.get("summary")]
        assert len(leaf_tasks) == 3


class TestParallelTasks:
    """Scenario 2: Three tasks without dependencies start on the same date."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  Task B @Bob 2d
  Task C @Carol 4d"""

    def test_all_start_same_date(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        b = find_task(tasks, "Task B")
        c = find_task(tasks, "Task C")
        assert a["start"] == b["start"]
        assert b["start"] == c["start"]


class TestDependencies:
    """Scenario 3: Task C depends on Task A."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  Task B @Bob 2d
  Task C @Carol 2d [depends Task A]"""

    def test_task_c_starts_after_task_a(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        c = find_task(tasks, "Task C")
        assert c["start"] >= a["finish"]

    def test_task_b_independent_of_a(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        b = find_task(tasks, "Task B")
        # Task B is parallel to Task A (same start)
        assert a["start"] == b["start"]


class TestMultiDependency:
    """Scenario 4: Task D depends on both Task A and Task B."""

    PLAN = """\
Phase 1
  Task A @Alice 5d
  Task B @Bob 3d
  Task D @Carol 2d [depends Task A, Task B]"""

    def test_task_d_starts_after_latest_dependency(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        b = find_task(tasks, "Task B")
        d = find_task(tasks, "Task D")
        latest_finish = max(a["finish"], b["finish"])
        assert d["start"] >= latest_finish


class TestLagTime:
    """Scenario 5: Dependency with +2d lag time."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  Task B @Bob 2d [depends Task A +2d]"""

    def test_lag_creates_gap(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        b = find_task(tasks, "Task B")
        # Task B should start at least 2 working days after Task A finishes
        # add_working_days(a['finish'], 2) gives the finish with 2 day lag
        expected_dep_with_lag = add_working_days(a["finish"], 2)
        assert b["start"] >= expected_dep_with_lag


class TestLeadTime:
    """Scenario 6: Dependency with -1d lead time (overlap)."""

    PLAN = """\
Phase 1
  Task A @Alice 5d
  Task B @Bob 3d [depends Task A -1d]"""

    def test_lead_creates_overlap(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        b = find_task(tasks, "Task B")
        # With -1d lead, Task B should start 1 working day before Task A finishes
        expected_dep_with_lead = add_working_days(a["finish"], -1)
        assert b["start"] >= expected_dep_with_lead
        # Task B should start before where it would without lead
        assert b["start"] < a["finish"] or b["start"] == expected_dep_with_lead


class TestZeroDurationMilestones:
    """Scenario 7: Milestone with 0d duration depends on a task."""

    PLAN = """\
Phase 1
  Task A @Alice 5d
  * Milestone 1 0d"""

    def test_milestone_aligns_with_dependency_finish(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        m = find_task(tasks, "Milestone 1")
        assert m["start"] == a["finish"]
        assert m["finish"] == a["finish"]

    def test_milestone_has_zero_duration(self):
        tasks = parse_and_schedule(self.PLAN)
        m = find_task(tasks, "Milestone 1")
        assert m["duration"] == timedelta(days=0)


class TestWeekendSkipping:
    """Scenario 8: A 5-day task starting on a Friday should skip the weekend."""

    def test_five_day_task_from_friday(self):
        # Friday 2026-02-13
        friday = datetime(2026, 2, 13)
        finish = add_working_days(friday, 5)
        # 5 working days from Friday: Fri, Mon, Tue, Wed, Thu
        # finish is exclusive (day after last working day) = Fri 2026-02-20
        expected = datetime(2026, 2, 20)
        # The 5 working days are: Feb 13 (Fri), Feb 16 (Mon), Feb 17 (Tue),
        # Feb 18 (Wed), Feb 19 (Thu). Next day is Feb 20.
        assert finish == expected

    def test_one_day_task_from_friday(self):
        friday = datetime(2026, 2, 13)
        finish = add_working_days(friday, 1)
        # 1 working day = Friday itself, finish exclusive = Saturday Feb 14
        expected = datetime(2026, 2, 14)
        assert finish == expected

    def test_weekend_start_moves_to_monday(self):
        saturday = datetime(2026, 2, 14)
        next_wd = get_next_working_day(saturday)
        assert next_wd == datetime(2026, 2, 16)  # Monday


class TestHolidaySkipping:
    """Scenario 9: Task spanning a defined holiday skips the holiday.

    Note: schedule_tasks does not currently pass holidays through,
    so we test add_working_days directly with holiday support.
    """

    def test_holiday_is_skipped(self):
        monday = datetime(2026, 2, 16)
        holiday = datetime(2026, 2, 18)  # Wednesday
        holidays = {holiday}
        finish = add_working_days(monday, 5, holidays)
        # Without holiday: Mon-Fri (5 days), finish = Sat Feb 21
        # With holiday Wed skipped: Mon, Tue, Thu, Fri, Mon -> finish = Tue Feb 24
        expected = datetime(2026, 2, 24)
        assert finish == expected

    def test_holiday_on_start_date(self):
        holiday = datetime(2026, 2, 16)  # Monday
        holidays = {holiday}
        next_wd = get_next_working_day(holiday, holidays)
        assert next_wd == datetime(2026, 2, 17)  # Tuesday


class TestPhaseSummaryDates:
    """Scenario 10: Phase start = earliest child start, phase finish = latest child finish."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  Task B @Bob 5d"""

    def test_phase_encompasses_children(self):
        tasks = parse_and_schedule(self.PLAN)
        phase = find_task(tasks, "Phase 1")
        a = find_task(tasks, "Task A")
        b = find_task(tasks, "Task B")
        assert phase.get("summary") is True
        assert phase["start"] == min(a["start"], b["start"])
        assert phase["finish"] == max(a["finish"], b["finish"])


# ===========================================================================
# Timesheet & Resource scenarios (11-13)
# ===========================================================================


class TestSingleResourceAllocation:
    """Scenario 11: One resource across 3 sequential tasks, verify total days."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  * Task B @Alice 2d
  * Task C @Alice 5d"""

    def test_resource_total_days(self):
        tasks = parse_and_schedule(self.PLAN)
        alice_tasks = [
            t for t in tasks
            if "alice" in t.get("resources", "").lower() and not t.get("summary")
        ]
        total_days = sum(t["duration"].days for t in alice_tasks)
        assert total_days == 10  # 3 + 2 + 5


class TestMultipleResources:
    """Scenario 12: Multiple resources on different tasks."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  Task B @Bob 5d
  Task C @Carol 2d"""

    def test_each_resource_total(self):
        tasks = parse_and_schedule(self.PLAN)
        leaf_tasks = [t for t in tasks if not t.get("summary")]
        resource_days = {}
        for t in leaf_tasks:
            res = t.get("resources", "")
            if res:
                for r in res.split(","):
                    r = r.strip().lower()
                    resource_days[r] = resource_days.get(r, 0) + t["duration"].days
        assert resource_days.get("alice", 0) == 3
        assert resource_days.get("bob", 0) == 5
        assert resource_days.get("carol", 0) == 2


class TestResourceOnOverlappingTasks:
    """Scenario 13: Same resource on parallel tasks."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  Task B @Alice 5d"""

    def test_resource_allocated_to_both(self):
        tasks = parse_and_schedule(self.PLAN)
        a = find_task(tasks, "Task A")
        b = find_task(tasks, "Task B")
        # Both tasks should be assigned to Alice
        assert "alice" in a.get("resources", "").lower()
        assert "alice" in b.get("resources", "").lower()
        # Both start at the same time (parallel)
        assert a["start"] == b["start"]
        # Total allocated days for Alice is 8 (3 + 5), even though they overlap
        alice_tasks = [
            t for t in tasks
            if "alice" in t.get("resources", "").lower() and not t.get("summary")
        ]
        total_days = sum(t["duration"].days for t in alice_tasks)
        assert total_days == 8


# ===========================================================================
# RAG Status scenarios (14-16)
# ===========================================================================


class TestGreenTask:
    """Scenario 14: 100% complete task should be Green."""

    def test_completed_task_is_green(self):
        task = {
            "start": datetime(2026, 1, 5),
            "finish": datetime(2026, 1, 16),
            "percent": 100,
            "duration": timedelta(days=10),
        }
        result = calculate_rag_status(task, datetime(2026, 2, 1))
        assert result == "Green"

    def test_future_task_is_green(self):
        task = {
            "start": datetime(2026, 6, 1),
            "finish": datetime(2026, 6, 12),
            "percent": 0,
            "duration": timedelta(days=10),
        }
        # Current date before start
        result = calculate_rag_status(task, datetime(2026, 1, 1))
        assert result == "Green"


class TestRedTask:
    """Scenario 15: 0% complete task past its start date should be Red."""

    def test_overdue_no_progress_is_red(self):
        task = {
            "start": datetime(2026, 1, 5),
            "finish": datetime(2026, 1, 16),
            "percent": 0,
            "duration": timedelta(days=10),
        }
        # Current date is after start with 0% progress
        result = calculate_rag_status(task, datetime(2026, 1, 12))
        assert result == "Red"


class TestAmberTask:
    """Scenario 16: Partially complete but behind schedule should be Amber."""

    def test_behind_schedule_is_amber(self):
        task = {
            "start": datetime(2026, 1, 5),
            "finish": datetime(2026, 1, 16),
            "percent": 50,
            "duration": timedelta(days=10),
        }
        # 8 out of 11 days elapsed (~73% expected) but only 50% done
        result = calculate_rag_status(task, datetime(2026, 1, 13))
        assert result == "Amber"


# ===========================================================================
# Completion scenarios (17-18)
# ===========================================================================


class TestPhaseCompletionFromChildren:
    """Scenario 17: Phase percent is average of child task percents."""

    PLAN = """\
Phase 1
  Task A @Alice 3d 50%
  Task B @Bob 3d 100%"""

    def test_phase_percent_is_average(self):
        tasks = parse_and_schedule(self.PLAN)
        phase = find_task(tasks, "Phase 1")
        assert phase.get("summary") is True
        # Average of 50 and 100 = 75
        assert phase["percent"] == 75


class TestComplexPlanEndToEnd:
    """Scenario 18: Full project with multiple phases, dependencies, milestones.

    Uses [depends ...] for cross-phase references and * for sequential tasks
    within the same phase. The parser strips '*' but leaves a leading space
    in the task name, so cross-phase dependencies reference non-sequential
    task names to ensure reliable name matching.
    """

    PLAN = """\
Planning
  Requirements @Alice 3d 100%
  Design @Bob 5d 50% [depends Requirements]
Development
  Backend @Alice 10d [depends Design]
  Frontend @Carol 8d [depends Design]
  Integration @Bob 3d [depends Backend, Frontend]
  * Release Milestone 0d"""

    def test_correct_task_count(self):
        tasks = parse_and_schedule(self.PLAN)
        leaf_tasks = [t for t in tasks if not t.get("summary")]
        assert len(leaf_tasks) == 6  # Requirements, Design, Backend, Frontend, Integration, Release Milestone

    def test_phases_are_summary_tasks(self):
        tasks = parse_and_schedule(self.PLAN)
        planning = find_task(tasks, "Planning")
        development = find_task(tasks, "Development")
        assert planning.get("summary") is True
        assert development.get("summary") is True

    def test_dependency_chain(self):
        tasks = parse_and_schedule(self.PLAN)
        design = find_task(tasks, "Design")
        backend = find_task(tasks, "Backend")
        frontend = find_task(tasks, "Frontend")
        integration = find_task(tasks, "Integration")
        # Backend and Frontend start after Design
        assert backend["start"] >= design["finish"]
        assert frontend["start"] >= design["finish"]
        # Integration starts after both Backend and Frontend
        assert integration["start"] >= backend["finish"]
        assert integration["start"] >= frontend["finish"]

    def test_milestone_at_end(self):
        tasks = parse_and_schedule(self.PLAN)
        integration = find_task(tasks, "Integration")
        milestone = find_task(tasks, "Release Milestone")
        assert milestone["start"] == integration["finish"]
        assert milestone["duration"] == timedelta(days=0)

    def test_sequential_design_after_requirements(self):
        tasks = parse_and_schedule(self.PLAN)
        req = find_task(tasks, "Requirements")
        design = find_task(tasks, "Design")
        # Design depends on Requirements
        assert design["start"] >= req["finish"]


# ===========================================================================
# Export scenarios (19-20)
# ===========================================================================


class TestExcelExport:
    """Scenario 19: Excel export produces a valid file."""

    PLAN = """\
Phase 1
  Task A @Alice 3d 50%
  Task B @Bob 2d
Phase 2
  Task C @Carol 4d [depends Task A]"""

    def test_excel_export_produces_valid_file(self):
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                tmp_path = tmp.name
            data = natural_language_to_yaml(self.PLAN, "Test Project")
            # export_to_excel expects the text and path
            export_to_excel(self.PLAN, tmp_path, is_yaml=False, project_name="Test Project")
            # Verify file is non-empty
            assert os.path.getsize(tmp_path) > 0
            # Verify it can be read by openpyxl
            wb = load_workbook(tmp_path)
            ws = wb.active
            assert ws.title == "Tasks"
            # Verify headers
            headers = [cell.value for cell in ws[1]]
            assert "ID" in headers
            assert "Task Name" in headers
            assert "Start" in headers
            assert "Finish" in headers
            assert "Duration (days)" in headers
            assert "Resources" in headers
            # Verify data rows exist (header + at least 1 data row)
            assert ws.max_row >= 2
            wb.close()
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def test_excel_export_has_correct_row_count(self):
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                tmp_path = tmp.name
            export_to_excel(self.PLAN, tmp_path, is_yaml=False, project_name="Test Project")
            wb = load_workbook(tmp_path)
            ws = wb.active
            # Should have header row + tasks (2 phases + 3 tasks = 5 rows)
            tasks = parse_and_schedule(self.PLAN)
            expected_rows = 1 + len(tasks)  # header + all tasks including phases
            assert ws.max_row == expected_rows
            wb.close()
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.unlink(tmp_path)


class TestCSVExport:
    """Scenario 20: CSV export produces valid output."""

    PLAN = """\
Phase 1
  Task A @Alice 3d
  Task B @Bob 2d"""

    def test_csv_export_produces_valid_output(self):
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as tmp:
                tmp_path = tmp.name
            export_to_csv(self.PLAN, tmp_path, is_yaml=False, project_name="Test Project")
            # Verify file is non-empty
            assert os.path.getsize(tmp_path) > 0
            # Read and parse CSV
            with open(tmp_path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                rows = list(reader)
            # Verify headers
            assert "ID" in reader.fieldnames
            assert "Task Name" in reader.fieldnames
            assert "Start" in reader.fieldnames
            assert "Finish" in reader.fieldnames
            assert "Duration (days)" in reader.fieldnames
            # Verify we have data rows
            assert len(rows) >= 2  # At least phase + tasks
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def test_csv_export_has_expected_columns(self):
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as tmp:
                tmp_path = tmp.name
            export_to_csv(self.PLAN, tmp_path, is_yaml=False, project_name="Test Project")
            with open(tmp_path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                expected_columns = {
                    "ID", "Task Name", "Start", "Finish",
                    "Duration (days)", "Resources", "% Complete", "RAG",
                    "Priority", "Bucket", "Comment"
                }
                assert set(reader.fieldnames) == expected_columns
        finally:
            if tmp_path and os.path.exists(tmp_path):
                os.unlink(tmp_path)


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
