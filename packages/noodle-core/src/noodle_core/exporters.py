"""Export functions: Excel, PowerPoint, CSV, PDF, and related helpers.

Depends on: date_math, metadata, renderers, and the scheduling functions
in scheduling_engine.
"""

import os
import re
import csv
import yaml
import logging
from datetime import datetime, timedelta
from dateutil.parser import parse as parse_date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from .format_converter import (
    extract_raid_log, parse_raid_markdown, extract_baseline,
    parse_baseline_markdown, extract_budget, parse_budget_markdown,
)
from .date_math import get_next_working_day, add_working_days
from .renderers import render_gantt_chart, render_resource_sheet, render_custom_timeline

logger = logging.getLogger(__name__)

# RAID Log Excel column widths (in characters)
RAID_COLUMN_WIDTHS = {
    'ID': 6,
    'Type': 14,
    'Title': 25,
    'Description': 35,
    'Raised By': 15,
    'Owner': 15,
    'Mitigation Actions': 35,
    'Impact': 10,
    'Likelihood': 12,
    'Score': 8,
    'Status': 14,
    'Priority': 12,
    'Target Date': 14,
}


def calculate_rag_status(task, current_date=None):
    """Calculate RAG status for a task with descriptive text.

    Args:
        task: Task dictionary with start, finish, percent, duration
        current_date: Current date for comparison (defaults to today)

    Returns:
        String: 'Complete', 'Not Started', 'Ahead of Schedule', 'On Track',
                'Behind Schedule', or 'Task Overdue'
    """
    if current_date is None:
        current_date = datetime.now().date()
    else:
        current_date = current_date.date() if hasattr(current_date, 'date') else current_date

    # Get task details
    start_date = task.get('start')
    finish_date = task.get('finish')
    percent_complete = task.get('percent')

    # Blue: Task is 100% complete (check first, regardless of dates)
    if percent_complete == 100:
        return 'Complete'

    # Red: deadline slippage (#877). A deadline is a fixed marker, distinct
    # from the on-track/behind-schedule comparison below and from the
    # start/finish dates that drive the schedule -- it never moves them. It
    # flags the task once the deadline date has passed while the task is
    # still incomplete (checked above), or once the computed finish
    # (start + duration) is later than the deadline.
    deadline = task.get('deadline')
    if deadline:
        deadline_date = deadline.date() if hasattr(deadline, 'date') else parse_date(deadline).date()
        deadline_finish = finish_date.date() if hasattr(finish_date, 'date') else finish_date
        if deadline_date < current_date or (deadline_finish is not None and deadline_finish > deadline_date):
            return 'Task Overdue'

    if not start_date or not finish_date:
        # Fallback when dates are missing: use percentage thresholds
        if percent_complete is None or percent_complete == 0:
            return 'Task Overdue'
        elif percent_complete < 50:
            return 'Task Overdue'
        elif percent_complete < 80:
            return 'Behind Schedule'
        else:
            return 'On Track'

    # Convert to date objects
    start_date = start_date.date() if hasattr(start_date, 'date') else start_date
    finish_date = finish_date.date() if hasattr(finish_date, 'date') else finish_date

    # Green: Task hasn't started yet (start date is in the future)
    if start_date > current_date:
        if percent_complete is not None and percent_complete > 0:
            return 'Ahead of Schedule'
        return 'Not Started'

    # Red: Start date is in the past and no progress or 0%
    if start_date <= current_date and (percent_complete is None or percent_complete == 0):
        return 'Task Overdue'

    # Calculate expected progress
    total_duration = (finish_date - start_date).days
    if total_duration <= 0:
        total_duration = 1  # Avoid division by zero

    elapsed_days = (current_date - start_date).days
    if elapsed_days < 0:
        elapsed_days = 0

    expected_percent = min(100, (elapsed_days / total_duration) * 100)

    # Amber: Actual progress is less than expected
    if percent_complete < expected_percent:
        return 'Behind Schedule'

    # Green: On track or ahead
    return 'On Track'


def rag_status_to_colour(rag_status):
    """Map a descriptive RAG status to its colour category.

    Args:
        rag_status: Descriptive status string from calculate_rag_status

    Returns:
        String: 'green', 'amber', 'red', 'blue', or 'grey'
    """
    mapping = {
        'not started': 'green',
        'on track': 'green',
        'ahead of schedule': 'green',
        'complete': 'blue',
        'behind schedule': 'amber',
        'task overdue': 'red',
        # Legacy values for backwards compatibility
        'green': 'green',
        'amber': 'amber',
        'red': 'red',
    }
    return mapping.get((rag_status or '').lower(), 'grey')


def _parse_non_working_suffix(suffix_str):
    """Parse a non-working [...] suffix into a set of datetime.date objects.

    Supports multiple formats:
    - Individual dates: ``2026-12-31``
    - Date ranges: ``2026-03-01:2026-04-01``
    - Named entries: ``Annual Leave: 2026-03-01:2026-03-14``
    - Named single dates: ``Doctor: 2026-04-01``

    Named entries have the format ``Name: date`` or ``Name: date:date``.
    The name is for display only; only dates are extracted.

    Args:
        suffix_str: The content inside the brackets,
            e.g. "Annual Leave: 2026-03-01:2026-03-14, Doctor: 2026-04-01"

    Returns:
        A set of datetime.date objects.
    """
    from datetime import datetime as _dt, timedelta as _td

    dates = set()
    if not suffix_str:
        return dates

    for part in suffix_str.split(','):
        part = part.strip()
        if not part:
            continue

        # Check for named entry: "Name: date" or "Name: date:date"
        named_match = re.match(
            r'([^:]+?):\s*(\d{4}-\d{2}-\d{2})(?::(\d{4}-\d{2}-\d{2}))?\s*$',
            part
        )
        if named_match:
            # Named entry -- the first group may be a name or a date
            # If group(1) looks like a date, treat as range (legacy format)
            if re.match(r'\d{4}-\d{2}-\d{2}$', named_match.group(1).strip()):
                # Legacy range format: "2026-03-01:2026-03-14"
                try:
                    start = _dt.strptime(named_match.group(1).strip(), '%Y-%m-%d').date()
                    end = _dt.strptime(named_match.group(2), '%Y-%m-%d').date()
                    current = start
                    while current <= end:
                        dates.add(current)
                        current += _td(days=1)
                except ValueError:
                    logger.warning("Invalid date range in non-working suffix: %s", part)
            else:
                # Named entry: "Name: start_date" or "Name: start_date:end_date"
                try:
                    start = _dt.strptime(named_match.group(2), '%Y-%m-%d').date()
                    if named_match.group(3):
                        end = _dt.strptime(named_match.group(3), '%Y-%m-%d').date()
                        current = start
                        while current <= end:
                            dates.add(current)
                            current += _td(days=1)
                    else:
                        dates.add(start)
                except ValueError:
                    logger.warning("Invalid date in named non-working entry: %s", part)
            continue

        # Legacy range: start:end (both are dates)
        if ':' in part:
            range_parts = part.split(':', 1)
            try:
                start = _dt.strptime(range_parts[0].strip(), '%Y-%m-%d').date()
                end = _dt.strptime(range_parts[1].strip(), '%Y-%m-%d').date()
                current = start
                while current <= end:
                    dates.add(current)
                    current += _td(days=1)
            except ValueError:
                logger.warning("Invalid date range in non-working suffix: %s", part)
        else:
            # Single date
            try:
                dates.add(_dt.strptime(part, '%Y-%m-%d').date())
            except ValueError:
                logger.warning("Invalid date in non-working suffix: %s", part)

    return dates


def _parse_named_non_working_suffix(suffix_str):
    """Parse a non-working [...] suffix preserving names and ranges.

    Returns a list of dicts with 'name', 'start', and optional 'finish' keys.
    """
    entries = []
    if not suffix_str:
        return entries

    for part in suffix_str.split(','):
        part = part.strip()
        if not part:
            continue

        named_match = re.match(
            r'([^:]+?):\s*(\d{4}-\d{2}-\d{2})(?::(\d{4}-\d{2}-\d{2}))?\s*$',
            part
        )
        if named_match:
            if re.match(r'\d{4}-\d{2}-\d{2}$', named_match.group(1).strip()):
                # Legacy range without name
                entries.append({
                    'name': '',
                    'start': named_match.group(1).strip(),
                    'finish': named_match.group(2),
                })
            else:
                entries.append({
                    'name': named_match.group(1).strip(),
                    'start': named_match.group(2),
                    'finish': named_match.group(3) or '',
                })
        else:
            # Single date without name
            date_match = re.match(r'(\d{4}-\d{2}-\d{2})\s*$', part)
            if date_match:
                entries.append({
                    'name': '',
                    'start': date_match.group(1),
                    'finish': '',
                })

    return entries


_CALENDAR_SUFFIX_RE = re.compile(r',?\s*calendar\s+([^,\[]+?)(?=\s*(?:,|non-working\s*\[|$))', re.IGNORECASE)

# The non-working [...] suffix, order-independent relative to `calendar
# <Name>` (issue #1136 lets either come first): stops at the next comma,
# the `calendar` keyword, or end of string, rather than requiring `\s*$`.
_NON_WORKING_SUFFIX_RE = re.compile(
    r',?\s*non-working\s*\[([^\]]*)\]\s*(?=,|\s+calendar\s+|$)', re.IGNORECASE
)


def _extract_non_working_suffix(full_info):
    """Remove a resource line's ``non-working [...]`` suffix, wherever it
    falls relative to a ``calendar <Name>`` suffix.

    Returns ``(full_info_without_the_suffix, raw_bracket_contents_or_none)``.
    """
    match = _NON_WORKING_SUFFIX_RE.search(full_info)
    if not match:
        return full_info, None
    remainder = (full_info[:match.start()] + full_info[match.end():]).strip().rstrip(',').strip()
    return remainder, match.group(1)


def _strip_calendar_suffix(full_info):
    """Remove a resource line's ``calendar <Name>`` suffix (issue #1136).

    Order-independent relative to the ``non-working [...]`` suffix (a
    resource line may write either one first) -- searched anywhere in the
    string rather than only at the end, stopping at the next comma or
    ``non-working [`` so a multi-word calendar name (``calendar Fortnight
    Ops``) doesn't swallow the fields after it.

    Returns ``(full_info_without_the_suffix, calendar_name_or_none)``.
    """
    match = _CALENDAR_SUFFIX_RE.search(full_info)
    if not match:
        return full_info, None
    calendar_name = match.group(1).strip()
    remainder = (full_info[:match.start()] + full_info[match.end():]).strip().rstrip(',').strip()
    return remainder, calendar_name


def parse_resource_calendars(original_text):
    """Parse each resource's assigned calendar name from its ``calendar
    <Name>`` suffix (issue #1136).

        - @kev: Kevin McAleer, PM calendar Gulf non-working [2026-06-10]

    Returns a dict mapping lowercase shortnames to the calendar name as
    written -- resolving that name to a Calendar (see calendar_model.py)
    against the plan's declared ``calendars:`` is the caller's job, since
    this function only sees one resource line at a time.
    """
    resource_calendars = {}
    if not original_text:
        return resource_calendars

    lines = original_text.split('\n')
    in_frontmatter = False
    in_resources = False

    for line in lines:
        if line.strip() == '---':
            if not in_frontmatter:
                in_frontmatter = True
            else:
                break
            continue

        if not in_frontmatter:
            continue

        if line.strip().startswith('Resources:'):
            in_resources = True
            continue

        if in_resources and line and not line.startswith(' ') and not line.startswith('-'):
            in_resources = False

        if in_resources and line.strip().startswith('-'):
            match = re.match(r'\s*-\s*@(\w+):\s*(.+)', line)
            if match:
                short_name = match.group(1)
                full_info = match.group(2).strip()
                full_info, _nwd = _extract_non_working_suffix(full_info)
                _, calendar_name = _strip_calendar_suffix(full_info)
                if calendar_name:
                    resource_calendars[short_name.lower()] = calendar_name

    return resource_calendars


def parse_resource_mappings(original_text):
    """Parse resource mappings from YAML front matter.

    Returns a tuple of (resource_map, resource_non_working_days).
    - resource_map: dict mapping short names (lowercase) to full names.
    - resource_non_working_days: dict mapping short names (lowercase) to sets of datetime.date.

    Example resource_map: {'andy': 'Andy McCarthy', 'bob': 'Bob Smith'}

    Resource lines may have an optional non-working days suffix:
        - @Jack: Jack Lloyd, Network Arch, non-working [2026-03-01:2026-04-01, 2026-12-31]
    """
    resource_map = {}
    resource_nwd = {}

    # Extract YAML front matter
    lines = original_text.split('\n')
    in_frontmatter = False
    in_resources = False

    for line in lines:
        if line.strip() == '---':
            if not in_frontmatter:
                in_frontmatter = True
            else:
                break  # End of front matter
            continue

        if not in_frontmatter:
            continue

        # Check if we're in the Resources section
        if line.strip().startswith('Resources:'):
            in_resources = True
            continue

        # Check if we've left the Resources section
        if in_resources and line and not line.startswith(' ') and not line.startswith('-'):
            in_resources = False

        # Parse resource line
        if in_resources and line.strip().startswith('-'):
            # Format: - @Andy: Andy McCarthy, Lead Developer, non-working [2026-03-01:2026-04-01]
            match = re.match(r'\s*-\s*@(\w+):\s*(.+)', line)
            if match:
                short_name = match.group(1)
                full_info = match.group(2).strip()

                # Extract non-working days suffix if present, wherever it
                # falls relative to a calendar suffix (issue #1136)
                full_info, nwd_raw = _extract_non_working_suffix(full_info)
                if nwd_raw:
                    nwd_dates = _parse_non_working_suffix(nwd_raw)
                    if nwd_dates:
                        resource_nwd[short_name.lower()] = nwd_dates

                # A calendar suffix (issue #1136) is metadata, not part of
                # the displayed name/role -- strip it here too.
                full_info, _calendar_name = _strip_calendar_suffix(full_info)

                # Extract just the name (before the first comma)
                name_only = full_info.split(',')[0].strip()
                # Store with lowercase key for case-insensitive lookup
                resource_map[short_name.lower()] = name_only

    return resource_map, resource_nwd


def parse_resource_roles(original_text):
    """Parse resource roles from YAML front matter.

    Returns a dict mapping short names (lowercase) to their role string.
    The role is the text after the first comma in the resource line.

    Example: ``- @Jack: Jack Lloyd, Network Arch`` returns ``{'jack': 'Network Arch'}``.
    """
    role_map = {}
    if not original_text:
        return role_map

    lines = original_text.split('\n')
    in_frontmatter = False
    in_resources = False

    for line in lines:
        if line.strip() == '---':
            if not in_frontmatter:
                in_frontmatter = True
            else:
                break
            continue

        if not in_frontmatter:
            continue

        if line.strip().startswith('Resources:'):
            in_resources = True
            continue

        if in_resources and line and not line.startswith(' ') and not line.startswith('-'):
            in_resources = False

        if in_resources and line.strip().startswith('-'):
            match = re.match(r'\s*-\s*@(\w+):\s*(.+)', line)
            if match:
                short_name = match.group(1)
                full_info = match.group(2).strip()

                # Remove non-working days suffix if present, wherever it
                # falls relative to a calendar suffix (issue #1136)
                full_info, _nwd = _extract_non_working_suffix(full_info)

                # A calendar suffix (issue #1136) is metadata, not the role.
                full_info, _calendar_name = _strip_calendar_suffix(full_info)

                # Parts after the first comma are the role
                parts = [p.strip() for p in full_info.split(',')]
                if len(parts) > 1:
                    role_map[short_name.lower()] = ', '.join(parts[1:])

    return role_map


def parse_stakeholders_from_frontmatter(original_text):
    """Parse stakeholder entries from YAML front matter.

    Stakeholders are stored in the front matter under a ``Stakeholders:`` or
    ``Key Stakeholders:`` header using the format::

        Stakeholders:
        - @Name: Role, interest:high, influence:low

    Returns a list of dicts with keys: name, role, interest, influence.
    """
    stakeholders = []
    if not original_text:
        return stakeholders

    lines = original_text.split('\n')
    in_frontmatter = False
    in_stakeholders = False

    for line in lines:
        if line.strip() == '---':
            if not in_frontmatter:
                in_frontmatter = True
            else:
                break  # End of front matter
            continue

        if not in_frontmatter:
            continue

        trimmed = line.strip()

        if trimmed.lower() in ('stakeholders:', 'key stakeholders:'):
            in_stakeholders = True
            continue

        # Check if we've left the Stakeholders section
        if in_stakeholders and trimmed and not trimmed.startswith('-') and re.match(r'^[a-zA-Z\s]+:', trimmed):
            in_stakeholders = False

        if in_stakeholders and trimmed.startswith('- @'):
            entry = trimmed[2:].strip()  # Remove "- "
            item = _parse_stakeholder_entry(entry)
            if item:
                stakeholders.append(item)

    return stakeholders


def _parse_stakeholder_entry(entry):
    """Parse a single stakeholder entry string.

    Format: ``@shortname: Full Name, Role, interest:high, influence:low``

    Returns a dict with keys: shortname, name, role, interest, influence.
    """
    if not entry or not entry.startswith('@'):
        return None

    colon_idx = entry.find(':')
    if colon_idx == -1:
        raw = entry[1:].strip()  # strip @
        return {'shortname': raw.split()[0].lower() if raw else raw, 'name': raw, 'role': '', 'interest': 'low', 'influence': 'low'}

    shortname = entry[1:colon_idx].strip().lower()  # after @ before :
    rest = entry[colon_idx + 1:].strip()

    parts = [p.strip() for p in rest.split(',')]

    text_parts = []
    interest = 'low'
    influence = 'low'

    for part in parts:
        kv_match = re.match(r'^(interest|influence):\s*(high|low)$', part, re.IGNORECASE)
        if kv_match:
            key = kv_match.group(1).lower()
            val = kv_match.group(2).lower()
            if key == 'interest':
                interest = val
            elif key == 'influence':
                influence = val
        else:
            if part:
                text_parts.append(part)

    # First text part is name, rest is role
    name = text_parts[0] if text_parts else shortname
    role = ', '.join(text_parts[1:]) if len(text_parts) > 1 else ''

    return {'shortname': shortname, 'name': name, 'role': role, 'interest': interest, 'influence': influence}


def calculate_evm(tasks, budget_items=None):
    """Calculate Earned Value Management metrics from task and budget data.

    This mirrors the JavaScript ``calculateEVM`` function in the frontend.

    Args:
        tasks: List of scheduled task dicts from ``schedule_tasks()``.
        budget_items: Optional list of budget item dicts (with ``forecast``,
            ``estimate``, ``total`` keys).

    Returns:
        A dict with EVM metrics, or ``None`` if insufficient data.
    """
    if not tasks:
        return None

    # Get non-summary, non-milestone tasks with dates
    work_tasks = [
        t for t in tasks
        if not t.get('summary')
        and t.get('start') and t.get('finish')
        and t.get('duration') and (
            isinstance(t['duration'], timedelta) and t['duration'].days > 0
            or isinstance(t['duration'], (int, float)) and t['duration'] > 0
        )
    ]

    if not work_tasks:
        return None

    # Calculate overall % complete (weighted by duration)
    total_duration_days = 0
    weighted_complete = 0
    for t in work_tasks:
        dur = t.get('duration', timedelta(days=0))
        if isinstance(dur, timedelta):
            dur_days = dur.days
        else:
            dur_days = int(dur)
        pct = float(t.get('percent', 0) or 0)
        total_duration_days += dur_days
        weighted_complete += dur_days * pct

    overall_percent_complete = weighted_complete / total_duration_days if total_duration_days > 0 else 0

    # Project date range
    starts = [t['start'] for t in work_tasks if t.get('start')]
    finishes = [t['finish'] for t in work_tasks if t.get('finish')]
    if not starts or not finishes:
        return None

    project_start = min(starts)
    project_end = max(finishes)
    total_project_days = (project_end - project_start).days
    if total_project_days <= 0:
        return None

    today = datetime.today().replace(hour=0, minute=0, second=0, microsecond=0)

    # Time elapsed fraction (clamped 0-1)
    elapsed_days = max(0, (today - project_start).days)
    time_elapsed_fraction = min(1.0, elapsed_days / total_project_days)

    # BAC = total forecast from budget items, or fall back to task duration
    bac = 0.0
    ac = 0.0
    has_budget_data = False

    if budget_items:
        for item in budget_items:
            bac += float(item.get('forecast', 0) or item.get('estimate', 0) or 0)
            ac += float(item.get('total', 0) or 0)

    has_budget_data = bac > 0
    if not has_budget_data:
        bac = float(total_duration_days)
        ac = float(total_duration_days) * (overall_percent_complete / 100.0)

    # Core EVM calculations
    pv = bac * time_elapsed_fraction
    ev = bac * (overall_percent_complete / 100.0)

    # Variances
    cv = ev - ac  # Cost Variance
    sv = ev - pv  # Schedule Variance

    # Indices (guard against division by zero)
    cpi = ev / ac if ac != 0 else 0
    spi = ev / pv if pv != 0 else 0

    # Forecasts
    eac = bac / cpi if cpi != 0 else bac  # Estimate at Completion
    etc = eac - ac  # Estimate to Complete
    vac = bac - eac  # Variance at Completion

    return {
        'BAC': bac,
        'PV': pv,
        'EV': ev,
        'AC': ac,
        'CV': cv,
        'SV': sv,
        'CPI': cpi,
        'SPI': spi,
        'EAC': eac,
        'ETC': etc,
        'VAC': vac,
        'overall_percent_complete': overall_percent_complete,
        'time_elapsed_fraction': time_elapsed_fraction,
        'project_start': project_start,
        'project_end': project_end,
        'has_budget_data': has_budget_data,
    }


def calculate_resource_allocation(tasks, start_date, finish_date):
    """Calculate daily resource allocation in hours per day.

    Returns a dict with:
    - resources: list of resource names (display names)
    - days: list of dates from start to finish
    - allocation: dict[resource][date] = hours
    """
    from collections import defaultdict

    # Collect all resources (case-insensitive)
    resources_lower = set()
    resource_display_names = {}
    for t in tasks:
        res = t.get('resources', '')
        if res:
            for r in res.split(','):
                r_stripped = r.strip()
                r_key = r_stripped.lower()
                resources_lower.add(r_key)
                # Keep first occurrence's case for display
                if r_key not in resource_display_names:
                    resource_display_names[r_key] = r_stripped

    resources = sorted([resource_display_names[r] for r in resources_lower])

    # Generate all days from start to finish
    days = []
    current = start_date
    while current <= finish_date:
        days.append(current)
        current += timedelta(days=1)

    # Calculate allocation for each resource on each day
    allocation = defaultdict(lambda: defaultdict(float))

    for t in tasks:
        # Skip summary tasks - their children carry the resource allocations
        if t.get('summary'):
            continue

        task_start = t.get('start')
        task_finish = t.get('finish')
        if not task_start or not task_finish:
            continue

        # Get resources for this task
        res = t.get('resources', '')
        if not res:
            continue

        task_resources = [r.strip() for r in res.split(',')]
        if not task_resources:
            continue

        # Calculate working days for this task (exclude weekends)
        task_days = []
        current = task_start
        while current < task_finish:
            # 0 = Monday, 6 = Sunday
            if current.weekday() < 5:  # Monday-Friday
                task_days.append(current)
            current += timedelta(days=1)

        if not task_days:
            task_days = [task_start]

        # Assume 8 hours per working day, split among resources
        hours_per_day = 8.0 / len(task_resources)

        for resource in task_resources:
            # Use display name for allocation key
            resource_key = resource.lower()
            display_name = resource_display_names.get(resource_key, resource)
            for day in task_days:
                allocation[display_name][day] += hours_per_day

    return {
        'resources': resources,
        'days': days,
        'allocation': allocation
    }


def analyze_plan(text, original_text=None):
    """Review a plan for common problems (#782).

    A thin wrapper over :func:`noodle_core.plan_quality.review_plan`, kept
    for existing callers. Pass the original plan text (with front matter)
    as ``original_text``; ``text`` is used when it is not given.

    Returns:
        The findings, most severe first: each a dict with ``type`` (the
        check's title), ``severity`` (``error`` / ``warning`` /
        ``suggestion``), ``message``, ``fix``, ``docs``, ``line`` and
        ``check``, plus ``fix_action`` where a one-click fix exists.
    """
    from .plan_quality import review_plan

    result = review_plan(original_text if original_text is not None else text)
    return [dict(finding, type=finding["title"]) for finding in result["findings"]]


def export_timeline_to_powerpoint(text, output_path, is_yaml=True, project_name="Project", original_text=None):
    """Export project timeline to PowerPoint format.

    Args:
        text: Input text (YAML or natural language)
        output_path: Path to save the PowerPoint file
        is_yaml: If True, parse as YAML; if False, parse as natural language
        project_name: Project name to use
        original_text: Original text before conversion (for extracting resource mappings)
    """
    # Import here to avoid circular imports
    from .scheduling_engine import natural_language_to_yaml, schedule_tasks

    # Parse the text to get tasks and project info
    if is_yaml:
        data = yaml.safe_load(text)
        project_name = list(data.keys())[0]
        phases_raw = data[project_name]
    else:
        # Parse natural language
        data = natural_language_to_yaml(text, project_name)
        phases_raw = data[project_name]

    # If phases_raw is a list, pass as-is; if dict, wrap in a list
    if isinstance(phases_raw, list):
        phases = phases_raw
    elif isinstance(phases_raw, dict):
        phases = [phases_raw]
    else:
        phases = []

    # Schedule tasks
    tasks = schedule_tasks(phases)

    if not tasks:
        logger.warning("No tasks found to export")
        return

    # Calculate project dates
    today = datetime.today()
    start_date = min([t.get('start', today) for t in tasks if 'start' in t])
    finish_date = max([t.get('finish', today) for t in tasks if 'finish' in t])
    total_days = (finish_date - start_date).days or 1

    # Compute phase timelines
    phase_dates = {}
    for t in tasks:
        phase = t.get('phase')
        start = t.get('start') or today
        duration = t.get('duration') if 'duration' in t else timedelta(days=1)
        finish = t.get('finish') or (start + duration)
        if phase:
            if phase not in phase_dates:
                phase_dates[phase] = {'start': start, 'end': finish}
            else:
                if start < phase_dates[phase]['start']:
                    phase_dates[phase]['start'] = start
                if finish > phase_dates[phase]['end']:
                    phase_dates[phase]['end'] = finish

    # Collect milestones
    milestones = []
    for t in tasks:
        duration = t.get('duration', timedelta(days=1))
        if isinstance(duration, timedelta) and duration.days == 0:
            milestone_name = t.get('description', t.get('name', ''))
            display_name = milestone_name.replace('_', ' ')
            milestones.append({'name': display_name, 'date': t.get('start')})

    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Add blank slide
    blank_slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(blank_slide_layout)

    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = f"{project_name} Timeline"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(54, 96, 146)

    # Timeline dimensions
    timeline_left = Inches(1)
    timeline_top = Inches(2.5)
    timeline_width = Inches(8)
    timeline_height = Inches(0.15)

    # Draw timeline base line
    timeline_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        timeline_left,
        timeline_top,
        timeline_width,
        timeline_height
    )
    timeline_line.fill.solid()
    timeline_line.fill.fore_color.rgb = RGBColor(54, 96, 146)
    timeline_line.line.color.rgb = RGBColor(54, 96, 146)

    # Add start and end dates
    start_text = slide.shapes.add_textbox(timeline_left, timeline_top - Inches(0.5), Inches(1.5), Inches(0.3))
    start_frame = start_text.text_frame
    start_frame.text = f"Start\n{start_date.strftime('%d %b %Y')}"
    start_frame.paragraphs[0].font.size = Pt(10)
    start_frame.paragraphs[0].font.bold = True

    end_text = slide.shapes.add_textbox(timeline_left + timeline_width - Inches(1.5), timeline_top - Inches(0.5), Inches(1.5), Inches(0.3))
    end_frame = end_text.text_frame
    end_frame.text = f"Finish\n{finish_date.strftime('%d %b %Y')}"
    end_frame.paragraphs[0].font.size = Pt(10)
    end_frame.paragraphs[0].font.bold = True
    end_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    # Add phase boxes
    for phase, dates in phase_dates.items():
        phase_start_pos = (dates['start'] - start_date).days / total_days
        phase_end_pos = (dates['end'] - start_date).days / total_days
        phase_width = (phase_end_pos - phase_start_pos) * timeline_width.inches

        if phase_width < 0.1:
            phase_width = 0.1

        phase_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            timeline_left + Inches(phase_start_pos * timeline_width.inches),
            timeline_top + Inches(0.4),
            Inches(phase_width),
            Inches(0.5)
        )
        phase_box.fill.solid()
        phase_box.fill.fore_color.rgb = RGBColor(180, 198, 231)
        phase_box.line.color.rgb = RGBColor(54, 96, 146)

        # Add phase name
        phase_frame = phase_box.text_frame
        phase_frame.text = phase.replace('_', ' ')
        phase_frame.paragraphs[0].font.size = Pt(9)
        phase_frame.paragraphs[0].font.bold = True
        phase_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        phase_frame.vertical_anchor = 1  # Middle

    # Add milestones
    for milestone in milestones:
        m_pos = (milestone['date'] - start_date).days / total_days
        m_left = timeline_left + Inches(m_pos * timeline_width.inches)

        # Add diamond marker
        diamond = slide.shapes.add_shape(
            MSO_SHAPE.DIAMOND,
            m_left - Inches(0.15),
            timeline_top - Inches(0.075),
            Inches(0.3),
            Inches(0.3)
        )
        diamond.fill.solid()
        diamond.fill.fore_color.rgb = RGBColor(192, 0, 0)
        diamond.line.color.rgb = RGBColor(128, 0, 0)

        # Add milestone label above
        m_label = slide.shapes.add_textbox(
            m_left - Inches(0.75),
            timeline_top - Inches(1.2),
            Inches(1.5),
            Inches(0.4)
        )
        m_frame = m_label.text_frame
        m_frame.text = milestone['name']
        m_frame.paragraphs[0].font.size = Pt(9)
        m_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Add milestone date below
        m_date = slide.shapes.add_textbox(
            m_left - Inches(0.5),
            timeline_top + timeline_height + Inches(0.2),
            Inches(1),
            Inches(0.3)
        )
        m_date_frame = m_date.text_frame
        m_date_frame.text = milestone['date'].strftime('%d %b')
        m_date_frame.paragraphs[0].font.size = Pt(8)
        m_date_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Add legend
    legend_top = Inches(6)
    legend_left = Inches(1)

    # Phase legend
    phase_legend_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        legend_left,
        legend_top,
        Inches(0.3),
        Inches(0.2)
    )
    phase_legend_box.fill.solid()
    phase_legend_box.fill.fore_color.rgb = RGBColor(180, 198, 231)
    phase_legend_box.line.color.rgb = RGBColor(54, 96, 146)

    phase_legend_text = slide.shapes.add_textbox(legend_left + Inches(0.4), legend_top, Inches(1), Inches(0.2))
    phase_legend_text.text_frame.text = "Phase"
    phase_legend_text.text_frame.paragraphs[0].font.size = Pt(10)

    # Milestone legend
    milestone_legend = slide.shapes.add_shape(
        MSO_SHAPE.DIAMOND,
        legend_left + Inches(2),
        legend_top - Inches(0.05),
        Inches(0.3),
        Inches(0.3)
    )
    milestone_legend.fill.solid()
    milestone_legend.fill.fore_color.rgb = RGBColor(192, 0, 0)
    milestone_legend.line.color.rgb = RGBColor(128, 0, 0)

    milestone_legend_text = slide.shapes.add_textbox(legend_left + Inches(2.4), legend_top, Inches(1), Inches(0.2))
    milestone_legend_text.text_frame.text = "Milestone"
    milestone_legend_text.text_frame.paragraphs[0].font.size = Pt(10)

    # Save presentation
    prs.save(output_path)
    logger.info(f"Exported timeline to PowerPoint: {output_path}")


def _draw_timeline_graphic(slide, timeline_tasks, left, top, width):
    """Draw a minimal timeline graphic on the slide using native shapes."""
    from pptx.util import Emu
    from pptx.oxml.ns import qn

    # Separate phases and milestones
    phases = [t for t in timeline_tasks
              if t.get('is_summary') and t.get('start') and t.get('finish')]
    milestones = [t for t in timeline_tasks
                  if t.get('duration_days', 0) == 0 and not t.get('is_summary')
                  and t.get('finish')]

    if not phases and not milestones:
        return Inches(0)

    def _parse(d):
        return datetime.strptime(d, '%Y-%m-%d')

    all_dates = []
    for t in phases + milestones:
        if t.get('start'):
            all_dates.append(_parse(t['start']))
        if t.get('finish'):
            all_dates.append(_parse(t['finish']))

    if not all_dates:
        return Inches(0)

    min_date = min(all_dates) - timedelta(days=7)
    max_date = max(all_dates) + timedelta(days=7)
    total_days = (max_date - min_date).days + 1

    if total_days <= 0:
        return Inches(0)

    BLUE_SHADES = [
        RGBColor(21, 101, 192),
        RGBColor(25, 118, 210),
        RGBColor(30, 136, 229),
        RGBColor(33, 150, 243),
        RGBColor(66, 165, 245),
        RGBColor(100, 181, 246),
        RGBColor(144, 202, 249),
    ]
    GREEN = RGBColor(76, 175, 80)
    MILESTONE_BLUE = RGBColor(25, 118, 210)
    GREY_LINE = RGBColor(204, 204, 204)
    DARK_TEXT = RGBColor(100, 100, 100)
    TODAY_RED = RGBColor(220, 53, 69)

    bar_height = Inches(0.07)
    row_gap = Inches(0.008)
    width_emu = int(width)

    sorted_phases = sorted(phases, key=lambda p: _parse(p['start']))
    row_ends = []
    phase_rows = []
    for phase in sorted_phases:
        start_day = (_parse(phase['start']) - min_date).days
        end_day = (_parse(phase['finish']) - min_date).days
        assigned = -1
        for r_idx, r_end in enumerate(row_ends):
            if start_day >= r_end:
                assigned = r_idx
                break
        if assigned == -1:
            assigned = len(row_ends)
            row_ends.append(0)
        row_ends[assigned] = end_day
        phase_rows.append((phase, assigned))

    num_rows = max(len(row_ends), 1) if phases else 0
    phase_area_height = int(num_rows * (bar_height + row_gap))

    for idx, (phase, row) in enumerate(phase_rows):
        start_day = (_parse(phase['start']) - min_date).days
        end_day = (_parse(phase['finish']) - min_date).days
        x_start = int(left + (start_day / total_days) * width_emu)
        x_end = int(left + (end_day / total_days) * width_emu)
        bar_w = max(Inches(0.05), x_end - x_start)
        y = int(top + row * (bar_height + row_gap))

        percent = float(phase.get('percent', 0))
        is_complete = percent >= 100
        bg_color = GREEN if is_complete else BLUE_SHADES[idx % len(BLUE_SHADES)]

        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x_start, y, bar_w, bar_height
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = bg_color
        bar.line.fill.background()
        bar.adjustments[0] = 0.15

        if 0 < percent < 100:
            progress_w = int(bar_w * (percent / 100))
            if progress_w > 0:
                prog = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE, x_start, y,
                    progress_w, bar_height
                )
                prog.fill.solid()
                prog.fill.fore_color.rgb = RGBColor(166, 215, 168)
                prog.line.fill.background()
                prog.adjustments[0] = 0.15

        if bar_w > Inches(0.5):
            tf = bar.text_frame
            tf.word_wrap = False
            tf.margin_left = Inches(0.02)
            tf.margin_right = Inches(0.02)
            tf.margin_top = Inches(0)
            tf.margin_bottom = Inches(0)
            p = tf.paragraphs[0]
            label = phase.get('name', '')
            if is_complete:
                label = '\u2713 ' + label
            p.text = label
            p.font.size = Pt(5)
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.font.bold = False

    backbone_thickness = Inches(0.01)
    line_y = int(top + phase_area_height + Inches(0.01))
    backbone = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, int(left), line_y, width_emu, backbone_thickness
    )
    backbone.fill.solid()
    backbone.fill.fore_color.rgb = GREY_LINE
    backbone.line.fill.background()

    diamond_size = Inches(0.07)
    backbone_centre = int(line_y + backbone_thickness // 2)
    milestone_y = int(backbone_centre - diamond_size // 2)
    for ms in milestones:
        ms_date = _parse(ms['finish'])
        day_offset = (ms_date - min_date).days
        x_pos = int(left + (day_offset / total_days) * width_emu
                     - diamond_size // 2)
        percent = float(ms.get('percent', 0))
        ms_color = GREEN if percent >= 100 else MILESTONE_BLUE

        diamond = slide.shapes.add_shape(
            MSO_SHAPE.DIAMOND, x_pos, milestone_y,
            diamond_size, diamond_size
        )
        diamond.fill.solid()
        diamond.fill.fore_color.rgb = ms_color
        diamond.line.fill.background()

    date_label_y = int(line_y + backbone_thickness + Inches(0.01))

    if total_days <= 60:
        interval_days = max(2, total_days // 7)
    elif total_days <= 365:
        interval_days = max(7, total_days // 12)
    elif total_days <= 730:
        interval_days = 30
    else:
        interval_days = 90

    d = min_date
    while d <= max_date:
        day_offset = (d - min_date).days
        x_pos = int(left + (day_offset / total_days) * width_emu)
        if total_days <= 365:
            label = d.strftime('%b %Y').lower()
        else:
            label = d.strftime("%b '%y").lower()

        tb = slide.shapes.add_textbox(x_pos, date_label_y,
                                      Inches(0.5), Inches(0.10))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(4)
        p.font.color.rgb = DARK_TEXT
        d += timedelta(days=interval_days)

    today = datetime.now()
    if min_date <= today <= max_date:
        today_offset = (today - min_date).days
        today_x = int(left + (today_offset / total_days) * width_emu)
        marker_top = int(top)
        marker_height = int(date_label_y - top)
        today_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, today_x, marker_top,
            Inches(0.01), marker_height
        )
        today_line.fill.solid()
        today_line.fill.fore_color.rgb = TODAY_RED
        today_line.line.fill.background()
        today_tb = slide.shapes.add_textbox(
            today_x - Inches(0.12), marker_top - Inches(0.09),
            Inches(0.3), Inches(0.09)
        )
        tp = today_tb.text_frame.paragraphs[0]
        tp.text = "Today"
        tp.font.size = Pt(4)
        tp.font.color.rgb = TODAY_RED
        tp.font.bold = True
        tp.alignment = PP_ALIGN.CENTER

    total_height = int(date_label_y + Inches(0.10) - top)
    return total_height


def export_report_to_powerpoint(output_path, report_data):
    """Export the weekly project report to PowerPoint with a highlight slide."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_report_slide(prs, report_data)
    _add_highlight_slide(prs, report_data)

    prs.save(output_path)
    logger.info(f"Exported weekly report to PowerPoint: {output_path}")


def _add_highlight_slide(prs, report_data):
    """Add a dedicated slide for the latest highlight entry."""
    import re

    highlight = report_data.get('highlight')
    if not highlight or not highlight.get('content'):
        return  # No highlight — skip the slide

    DARK_BLUE = RGBColor(33, 60, 114)
    BLACK = RGBColor(0, 0, 0)

    def _sanitise(text):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(text))

    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Title bar
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.6))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    project_name = report_data.get('project_name', 'Project')
    p.text = _sanitise(f"{project_name} — Latest Highlight")
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    # Date and author meta
    meta_parts = []
    if highlight.get('date'):
        meta_parts.append(highlight['date'])
    if highlight.get('author'):
        meta_parts.append(f"@{highlight['author']}")
    if meta_parts:
        meta_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(12), Inches(0.3))
        mf = meta_box.text_frame
        mp = mf.paragraphs[0]
        mp.text = _sanitise("  ".join(meta_parts))
        mp.font.size = Pt(11)
        mp.font.color.rgb = RGBColor(100, 100, 100)
        mp.font.italic = True

    # Highlight content
    content_top = Inches(1.4)
    content_box = slide.shapes.add_textbox(
        Inches(0.5), content_top, Inches(12), Inches(5.5)
    )
    cf = content_box.text_frame
    cf.word_wrap = True

    content = highlight.get('content', '')
    lines = content.split('\n')
    for idx, line in enumerate(lines):
        if idx == 0:
            cp = cf.paragraphs[0]
        else:
            cp = cf.add_paragraph()
        clean = line.strip()
        # Convert markdown bold **text** to plain (PPTX doesn't render markdown)
        clean = re.sub(r'\*\*(.+?)\*\*', r'\1', clean)
        if clean.startswith('- '):
            clean = '\u2022 ' + clean[2:]
        is_heading = clean.startswith('**') or (line.strip().startswith('**') and line.strip().endswith('**'))
        cp.text = _sanitise(clean)
        cp.font.size = Pt(14) if is_heading else Pt(12)
        cp.font.bold = is_heading
        cp.font.color.rgb = DARK_BLUE if is_heading else BLACK
        cp.space_after = Pt(4)


def _add_report_slide(prs, report_data, include_footer=True):
    """Add a single project report slide to an existing Presentation."""
    DARK_BLUE = RGBColor(33, 60, 114)
    MID_BLUE = RGBColor(54, 96, 146)
    LIGHT_BLUE = RGBColor(180, 198, 231)
    WHITE = RGBColor(255, 255, 255)
    BLACK = RGBColor(0, 0, 0)
    LIGHT_GREY = RGBColor(242, 242, 242)
    RED = RGBColor(192, 0, 0)
    AMBER = RGBColor(218, 165, 32)
    GREEN = RGBColor(0, 128, 0)
    COMPLETE_BLUE = RGBColor(25, 118, 210)

    def _rag_colour(rag_str):
        colour = rag_status_to_colour(rag_str)
        if colour == 'red':
            return RED
        if colour == 'amber':
            return AMBER
        if colour == 'green':
            return GREEN
        if colour == 'blue':
            return COMPLETE_BLUE
        return BLACK

    def _sanitise_text(text):
        import re
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(text))

    def _set_cell_text(cell, text, font_size=8, bold=False, colour=None,
                       alignment=PP_ALIGN.LEFT):
        cell.text = _sanitise_text(text)
        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(font_size)
            para.font.bold = bold
            if colour:
                para.font.color.rgb = colour
            para.alignment = alignment
        cell.text_frame.word_wrap = True
        cell.text_frame.margin_top = Inches(0.02)
        cell.text_frame.margin_bottom = Inches(0.02)
        cell.text_frame.margin_left = Inches(0.04)
        cell.text_frame.margin_right = Inches(0.04)

    def _shade_header_row(table, col_count):
        for c in range(col_count):
            cell = table.cell(0, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = DARK_BLUE

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    project_name = report_data.get('project_name', 'Project')
    manager = report_data.get('manager', '')
    sponsor = report_data.get('sponsor', '')
    budget = report_data.get('budget', '')
    report_date = report_data.get('date', '')
    status = report_data.get('status', '')
    percent_complete = report_data.get('percent_complete', None)

    title_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
        Inches(13.333), Inches(0.85)
    )
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = DARK_BLUE
    title_bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.08),
                                  Inches(6), Inches(0.45))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = _sanitise_text(project_name)
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE

    details_parts = []
    if manager:
        details_parts.append(f"PM: {manager}")
    if sponsor:
        details_parts.append(f"Sponsor: {sponsor}")
    if budget:
        details_parts.append(f"Budget: {budget}")
    if report_date:
        details_parts.append(f"Date: {report_date}")
    if percent_complete is not None:
        details_parts.append(f"{percent_complete}% complete")
    detail_text = "   |   ".join(details_parts) if details_parts else ""

    if detail_text:
        dtb = slide.shapes.add_textbox(Inches(0.4), Inches(0.50),
                                       Inches(10), Inches(0.30))
        dtf = dtb.text_frame
        dtf.word_wrap = True
        dp = dtf.paragraphs[0]
        dp.text = _sanitise_text(detail_text)
        dp.font.size = Pt(10)
        dp.font.color.rgb = WHITE

    if status:
        status_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(11.5), Inches(0.15), Inches(1.6), Inches(0.55)
        )
        status_bg.fill.solid()
        status_bg.fill.fore_color.rgb = WHITE
        status_bg.line.fill.background()

        stb = slide.shapes.add_textbox(Inches(11.5), Inches(0.15),
                                       Inches(1.6), Inches(0.55))
        stf = stb.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        sp.text = _sanitise_text(status.upper())
        sp.font.size = Pt(16)
        sp.font.bold = True
        sp.font.color.rgb = _rag_colour(status)
        sp.alignment = PP_ALIGN.CENTER

    timeline_height_used = Inches(0)
    timeline_image_b64 = report_data.get('timeline_image')
    if timeline_image_b64:
        import base64
        import io
        try:
            img_data = base64.b64decode(timeline_image_b64)
            img_stream = io.BytesIO(img_data)
            tl_width = Inches(12.533)
            max_tl_height = Inches(1.8)
            pic = slide.shapes.add_picture(
                img_stream, Inches(0.4), Inches(1.05), width=tl_width
            )
            if pic.height > max_tl_height:
                aspect = pic.width / pic.height
                pic.height = max_tl_height
                pic.width = int(max_tl_height * aspect)
            timeline_height_used = pic.height + Inches(0.15)
        except Exception:
            logger.warning("Failed to embed timeline image in PPTX report",
                           exc_info=True)
            timeline_height_used = Inches(0)

    left_margin = Inches(0.4)
    right_margin = Inches(0.4)
    quad_gap = Inches(0.3)
    total_width = Inches(13.333) - left_margin - right_margin
    col_width = (total_width - quad_gap) // 2

    top_row_top = Inches(1.05) + timeline_height_used
    row_height = Inches(3.0)
    bottom_row_top = top_row_top + row_height + quad_gap
    right_col_left = left_margin + col_width + quad_gap

    def _add_section_heading(text, left, top, width):
        hbox = slide.shapes.add_textbox(left, top, width, Inches(0.30))
        hf = hbox.text_frame
        hf.word_wrap = True
        hp = hf.paragraphs[0]
        hp.text = text
        hp.font.size = Pt(13)
        hp.font.bold = True
        hp.font.color.rgb = MID_BLUE

    # TOP-LEFT: Milestones
    _add_section_heading("Milestones", left_margin, top_row_top, col_width)
    milestones = report_data.get('milestones', [])
    ms_table_top = top_row_top + Inches(0.35)

    if milestones:
        ms_rows = min(len(milestones), 10) + 1
        ms_table = slide.shapes.add_table(
            ms_rows, 3, left_margin, ms_table_top,
            col_width, Inches(0.26 * ms_rows)
        ).table
        w0 = int(col_width * 55 // 100)
        w1 = int(col_width * 30 // 100)
        w2 = int(col_width) - w0 - w1
        ms_table.columns[0].width = w0
        ms_table.columns[1].width = w1
        ms_table.columns[2].width = w2
        _shade_header_row(ms_table, 3)
        _set_cell_text(ms_table.cell(0, 0), "Milestone", 9, True, WHITE)
        _set_cell_text(ms_table.cell(0, 1), "Date", 9, True, WHITE, PP_ALIGN.CENTER)
        _set_cell_text(ms_table.cell(0, 2), "RAG", 9, True, WHITE, PP_ALIGN.CENTER)
        for i, ms in enumerate(milestones[:10]):
            row_idx = i + 1
            _set_cell_text(ms_table.cell(row_idx, 0), ms.get('name', ''))
            _set_cell_text(ms_table.cell(row_idx, 1), ms.get('date', ''),
                           8, False, None, PP_ALIGN.CENTER)
            rag = ms.get('rag', '')
            _set_cell_text(ms_table.cell(row_idx, 2), rag, 8, True,
                           _rag_colour(rag), PP_ALIGN.CENTER)
            if row_idx % 2 == 0:
                for c in range(3):
                    ms_table.cell(row_idx, c).fill.solid()
                    ms_table.cell(row_idx, c).fill.fore_color.rgb = LIGHT_GREY
    else:
        nb = slide.shapes.add_textbox(left_margin, ms_table_top, col_width, Inches(0.3))
        nb.text_frame.text = "No upcoming milestones."
        nb.text_frame.paragraphs[0].font.size = Pt(9)
        nb.text_frame.paragraphs[0].font.color.rgb = RGBColor(128, 128, 128)

    # TOP-RIGHT: Up Next
    _add_section_heading("Up Next", right_col_left, top_row_top, col_width)
    up_next = report_data.get('up_next', [])
    un_table_top = top_row_top + Inches(0.35)

    if up_next:
        un_rows = min(len(up_next), 10) + 1
        un_table = slide.shapes.add_table(
            un_rows, 4, right_col_left, un_table_top,
            col_width, Inches(0.26 * un_rows)
        ).table
        w0 = int(col_width * 40 // 100)
        w1 = int(col_width * 18 // 100)
        w2 = int(col_width * 18 // 100)
        w3 = int(col_width) - w0 - w1 - w2
        un_table.columns[0].width = w0
        un_table.columns[1].width = w1
        un_table.columns[2].width = w2
        un_table.columns[3].width = w3
        _shade_header_row(un_table, 4)
        _set_cell_text(un_table.cell(0, 0), "Task", 9, True, WHITE)
        _set_cell_text(un_table.cell(0, 1), "Start", 9, True, WHITE, PP_ALIGN.CENTER)
        _set_cell_text(un_table.cell(0, 2), "Finish", 9, True, WHITE, PP_ALIGN.CENTER)
        _set_cell_text(un_table.cell(0, 3), "RAG", 9, True, WHITE, PP_ALIGN.CENTER)
        for i, item in enumerate(up_next[:10]):
            row_idx = i + 1
            _set_cell_text(un_table.cell(row_idx, 0), item.get('name', ''))
            _set_cell_text(un_table.cell(row_idx, 1), item.get('start', ''),
                           8, False, None, PP_ALIGN.CENTER)
            _set_cell_text(un_table.cell(row_idx, 2), item.get('finish', ''),
                           8, False, None, PP_ALIGN.CENTER)
            rag_text = item.get('rag', '')
            _set_cell_text(un_table.cell(row_idx, 3), rag_text, 8, True,
                           _rag_colour(rag_text), PP_ALIGN.CENTER)
            if row_idx % 2 == 0:
                for c in range(4):
                    un_table.cell(row_idx, c).fill.solid()
                    un_table.cell(row_idx, c).fill.fore_color.rgb = LIGHT_GREY
    else:
        nb = slide.shapes.add_textbox(right_col_left, un_table_top, col_width, Inches(0.3))
        nb.text_frame.text = "No upcoming tasks in the next 2 weeks."
        nb.text_frame.paragraphs[0].font.size = Pt(9)
        nb.text_frame.paragraphs[0].font.color.rgb = RGBColor(128, 128, 128)

    # BOTTOM-LEFT: Latest Highlight
    _add_section_heading("Latest Highlight", left_margin, bottom_row_top, col_width)
    highlight = report_data.get('highlight')
    hl_top = bottom_row_top + Inches(0.35)

    if highlight and highlight.get('content'):
        meta_parts = []
        if highlight.get('date'):
            meta_parts.append(highlight['date'])
        if highlight.get('author'):
            meta_parts.append(f"@{highlight['author']}")
        meta_text = "  ".join(meta_parts)
        if meta_text:
            mb = slide.shapes.add_textbox(left_margin, hl_top, col_width, Inches(0.25))
            mf = mb.text_frame
            mf.word_wrap = True
            mp = mf.paragraphs[0]
            mp.text = _sanitise_text(meta_text)
            mp.font.size = Pt(8)
            mp.font.color.rgb = RGBColor(100, 100, 100)
            mp.font.italic = True
            hl_top += Inches(0.25)
        cb = slide.shapes.add_textbox(left_margin, hl_top, col_width, Inches(2.3))
        cf = cb.text_frame
        cf.word_wrap = True
        content = highlight.get('content', '')
        lines = content.split('\n')
        for idx, line in enumerate(lines):
            if idx == 0:
                cp = cf.paragraphs[0]
            else:
                cp = cf.add_paragraph()
            clean = line.strip()
            if clean.startswith('- '):
                clean = '\u2022 ' + clean[2:]
            cp.text = _sanitise_text(clean)
            cp.font.size = Pt(9)
            cp.font.color.rgb = BLACK
    else:
        nb = slide.shapes.add_textbox(left_margin, hl_top, col_width, Inches(0.3))
        nb.text_frame.text = "No highlights recorded yet."
        nb.text_frame.paragraphs[0].font.size = Pt(9)
        nb.text_frame.paragraphs[0].font.color.rgb = RGBColor(128, 128, 128)

    # BOTTOM-RIGHT: Risks & Issues
    _add_section_heading("Risks & Issues", right_col_left, bottom_row_top, col_width)
    risks_issues = report_data.get('risks_issues', [])
    ri_table_top = bottom_row_top + Inches(0.35)

    if risks_issues:
        ri_cols = 5
        ri_rows = min(len(risks_issues), 10) + 1
        ri_table = slide.shapes.add_table(
            ri_rows, ri_cols, right_col_left, ri_table_top,
            col_width, Inches(0.26 * ri_rows)
        ).table
        w0 = int(col_width * 10 // 100)
        w1 = int(col_width * 20 // 100)
        w2 = int(col_width * 28 // 100)
        w3 = int(col_width * 28 // 100)
        w4 = int(col_width) - w0 - w1 - w2 - w3
        ri_table.columns[0].width = w0
        ri_table.columns[1].width = w1
        ri_table.columns[2].width = w2
        ri_table.columns[3].width = w3
        ri_table.columns[4].width = w4
        _shade_header_row(ri_table, ri_cols)
        _set_cell_text(ri_table.cell(0, 0), "Type", 9, True, WHITE, PP_ALIGN.CENTER)
        _set_cell_text(ri_table.cell(0, 1), "Title", 9, True, WHITE)
        _set_cell_text(ri_table.cell(0, 2), "Description", 9, True, WHITE)
        _set_cell_text(ri_table.cell(0, 3), "Mitigation", 9, True, WHITE)
        _set_cell_text(ri_table.cell(0, 4), "Score", 9, True, WHITE, PP_ALIGN.CENTER)
        for i, item in enumerate(risks_issues[:10]):
            row_idx = i + 1
            item_type = item.get('type', '')
            _set_cell_text(ri_table.cell(row_idx, 0), item_type.capitalize(),
                           8, True, None, PP_ALIGN.CENTER)
            _set_cell_text(ri_table.cell(row_idx, 1), item.get('title', ''))
            _set_cell_text(ri_table.cell(row_idx, 2),
                           item.get('description', ''))
            _set_cell_text(ri_table.cell(row_idx, 3),
                           item.get('mitigation', ''))
            score = item.get('score', 0)
            score_colour = RED if score >= 16 else (AMBER if score >= 6 else GREEN)
            _set_cell_text(ri_table.cell(row_idx, 4), str(score), 8,
                           True, score_colour, PP_ALIGN.CENTER)
            if row_idx % 2 == 0:
                for c in range(ri_cols):
                    ri_table.cell(row_idx, c).fill.solid()
                    ri_table.cell(row_idx, c).fill.fore_color.rgb = LIGHT_GREY
    else:
        nb = slide.shapes.add_textbox(right_col_left, ri_table_top, col_width, Inches(0.3))
        nb.text_frame.text = "No open risks or issues."
        nb.text_frame.paragraphs[0].font.size = Pt(9)
        nb.text_frame.paragraphs[0].font.color.rgb = RGBColor(128, 128, 128)

    if include_footer:
        fb = slide.shapes.add_textbox(Inches(0.4), Inches(7.1),
                                      Inches(4), Inches(0.25))
        ff = fb.text_frame
        fp = ff.paragraphs[0]
        fp.text = f"Generated by Noodle Planner  |  {report_date}"
        fp.font.size = Pt(7)
        fp.font.color.rgb = RGBColor(160, 160, 160)

    return slide


def _parse_budget_value(budget_str):
    """Parse a budget string into a numeric value."""
    import re as _re
    if not budget_str:
        return None
    text = str(budget_str).strip()
    text = _re.sub(r'[\u00a3$\u20ac\u00a5\s,]', '', text)
    if not text:
        return None
    multiplier = 1
    if text[-1].lower() == 'k':
        multiplier = 1_000
        text = text[:-1]
    elif text[-1].lower() == 'm':
        multiplier = 1_000_000
        text = text[:-1]
    elif text[-1].lower() == 'b':
        multiplier = 1_000_000_000
        text = text[:-1]
    try:
        return float(text) * multiplier
    except ValueError:
        return None


def _format_budget_total(total):
    """Format a numeric budget total as a human-readable string."""
    if total >= 1_000_000_000:
        formatted = f"${total / 1_000_000_000:,.1f}B"
    elif total >= 1_000_000:
        formatted = f"${total / 1_000_000:,.1f}M"
    elif total >= 1_000:
        formatted = f"${total:,.0f}"
    else:
        formatted = f"${total:,.0f}"
    return formatted


def _calculate_total_portfolio_budget(projects):
    """Calculate the total budget across all projects."""
    total = 0
    has_any = False
    for proj in projects:
        value = _parse_budget_value(proj.get('budget', ''))
        if value is not None:
            total += value
            has_any = True
    return total if has_any else None


def _add_portfolio_overview_slide(prs, portfolio_data):
    """Add a portfolio overview slide with status dashboard and timeline."""
    DARK_BLUE = RGBColor(33, 60, 114)
    MID_BLUE = RGBColor(54, 96, 146)
    WHITE = RGBColor(255, 255, 255)
    BLACK = RGBColor(0, 0, 0)
    LIGHT_GREY = RGBColor(242, 242, 242)
    RED = RGBColor(192, 0, 0)
    AMBER = RGBColor(218, 165, 32)
    GREEN = RGBColor(0, 128, 0)
    COMPLETE_BLUE = RGBColor(25, 118, 210)

    import re as _re

    def _sanitise_text(text):
        return _re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(text))

    def _rag_rgb(rag_str):
        val = str(rag_str).lower().strip()
        if val in ('red', 'r'):
            return RED
        if val in ('amber', 'a', 'yellow'):
            return AMBER
        if val in ('green', 'g'):
            return GREEN
        if val in ('blue', 'complete'):
            return COMPLETE_BLUE
        return BLACK

    def _set_cell_text(cell, text, font_size=8, bold=False, colour=None,
                       alignment=PP_ALIGN.LEFT):
        cell.text = _sanitise_text(text)
        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(font_size)
            para.font.bold = bold
            if colour:
                para.font.color.rgb = colour
            para.alignment = alignment
        cell.text_frame.word_wrap = True
        cell.text_frame.margin_top = Inches(0.02)
        cell.text_frame.margin_bottom = Inches(0.02)
        cell.text_frame.margin_left = Inches(0.04)
        cell.text_frame.margin_right = Inches(0.04)

    def _shade_header_row(table, col_count):
        for c in range(col_count):
            cell = table.cell(0, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = DARK_BLUE

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    portfolio_name = portfolio_data.get('portfolio_name', 'Portfolio')
    report_date = portfolio_data.get('date', '')
    projects = portfolio_data.get('projects', [])

    title_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
        Inches(13.333), Inches(0.85)
    )
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = DARK_BLUE
    title_bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.08),
                                  Inches(8), Inches(0.45))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = _sanitise_text(f"{portfolio_name} Overview")
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE

    total_budget = _calculate_total_portfolio_budget(projects)

    if report_date:
        date_text = f"Date: {report_date}"
        if total_budget is not None:
            date_text += f"    |    Total Budget: {_format_budget_total(total_budget)}"
        dtb = slide.shapes.add_textbox(Inches(0.4), Inches(0.50),
                                       Inches(8), Inches(0.30))
        dtf = dtb.text_frame
        dtf.word_wrap = True
        dp = dtf.paragraphs[0]
        dp.text = _sanitise_text(date_text)
        dp.font.size = Pt(10)
        dp.font.color.rgb = WHITE
    elif total_budget is not None:
        dtb = slide.shapes.add_textbox(Inches(0.4), Inches(0.50),
                                       Inches(8), Inches(0.30))
        dtf = dtb.text_frame
        dtf.word_wrap = True
        dp = dtf.paragraphs[0]
        dp.text = _sanitise_text(f"Total Budget: {_format_budget_total(total_budget)}")
        dp.font.size = Pt(10)
        dp.font.color.rgb = WHITE

    rag_counts = {'red': 0, 'amber': 0, 'green': 0}
    for proj in projects:
        rag = str(proj.get('rag', '')).lower()
        if rag in rag_counts:
            rag_counts[rag] += 1

    badge_bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(0.15),
        Inches(4.3), Inches(0.55)
    )
    badge_bg.fill.solid()
    badge_bg.fill.fore_color.rgb = WHITE
    badge_bg.line.fill.background()

    stb = slide.shapes.add_textbox(Inches(8.9), Inches(0.25),
                                   Inches(4.1), Inches(0.40))
    stf = stb.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.alignment = PP_ALIGN.RIGHT

    run_proj = sp.add_run()
    run_proj.text = _sanitise_text(f"{len(projects)} Projects   ")
    run_proj.font.size = Pt(11)
    run_proj.font.bold = True
    run_proj.font.color.rgb = DARK_BLUE

    run_green = sp.add_run()
    run_green.text = _sanitise_text(f"Green: {rag_counts['green']}  ")
    run_green.font.size = Pt(11)
    run_green.font.bold = True
    run_green.font.color.rgb = GREEN

    run_amber = sp.add_run()
    run_amber.text = _sanitise_text(f"Amber: {rag_counts['amber']}  ")
    run_amber.font.size = Pt(11)
    run_amber.font.bold = True
    run_amber.font.color.rgb = AMBER

    run_red = sp.add_run()
    run_red.text = _sanitise_text(f"Red: {rag_counts['red']}")
    run_red.font.size = Pt(11)
    run_red.font.bold = True
    run_red.font.color.rgb = RED

    dashboard_heading = slide.shapes.add_textbox(
        Inches(0.4), Inches(1.05), Inches(6), Inches(0.30))
    dhf = dashboard_heading.text_frame
    dhf.word_wrap = True
    dhp = dhf.paragraphs[0]
    dhp.text = "Project Status Dashboard"
    dhp.font.size = Pt(14)
    dhp.font.bold = True
    dhp.font.color.rgb = MID_BLUE

    if projects:
        table_top = Inches(1.45)
        num_rows = min(len(projects), 20) + 1
        num_cols = 6
        table_width = Inches(12.533)
        table_height = Inches(0.28 * num_rows)

        tbl = slide.shapes.add_table(
            num_rows, num_cols, Inches(0.4), table_top,
            table_width, table_height
        ).table

        tbl.columns[0].width = int(table_width * 28 // 100)
        tbl.columns[1].width = int(table_width * 12 // 100)
        tbl.columns[2].width = int(table_width * 18 // 100)
        tbl.columns[3].width = int(table_width * 18 // 100)
        tbl.columns[4].width = int(table_width * 10 // 100)
        tbl.columns[5].width = (int(table_width)
                                - tbl.columns[0].width
                                - tbl.columns[1].width
                                - tbl.columns[2].width
                                - tbl.columns[3].width
                                - tbl.columns[4].width)

        _shade_header_row(tbl, num_cols)
        _set_cell_text(tbl.cell(0, 0), "Project", 9, True, WHITE)
        _set_cell_text(tbl.cell(0, 1), "Budget", 9, True, WHITE, PP_ALIGN.CENTER)
        _set_cell_text(tbl.cell(0, 2), "Status", 9, True, WHITE, PP_ALIGN.CENTER)
        _set_cell_text(tbl.cell(0, 3), "Progress", 9, True, WHITE, PP_ALIGN.CENTER)
        _set_cell_text(tbl.cell(0, 4), "RAG", 9, True, WHITE, PP_ALIGN.CENTER)
        _set_cell_text(tbl.cell(0, 5), "Open Risks", 9, True, WHITE, PP_ALIGN.CENTER)

        for i, proj in enumerate(projects[:20]):
            row_idx = i + 1
            _set_cell_text(tbl.cell(row_idx, 0), proj.get('name', ''), 9)
            _set_cell_text(tbl.cell(row_idx, 1), proj.get('budget', ''),
                           8, False, None, PP_ALIGN.CENTER)
            _set_cell_text(tbl.cell(row_idx, 2), proj.get('status', ''),
                           8, False, None, PP_ALIGN.CENTER)
            completion = proj.get('completion', 0)
            _set_cell_text(tbl.cell(row_idx, 3), f"{completion}%",
                           8, False, None, PP_ALIGN.CENTER)
            rag = proj.get('rag', '')
            _set_cell_text(tbl.cell(row_idx, 4), rag.upper() if rag else '',
                           8, True, _rag_rgb(rag), PP_ALIGN.CENTER)
            _set_cell_text(tbl.cell(row_idx, 5),
                           str(proj.get('risk_count', 0)),
                           8, False, None, PP_ALIGN.CENTER)
            if row_idx % 2 == 0:
                for c in range(num_cols):
                    tbl.cell(row_idx, c).fill.solid()
                    tbl.cell(row_idx, c).fill.fore_color.rgb = LIGHT_GREY
            tbl.cell(row_idx, 4).fill.solid()
            tbl.cell(row_idx, 4).fill.fore_color.rgb = WHITE

        timeline_top = table_top + table_height + Inches(0.3)
    else:
        nb = slide.shapes.add_textbox(Inches(0.4), Inches(1.45),
                                      Inches(6), Inches(0.3))
        nb.text_frame.text = "No projects in portfolio."
        nb.text_frame.paragraphs[0].font.size = Pt(9)
        nb.text_frame.paragraphs[0].font.color.rgb = RGBColor(128, 128, 128)
        timeline_top = Inches(2.0)

    timeline_image_b64 = portfolio_data.get('timeline_image')
    if timeline_image_b64:
        import base64
        import io
        try:
            img_data = base64.b64decode(timeline_image_b64)
            img_stream = io.BytesIO(img_data)
            tl_left = Inches(0.4)
            tl_top = timeline_top
            tl_width = Inches(12.533)
            max_tl_height = Inches(1.8)
            pic = slide.shapes.add_picture(
                img_stream, tl_left, tl_top, width=tl_width
            )
            if pic.height > max_tl_height:
                aspect = pic.width / pic.height
                pic.height = max_tl_height
                pic.width = int(max_tl_height * aspect)
        except Exception:
            logger.warning("Failed to embed portfolio timeline image",
                           exc_info=True)

    fb = slide.shapes.add_textbox(Inches(0.4), Inches(7.1),
                                  Inches(4), Inches(0.25))
    ff = fb.text_frame
    fp = ff.paragraphs[0]
    fp.text = f"Generated by Noodle Planner  |  {report_date}"
    fp.font.size = Pt(7)
    fp.font.color.rgb = RGBColor(160, 160, 160)

    return slide


def _collect_portfolio_risks(project_reports):
    """Collect Medium and High open risks/issues across all projects."""
    risks = []
    seen = set()
    for report in project_reports:
        project_name = report.get('project_name', '')
        for item in report.get('risks_issues', []):
            score = item.get('score', 0)
            if score < 6:
                continue
            # Deduplicate by project, type, and title to prevent the
            # same risk from appearing multiple times.
            dedup_key = (
                project_name,
                item.get('type', ''),
                item.get('title', ''),
            )
            if dedup_key in seen:
                continue
            seen.add(dedup_key)
            rag = 'red' if score >= 16 else 'amber'
            risks.append({
                'project_name': project_name,
                'type': item.get('type', ''),
                'title': item.get('title', ''),
                'description': item.get('description', ''),
                'mitigation': item.get('mitigation', ''),
                'score': score,
                'rag': rag,
            })
    risks.sort(key=lambda r: r['score'], reverse=True)
    return risks


def _add_portfolio_risk_slides(prs, portfolio_data, risks):
    """Add one or more portfolio risk slides showing Medium and High risks."""
    DARK_BLUE = RGBColor(33, 60, 114)
    MID_BLUE = RGBColor(54, 96, 146)
    WHITE = RGBColor(255, 255, 255)
    BLACK = RGBColor(0, 0, 0)
    LIGHT_GREY = RGBColor(242, 242, 242)
    RED = RGBColor(192, 0, 0)
    AMBER = RGBColor(218, 165, 32)
    GREEN = RGBColor(0, 128, 0)

    import re as _re

    def _sanitise_text(text):
        return _re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(text))

    def _set_cell_text(cell, text, font_size=8, bold=False, colour=None,
                       alignment=PP_ALIGN.LEFT):
        cell.text = _sanitise_text(text)
        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(font_size)
            para.font.bold = bold
            if colour:
                para.font.color.rgb = colour
            para.alignment = alignment
        cell.text_frame.word_wrap = True
        cell.text_frame.margin_top = Inches(0.02)
        cell.text_frame.margin_bottom = Inches(0.02)
        cell.text_frame.margin_left = Inches(0.04)
        cell.text_frame.margin_right = Inches(0.04)

    def _shade_header_row(table, col_count):
        for c in range(col_count):
            cell = table.cell(0, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = DARK_BLUE

    def _score_colour(score):
        if score >= 16:
            return RED
        if score >= 6:
            return AMBER
        return GREEN

    portfolio_name = portfolio_data.get('portfolio_name', 'Portfolio')
    report_date = portfolio_data.get('date', '')

    max_rows_per_slide = 18
    slides = []
    total_pages = max(1, (len(risks) + max_rows_per_slide - 1) // max_rows_per_slide)

    for page_idx in range(total_pages):
        page_risks = risks[page_idx * max_rows_per_slide:(page_idx + 1) * max_rows_per_slide]

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slides.append(slide)

        title_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
            Inches(13.333), Inches(0.85)
        )
        title_bar.fill.solid()
        title_bar.fill.fore_color.rgb = DARK_BLUE
        title_bar.line.fill.background()

        title_text = f"{portfolio_name} Risk Register"
        if total_pages > 1:
            title_text += f" ({page_idx + 1}/{total_pages})"

        tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.08),
                                      Inches(8), Inches(0.45))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = _sanitise_text(title_text)
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = WHITE

        if report_date:
            dtb = slide.shapes.add_textbox(Inches(0.4), Inches(0.50),
                                           Inches(6), Inches(0.30))
            dtf = dtb.text_frame
            dtf.word_wrap = True
            dp = dtf.paragraphs[0]
            dp.text = _sanitise_text(f"Date: {report_date}")
            dp.font.size = Pt(10)
            dp.font.color.rgb = WHITE

        if page_idx == 0:
            high_count = sum(1 for r in risks if r['rag'] == 'red')
            medium_count = sum(1 for r in risks if r['rag'] == 'amber')
            summary_text = (f"{len(risks)} Risks/Issues  |  "
                            f"High: {high_count}  Medium: {medium_count}")
            stb = slide.shapes.add_textbox(Inches(9), Inches(0.25),
                                           Inches(4), Inches(0.40))
            stf = stb.text_frame
            stf.word_wrap = True
            sp = stf.paragraphs[0]
            sp.text = _sanitise_text(summary_text)
            sp.font.size = Pt(11)
            sp.font.bold = True
            sp.font.color.rgb = WHITE
            sp.alignment = PP_ALIGN.RIGHT

        heading_text = "Open Risks & Issues (Medium and High)"
        heading_box = slide.shapes.add_textbox(
            Inches(0.4), Inches(1.05), Inches(6), Inches(0.30))
        hf = heading_box.text_frame
        hf.word_wrap = True
        hp = hf.paragraphs[0]
        hp.text = heading_text
        hp.font.size = Pt(14)
        hp.font.bold = True
        hp.font.color.rgb = MID_BLUE

        if page_risks:
            table_top = Inches(1.45)
            num_rows = len(page_risks) + 1
            num_cols = 7
            table_width = Inches(12.533)
            table_height = Inches(0.28 * num_rows)

            tbl = slide.shapes.add_table(
                num_rows, num_cols, Inches(0.4), table_top,
                table_width, table_height
            ).table

            tbl.columns[0].width = int(table_width * 12 // 100)
            tbl.columns[1].width = int(table_width * 7 // 100)
            tbl.columns[2].width = int(table_width * 18 // 100)
            tbl.columns[3].width = int(table_width * 24 // 100)
            tbl.columns[4].width = int(table_width * 24 // 100)
            tbl.columns[5].width = int(table_width * 7 // 100)
            tbl.columns[6].width = (int(table_width) - tbl.columns[0].width
                                    - tbl.columns[1].width - tbl.columns[2].width
                                    - tbl.columns[3].width - tbl.columns[4].width
                                    - tbl.columns[5].width)

            _shade_header_row(tbl, num_cols)
            _set_cell_text(tbl.cell(0, 0), "Project", 9, True, WHITE)
            _set_cell_text(tbl.cell(0, 1), "Type", 9, True, WHITE, PP_ALIGN.CENTER)
            _set_cell_text(tbl.cell(0, 2), "Title", 9, True, WHITE)
            _set_cell_text(tbl.cell(0, 3), "Description", 9, True, WHITE)
            _set_cell_text(tbl.cell(0, 4), "Mitigation", 9, True, WHITE)
            _set_cell_text(tbl.cell(0, 5), "Score", 9, True, WHITE, PP_ALIGN.CENTER)
            _set_cell_text(tbl.cell(0, 6), "RAG", 9, True, WHITE, PP_ALIGN.CENTER)

            for i, risk in enumerate(page_risks):
                row_idx = i + 1
                _set_cell_text(tbl.cell(row_idx, 0),
                               risk.get('project_name', ''), 8)
                _set_cell_text(tbl.cell(row_idx, 1),
                               risk.get('type', '').capitalize(),
                               8, True, None, PP_ALIGN.CENTER)
                _set_cell_text(tbl.cell(row_idx, 2),
                               risk.get('title', ''), 8)
                _set_cell_text(tbl.cell(row_idx, 3),
                               risk.get('description', ''), 8)
                _set_cell_text(tbl.cell(row_idx, 4),
                               risk.get('mitigation', ''), 8)
                score = risk.get('score', 0)
                _set_cell_text(tbl.cell(row_idx, 5), str(score),
                               8, True, _score_colour(score), PP_ALIGN.CENTER)
                rag = risk.get('rag', '')
                rag_label = 'HIGH' if rag == 'red' else 'MEDIUM'
                rag_colour = RED if rag == 'red' else AMBER
                _set_cell_text(tbl.cell(row_idx, 6), rag_label,
                               8, True, rag_colour, PP_ALIGN.CENTER)
                if row_idx % 2 == 0:
                    for c in range(num_cols):
                        tbl.cell(row_idx, c).fill.solid()
                        tbl.cell(row_idx, c).fill.fore_color.rgb = LIGHT_GREY
        else:
            nb = slide.shapes.add_textbox(Inches(0.4), Inches(1.45),
                                          Inches(6), Inches(0.3))
            nb.text_frame.text = "No medium or high risks/issues found."
            nb.text_frame.paragraphs[0].font.size = Pt(9)
            nb.text_frame.paragraphs[0].font.color.rgb = RGBColor(128, 128, 128)

        fb = slide.shapes.add_textbox(Inches(0.4), Inches(7.1),
                                      Inches(4), Inches(0.25))
        ff = fb.text_frame
        fp = ff.paragraphs[0]
        fp.text = f"Generated by Noodle Planner  |  {report_date}"
        fp.font.size = Pt(7)
        fp.font.color.rgb = RGBColor(160, 160, 160)

    return slides


def _add_portfolio_deliverables_slides(prs, project_reports):
    """Add deliverables matrix slides to the portfolio appendix.

    For each project that has deliverables data, a slide is created with a
    table showing deliverable tasks, their dates, status, role assignments
    (P/R/A), and a QA column indicating whether all three roles are present.

    Args:
        prs: python-pptx Presentation object.
        project_reports: List of per-project report dicts.  Each may contain
            a ``deliverables`` key with sub-keys ``items``, ``people``, and
            ``role_map``.
    """
    from pptx.oxml.ns import qn

    DARK_BLUE = RGBColor(33, 60, 114)
    WHITE = RGBColor(255, 255, 255)
    BLACK = RGBColor(0, 0, 0)
    LIGHT_GREY = RGBColor(242, 242, 242)

    # Role colours matching the web / Excel exports
    ROLE_COLOURS = {
        'P': RGBColor(0x44, 0x72, 0xC4),   # Blue
        'R': RGBColor(0xED, 0x7D, 0x31),   # Orange
        'A': RGBColor(0x70, 0xAD, 0x47),   # Green
    }
    QA_GREEN = RGBColor(0x70, 0xAD, 0x47)

    import re as _re

    def _sanitise(text):
        return _re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(text))

    def _set_cell(cell, text, font_size=8, bold=False, colour=None,
                  alignment=PP_ALIGN.LEFT):
        cell.text = _sanitise(text)
        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(font_size)
            para.font.bold = bold
            if colour:
                para.font.color.rgb = colour
            para.alignment = alignment
        cell.text_frame.word_wrap = True
        cell.text_frame.margin_top = Inches(0.02)
        cell.text_frame.margin_bottom = Inches(0.02)
        cell.text_frame.margin_left = Inches(0.04)
        cell.text_frame.margin_right = Inches(0.04)

    def _set_cell_vertical(cell):
        """Set vertical (bottom-to-top) text direction on a table cell."""
        tc = cell._tc
        tcPr = tc.find(qn('a:tcPr'))
        if tcPr is None:
            tcPr = tc.makeelement(qn('a:tcPr'), {})
            tc.insert(0, tcPr)
        tcPr.set('vert', 'vert270')

    for report in project_reports:
        deliverables_data = report.get('deliverables')
        if not deliverables_data:
            continue

        items = deliverables_data.get('items', [])
        if not items:
            continue

        people = deliverables_data.get('people', [])
        role_map = deliverables_data.get('role_map', {})

        project_name = report.get('project_name', 'Project')

        # Fixed columns: ID, Deliverable, Dates, Status
        fixed_count = 4
        person_count = len(people)
        total_cols = fixed_count + person_count + 1  # +1 for QA

        # Limit rows to avoid overflowing the slide
        max_rows = min(len(items), 18)
        num_rows = max_rows + 1  # +1 for header

        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Title bar
        title_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
            Inches(13.333), Inches(0.85)
        )
        title_bar.fill.solid()
        title_bar.fill.fore_color.rgb = DARK_BLUE
        title_bar.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.08),
                                      Inches(10), Inches(0.45))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = _sanitise(f"Deliverables Matrix \u2014 {project_name}")
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = WHITE

        # Build table
        table_left = Inches(0.4)
        table_top = Inches(1.1)
        table_width = Inches(12.533)
        row_height = Inches(0.28)
        table_height = row_height * num_rows

        shape = slide.shapes.add_table(
            num_rows, total_cols, table_left, table_top,
            table_width, table_height
        )
        tbl = shape.table

        # Column widths
        id_width = int(table_width * 4 / 100)
        deliv_width = int(table_width * 22 / 100)
        dates_width = int(table_width * 18 / 100)
        status_width = int(table_width * 10 / 100)
        qa_width = int(table_width * 4 / 100)
        fixed_total = id_width + deliv_width + dates_width + status_width + qa_width
        remaining = int(table_width) - fixed_total
        person_col_width = max(remaining // max(person_count, 1), Inches(0.3))

        tbl.columns[0].width = id_width
        tbl.columns[1].width = deliv_width
        tbl.columns[2].width = dates_width
        tbl.columns[3].width = status_width
        for pi in range(person_count):
            tbl.columns[fixed_count + pi].width = person_col_width
        tbl.columns[total_cols - 1].width = qa_width

        # Header row
        tbl.rows[0].height = Inches(0.9)
        header_labels = ['ID', 'Deliverable', 'Dates', 'Status']
        person_headers = [role_map.get(short, short) for short in people]
        header_labels += person_headers + ['QA']

        for ci, label in enumerate(header_labels):
            cell = tbl.cell(0, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = DARK_BLUE
            is_person_col = fixed_count <= ci < fixed_count + person_count
            _set_cell(cell, label, 8, True, WHITE, PP_ALIGN.CENTER)
            if is_person_col:
                _set_cell_vertical(cell)

        # Data rows
        for ri, item in enumerate(items[:max_rows]):
            row_idx = ri + 1
            _set_cell(tbl.cell(row_idx, 0), str(ri + 1), 8, False, None, PP_ALIGN.CENTER)
            _set_cell(tbl.cell(row_idx, 1), item.get('name', ''), 8)

            dates_str = ''
            d_start = item.get('start', '')
            d_finish = item.get('finish', '')
            if d_start and d_finish:
                dates_str = f"{d_start} \u2013 {d_finish}"
            elif d_start:
                dates_str = str(d_start)
            _set_cell(tbl.cell(row_idx, 2), dates_str, 8, False, None, PP_ALIGN.CENTER)

            status = item.get('status', '')
            status_cell = tbl.cell(row_idx, 3)
            _set_cell(status_cell, status, 8, False, None, PP_ALIGN.CENTER)
            if status == 'Complete':
                status_cell.fill.solid()
                status_cell.fill.fore_color.rgb = QA_GREEN
                for para in status_cell.text_frame.paragraphs:
                    para.font.color.rgb = WHITE
            elif status == 'In Progress':
                status_cell.fill.solid()
                status_cell.fill.fore_color.rgb = RGBColor(0xFF, 0xC0, 0x00)

            # Role columns
            merged_roles = item.get('roles', {})
            role_letters = []
            for pi, shortname in enumerate(people):
                col_idx = fixed_count + pi
                role = merged_roles.get(shortname, '')
                cell = tbl.cell(row_idx, col_idx)
                if role:
                    _set_cell(cell, role, 8, True, WHITE, PP_ALIGN.CENTER)
                    if role in ROLE_COLOURS:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = ROLE_COLOURS[role]
                    role_letters.append(role)

            # QA column
            has_p = 'P' in role_letters
            has_r = 'R' in role_letters
            has_a = 'A' in role_letters
            qa_cell = tbl.cell(row_idx, total_cols - 1)
            if has_p and has_r and has_a:
                _set_cell(qa_cell, '\u2713', 8, True, WHITE, PP_ALIGN.CENTER)
                qa_cell.fill.solid()
                qa_cell.fill.fore_color.rgb = QA_GREEN
            else:
                _set_cell(qa_cell, '', 8)

            # Zebra striping
            if row_idx % 2 == 0:
                for ci in range(total_cols):
                    c = tbl.cell(row_idx, ci)
                    # Only shade cells that don't already have a role fill
                    if ci < fixed_count or (ci == total_cols - 1 and not (has_p and has_r and has_a)):
                        if not (ci == 3 and status in ('Complete', 'In Progress')):
                            c.fill.solid()
                            c.fill.fore_color.rgb = LIGHT_GREY

        # Footer
        fb = slide.shapes.add_textbox(Inches(0.4), Inches(7.1),
                                      Inches(4), Inches(0.25))
        ff = fb.text_frame
        fp = ff.paragraphs[0]
        fp.text = _sanitise(f"Generated by Noodle Planner  |  Appendix")
        fp.font.size = Pt(7)
        fp.font.color.rgb = RGBColor(160, 160, 160)


def export_portfolio_to_powerpoint(output_path, portfolio_data, project_reports):
    """Export a portfolio report as a multi-slide PowerPoint deck."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_portfolio_overview_slide(prs, portfolio_data)

    portfolio_risks = _collect_portfolio_risks(project_reports)
    _add_portfolio_risk_slides(prs, portfolio_data, portfolio_risks)

    for report_data in project_reports:
        _add_report_slide(prs, report_data, include_footer=False)

    # Appendix: deliverables matrix slides
    _add_portfolio_deliverables_slides(prs, project_reports)

    prs.save(output_path)
    logger.info(f"Exported portfolio report to PowerPoint: {output_path}")


def export_to_excel(text, output_path, is_yaml=True, project_name="Project", original_text=None):
    """Export project data to Excel format."""
    # Import here to avoid circular imports
    from .scheduling_engine import natural_language_to_yaml, schedule_tasks

    resource_map = {}
    baseline_lookup = {}
    if original_text:
        resource_map, _ = parse_resource_mappings(original_text)
        baseline_text = extract_baseline(original_text)
        if baseline_text:
            baseline_items_list = parse_baseline_markdown(baseline_text)
            baseline_lookup = {item['name']: item for item in baseline_items_list}

    if is_yaml:
        data = yaml.safe_load(text)
        project_name = list(data.keys())[0]
        phases_raw = data[project_name]
    else:
        data = natural_language_to_yaml(text, project_name)
        phases_raw = data[project_name]

    if isinstance(phases_raw, list):
        phases = phases_raw
    elif isinstance(phases_raw, dict):
        phases = [phases_raw]
    else:
        phases = []

    tasks = schedule_tasks(phases)

    wb = Workbook()

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")

    ws_tasks = wb.active
    ws_tasks.title = "Tasks"

    task_headers = ['ID', 'Task Name', 'Start', 'Finish', 'Duration (days)',
                    'Resources', '% Complete', 'RAG', 'Priority', 'Bucket',
                    'Dependencies', 'Comment']
    ws_tasks.append(task_headers)

    for col_num, header in enumerate(task_headers, 1):
        cell = ws_tasks.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    name_to_id = {}
    for _idx, _task in enumerate(tasks, start=1):
        _name = _task.get('name', '')
        if _name:
            name_to_id[_name.lower()] = _idx

    task_row_num = 2
    for idx, task in enumerate(tasks, start=1):
        duration = task.get('duration')
        if isinstance(duration, timedelta):
            duration_days = duration.days
        elif task.get('start') and task.get('finish'):
            duration_days = (task.get('finish') - task.get('start')).days
        else:
            duration_days = 0

        task_name = task.get('description', task.get('name', ''))
        task_name = task_name.replace('_', ' ')

        level = task.get('level', 0)
        if level > 0:
            indent = '  ' * level
            task_name = f"{indent}{task_name}"

        resources_str = task.get('resources', '')
        if resources_str and resource_map:
            resource_list = [r.strip() for r in resources_str.split(',')]
            mapped_resources = [resource_map.get(r.lower(), r) for r in resource_list]
            resources_str = ', '.join(mapped_resources)

        if task.get('summary'):
            rag_status = ''
        else:
            rag_status = calculate_rag_status(task)

        deps_str = ''
        if task.get('depends'):
            dep_type_map = task.get('dependency_types', {})
            lag_lead_map = task.get('lag_lead', {})
            dep_parts = []
            for dep_name in task['depends']:
                dep_id = name_to_id.get(dep_name.lower())
                if dep_id is None:
                    continue
                dep_type = dep_type_map.get(dep_name, 'FS')
                entry = f"{dep_id}{dep_type}"
                lag = lag_lead_map.get(dep_name, '')
                if lag:
                    entry += lag
                dep_parts.append(entry)
            deps_str = ', '.join(dep_parts)

        row = [
            idx,
            task_name,
            task.get('start').strftime('%Y-%m-%d') if task.get('start') else '',
            task.get('finish').strftime('%Y-%m-%d') if task.get('finish') else '',
            duration_days,
            resources_str,
            task.get('percent', 0) if task.get('percent') else '',
            rag_status,
            task.get('priority', 'Low'),
            task.get('bucket', ''),
            deps_str,
            task.get('comment', '')
        ]
        ws_tasks.append(row)

        if task.get('summary'):
            for col_num in range(1, len(task_headers) + 1):
                cell = ws_tasks.cell(row=task_row_num, column=col_num)
                cell.font = Font(bold=True)

        if rag_status:
            rag_col_num = 8
            rag_cell = ws_tasks.cell(row=task_row_num, column=rag_col_num)
            colour = rag_status_to_colour(rag_status)
            if colour == 'green':
                rag_cell.fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
            elif colour == 'amber':
                rag_cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
            elif colour == 'red':
                rag_cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
                rag_cell.font = Font(color="FFFFFF", bold=True)
            elif colour == 'blue':
                rag_cell.fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
                rag_cell.font = Font(color="FFFFFF", bold=True)
            elif colour == 'grey':
                rag_cell.fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
                rag_cell.font = Font(color="FFFFFF", bold=True)

        task_row_num += 1

    for col_num, header in enumerate(task_headers, 1):
        column_letter = get_column_letter(col_num)
        max_length = len(header)
        for row in ws_tasks.iter_rows(min_row=2, max_col=col_num, max_row=ws_tasks.max_row):
            cell_value = str(row[col_num-1].value) if row[col_num-1].value else ''
            max_length = max(max_length, len(cell_value))
        ws_tasks.column_dimensions[column_letter].width = min(max_length + 2, 50)

    ws_milestones = wb.create_sheet("Milestones")
    has_baseline = len(baseline_lookup) > 0
    milestone_headers = ['Milestone', 'Type', 'Date']
    if has_baseline:
        milestone_headers.extend(['Baseline Finish', 'Variance (days)'])
    ws_milestones.append(milestone_headers)

    for col_num, header in enumerate(milestone_headers, 1):
        cell = ws_milestones.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    phase_dates = {}
    for t in tasks:
        phase = t.get('phase')
        if phase:
            if phase not in phase_dates:
                phase_dates[phase] = {'start': t.get('start'), 'end': t.get('finish')}
            else:
                if t.get('start') and (not phase_dates[phase]['start'] or t.get('start') < phase_dates[phase]['start']):
                    phase_dates[phase]['start'] = t.get('start')
                if t.get('finish') and (not phase_dates[phase]['end'] or t.get('finish') > phase_dates[phase]['end']):
                    phase_dates[phase]['end'] = t.get('finish')

    milestone_entries = []
    for phase, dates in phase_dates.items():
        milestone_entries.append({
            'name': phase,
            'type': 'Phase',
            'date': dates['end']
        })

    for t in tasks:
        if t.get('summary'):
            milestone_entries.append({
                'name': t.get('description', t.get('name', '')),
                'type': 'Summary',
                'date': t.get('finish')
            })

    for t in tasks:
        duration = t.get('duration', timedelta(days=1))
        if isinstance(duration, timedelta) and duration.days == 0:
            milestone_entries.append({
                'name': t.get('description', t.get('name', '')),
                'type': 'Milestone',
                'date': t.get('start')
            })

    seen = set()
    unique_milestones = []
    for entry in milestone_entries:
        key = (entry['name'], entry['date'])
        if key not in seen:
            seen.add(key)
            unique_milestones.append(entry)

    unique_milestones.sort(key=lambda x: x['date'] if x['date'] else datetime.now())

    for milestone in unique_milestones:
        row = [
            milestone['name'],
            milestone['type'],
            milestone['date'].strftime('%Y-%m-%d') if milestone['date'] else ''
        ]
        if has_baseline:
            bl = baseline_lookup.get(milestone['name'])
            bl_finish = bl.get('finish', '') if bl else ''
            row.append(bl_finish)
            if bl_finish and milestone['date']:
                try:
                    from datetime import date as date_type
                    bl_date = datetime.strptime(bl_finish, '%Y-%m-%d').date() if isinstance(bl_finish, str) else bl_finish
                    current_date = milestone['date'].date() if isinstance(milestone['date'], datetime) else milestone['date']
                    diff = (current_date - bl_date).days
                    if diff > 0:
                        row.append(f'+{diff}')
                    elif diff < 0:
                        row.append(str(diff))
                    else:
                        row.append('0')
                except (ValueError, TypeError):
                    row.append('')
            else:
                row.append('New' if not bl else '')
        ws_milestones.append(row)

    for col_num, header in enumerate(milestone_headers, 1):
        column_letter = get_column_letter(col_num)
        max_length = len(header)
        for row in ws_milestones.iter_rows(min_row=2, max_col=col_num, max_row=ws_milestones.max_row):
            cell_value = str(row[col_num-1].value) if row[col_num-1].value else ''
            max_length = max(max_length, len(cell_value))
        ws_milestones.column_dimensions[column_letter].width = min(max_length + 2, 50)

    if tasks:
        tasks_with_dates = [t for t in tasks if t.get('start') and t.get('finish') and t.get('resources') and not t.get('summary')]

        if tasks_with_dates:
            start_date = min([t.get('start') for t in tasks_with_dates])
            finish_date = max([t.get('finish') for t in tasks_with_dates])

            ws_resources = wb.create_sheet("Resources")

            date_range = []
            current = start_date
            while current <= finish_date:
                date_range.append(current)
                current += timedelta(days=1)

            resource_headers = ['Resource']
            for date in date_range:
                day_of_week = date.strftime('%a')
                day = date.strftime('%d')
                month = date.strftime('%b').lower()
                resource_headers.append(f"{day_of_week} {day} {month}")

            ws_resources.append(resource_headers)

            for col_num, header in enumerate(resource_headers, 1):
                cell = ws_resources.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment

                if col_num > 1:
                    date = date_range[col_num - 2]
                    if date.weekday() in [5, 6]:
                        cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
                        cell.font = Font(color="000000", bold=True)

            resource_daily_hours = {}

            for task in tasks_with_dates:
                resources_str = task.get('resources', '')
                if resources_str:
                    task_resources = [
                        resource_map.get(resource.lower(), resource)
                        for resource in (r.strip() for r in resources_str.split(','))
                        if resource
                    ]
                    task_start = task.get('start')
                    task_finish = task.get('finish')

                    task_duration_days = (task_finish - task_start).days + 1

                    duration = task.get('duration', timedelta(days=0))
                    if isinstance(duration, timedelta):
                        effort_days = duration.days
                    else:
                        effort_days = 0

                    if task_duration_days > 0 and len(task_resources) > 0:
                        hours_per_day = (effort_days * 8.0) / task_duration_days / len(task_resources)

                        current = task_start
                        while current <= task_finish:
                            date_key = current.strftime('%Y-%m-%d')

                            for resource in task_resources:
                                resource_key = resource.lower()
                                if resource_key not in resource_daily_hours:
                                    resource_daily_hours[resource_key] = {'display_name': resource, 'hours': {}}

                                if date_key not in resource_daily_hours[resource_key]['hours']:
                                    resource_daily_hours[resource_key]['hours'][date_key] = 0

                                resource_daily_hours[resource_key]['hours'][date_key] += hours_per_day

                            current += timedelta(days=1)

            for resource_key in sorted(resource_daily_hours.keys()):
                display_name = resource_daily_hours[resource_key]['display_name']
                hours_dict = resource_daily_hours[resource_key]['hours']

                row = [display_name]
                for date in date_range:
                    date_key = date.strftime('%Y-%m-%d')
                    hours = hours_dict.get(date_key, 0)
                    if hours > 0:
                        row.append(round(hours, 1))
                    else:
                        row.append('-')

                ws_resources.append(row)

                row_num = ws_resources.max_row
                for col_num in range(2, len(resource_headers) + 1):
                    cell = ws_resources.cell(row=row_num, column=col_num)
                    cell.alignment = Alignment(horizontal="center")

                    date = date_range[col_num - 2]
                    is_weekend = date.weekday() in [5, 6]

                    if cell.value == '-':
                        if is_weekend:
                            cell.fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
                    else:
                        hours = float(cell.value)
                        if hours <= 4:
                            cell.fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
                            cell.font = Font(color="2E7D32")
                        elif hours <= 8:
                            cell.fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
                            cell.font = Font(color="E65100")
                        else:
                            cell.fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
                            cell.font = Font(color="C62828", bold=True)

                        if is_weekend:
                            current_color = cell.font.color
                            cell.fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

            for col_num, header in enumerate(resource_headers, 1):
                column_letter = get_column_letter(col_num)
                max_length = len(header)
                for row in ws_resources.iter_rows(min_row=2, max_col=col_num, max_row=ws_resources.max_row):
                    cell_value = str(row[col_num-1].value) if row[col_num-1].value else ''
                    max_length = max(max_length, len(cell_value))
                ws_resources.column_dimensions[column_letter].width = min(max_length + 2, 50)

    if original_text:
        budget_text = extract_budget(original_text)
        if budget_text:
            budget_items = parse_budget_markdown(budget_text)
            if budget_items:
                ws_budget = wb.create_sheet("Budget")

                budget_headers = [
                    'ID', 'Description', 'Estimate', 'Forecast', 'Type',
                    'Invoice', 'PO', 'Supplier', 'Total', 'Ordered',
                    'Received', 'Category'
                ]

                ws_budget.append(budget_headers)

                budget_header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
                budget_header_font = Font(bold=True, color="FFFFFF", size=11)

                for col_num, header in enumerate(budget_headers, 1):
                    cell = ws_budget.cell(row=1, column=col_num)
                    cell.fill = budget_header_fill
                    cell.font = budget_header_font
                    cell.alignment = Alignment(horizontal='center')

                currency_cols = {3, 4, 9}

                for row_idx, item in enumerate(budget_items, 2):
                    ws_budget.cell(row=row_idx, column=1, value=item.get('id', ''))
                    ws_budget.cell(row=row_idx, column=2, value=item.get('description', ''))

                    for col_num in currency_cols:
                        field = budget_headers[col_num - 1].lower()
                        val = item.get(field, 0)
                        cell = ws_budget.cell(row=row_idx, column=col_num, value=val)
                        cell.number_format = '#,##0.00'

                    ws_budget.cell(row=row_idx, column=5, value=item.get('type', ''))
                    ws_budget.cell(row=row_idx, column=6, value=item.get('invoice', ''))
                    ws_budget.cell(row=row_idx, column=7, value=item.get('po', ''))
                    ws_budget.cell(row=row_idx, column=8, value=item.get('supplier', ''))
                    ws_budget.cell(row=row_idx, column=10, value=item.get('date_ordered', ''))
                    ws_budget.cell(row=row_idx, column=11, value=item.get('date_received', ''))
                    ws_budget.cell(row=row_idx, column=12, value=item.get('category', ''))

                last_row = len(budget_items) + 2
                total_estimate = sum(item.get('estimate', 0) for item in budget_items)
                total_forecast = sum(item.get('forecast', 0) for item in budget_items)
                total_spend = sum(item.get('total', 0) for item in budget_items)

                ws_budget.cell(row=last_row, column=2, value='TOTALS').font = Font(bold=True)
                for col_num, val in [(3, total_estimate), (4, total_forecast), (9, total_spend)]:
                    cell = ws_budget.cell(row=last_row, column=col_num, value=val)
                    cell.font = Font(bold=True)
                    cell.number_format = '#,##0.00'

                budget_col_widths = [6, 30, 12, 12, 10, 15, 12, 20, 12, 12, 12, 18]
                for col_num, width in enumerate(budget_col_widths, 1):
                    ws_budget.column_dimensions[get_column_letter(col_num)].width = width

    if original_text:
        raid_log_text = extract_raid_log(original_text)
        if raid_log_text:
            raid_items = parse_raid_markdown(raid_log_text)
            if raid_items:
                ws_raid = wb.create_sheet("RAID Log")

                raid_headers = [
                    'ID', 'Type', 'Title', 'Description', 'Raised By', 'Owner',
                    'Mitigation Actions', 'Impact', 'Likelihood', 'Score', 'Status',
                    'Priority', 'Target Date'
                ]

                ws_raid.append(raid_headers)

                raid_header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
                raid_header_font = Font(bold=True, color="FFFFFF", size=11)

                for col_num, header in enumerate(raid_headers, 1):
                    cell = ws_raid.cell(row=1, column=col_num)
                    cell.fill = raid_header_fill
                    cell.font = raid_header_font
                    cell.alignment = Alignment(horizontal='center')

                for row_idx, item in enumerate(raid_items, 2):
                    ws_raid.cell(row=row_idx, column=1, value=item.get('id', ''))
                    ws_raid.cell(row=row_idx, column=2, value=item.get('type', '').capitalize())
                    ws_raid.cell(row=row_idx, column=3, value=item.get('title', ''))
                    ws_raid.cell(row=row_idx, column=4, value=item.get('description', ''))
                    ws_raid.cell(row=row_idx, column=5, value=item.get('raised_by', ''))
                    ws_raid.cell(row=row_idx, column=6, value=item.get('owner', ''))
                    ws_raid.cell(row=row_idx, column=7, value=item.get('mitigation_actions', ''))
                    ws_raid.cell(row=row_idx, column=8, value=item.get('impact', ''))
                    ws_raid.cell(row=row_idx, column=9, value=item.get('likelihood', ''))

                    score = item.get('score', 0)
                    score_cell = ws_raid.cell(row=row_idx, column=10, value=score)
                    if score >= 16:
                        score_cell.fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
                    elif score >= 6:
                        score_cell.fill = PatternFill(start_color="FFF3BF", end_color="FFF3BF", fill_type="solid")
                    else:
                        score_cell.fill = PatternFill(start_color="D3F9D8", end_color="D3F9D8", fill_type="solid")

                    ws_raid.cell(row=row_idx, column=11, value=item.get('status', '').capitalize())
                    ws_raid.cell(row=row_idx, column=12, value=item.get('priority', ''))
                    ws_raid.cell(row=row_idx, column=13, value=item.get('target_date', item.get('date', '')))

                for col_num, header in enumerate(raid_headers, 1):
                    width = RAID_COLUMN_WIDTHS.get(header, 15)
                    ws_raid.column_dimensions[get_column_letter(col_num)].width = width

    if original_text:
        stakeholder_items = parse_stakeholders_from_frontmatter(original_text)
        if stakeholder_items:
            ws_stakeholders = wb.create_sheet("Stakeholders")

            stakeholder_headers = ['Name', 'Role', 'Interest', 'Influence']
            ws_stakeholders.append(stakeholder_headers)

            stakeholder_header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
            stakeholder_header_font = Font(bold=True, color="FFFFFF", size=11)

            for col_num, header_text in enumerate(stakeholder_headers, 1):
                cell = ws_stakeholders.cell(row=1, column=col_num)
                cell.fill = stakeholder_header_fill
                cell.font = stakeholder_header_font
                cell.alignment = Alignment(horizontal='center')

            for row_idx, item in enumerate(stakeholder_items, 2):
                ws_stakeholders.cell(row=row_idx, column=1, value=item.get('name', ''))
                ws_stakeholders.cell(row=row_idx, column=2, value=item.get('role', ''))

                interest_val = item.get('interest', 'low').capitalize()
                interest_cell = ws_stakeholders.cell(row=row_idx, column=3, value=interest_val)
                if interest_val.lower() == 'high':
                    interest_cell.fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
                else:
                    interest_cell.fill = PatternFill(start_color="D3F9D8", end_color="D3F9D8", fill_type="solid")

                influence_val = item.get('influence', 'low').capitalize()
                influence_cell = ws_stakeholders.cell(row=row_idx, column=4, value=influence_val)
                if influence_val.lower() == 'high':
                    influence_cell.fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
                else:
                    influence_cell.fill = PatternFill(start_color="D3F9D8", end_color="D3F9D8", fill_type="solid")

            stakeholder_col_widths = [25, 30, 12, 12]
            for col_num, width in enumerate(stakeholder_col_widths, 1):
                ws_stakeholders.column_dimensions[get_column_letter(col_num)].width = width

    if tasks:
        budget_items_for_evm = None
        if original_text:
            budget_text = extract_budget(original_text)
            if budget_text:
                budget_items_for_evm = parse_budget_markdown(budget_text)

        evm_data = calculate_evm(tasks, budget_items_for_evm)
        if evm_data:
            ws_evm = wb.create_sheet("EVM")

            evm_header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
            evm_header_font = Font(bold=True, color="FFFFFF", size=11)

            evm_headers = ['Metric', 'Value']
            ws_evm.append(evm_headers)

            for col_num, header_text in enumerate(evm_headers, 1):
                cell = ws_evm.cell(row=1, column=col_num)
                cell.fill = evm_header_fill
                cell.font = evm_header_font
                cell.alignment = Alignment(horizontal='center')

            fmt = lambda v: round(v, 2)
            fmt_pct = lambda v: f"{round(v, 1)}%"
            fmt_idx = lambda v: round(v, 2)

            evm_rows = [
                ('% Complete', fmt_pct(evm_data['overall_percent_complete'])),
                ('BAC (Budget at Completion)', fmt(evm_data['BAC'])),
                ('PV (Planned Value)', fmt(evm_data['PV'])),
                ('EV (Earned Value)', fmt(evm_data['EV'])),
                ('AC (Actual Cost)', fmt(evm_data['AC'])),
                ('CV (Cost Variance)', fmt(evm_data['CV'])),
                ('SV (Schedule Variance)', fmt(evm_data['SV'])),
                ('CPI (Cost Performance Index)', fmt_idx(evm_data['CPI'])),
                ('SPI (Schedule Performance Index)', fmt_idx(evm_data['SPI'])),
                ('EAC (Estimate at Completion)', fmt(evm_data['EAC'])),
                ('ETC (Estimate to Complete)', fmt(evm_data['ETC'])),
                ('VAC (Variance at Completion)', fmt(evm_data['VAC'])),
                ('Project Start', evm_data['project_start'].strftime('%Y-%m-%d') if evm_data.get('project_start') else ''),
                ('Project End', evm_data['project_end'].strftime('%Y-%m-%d') if evm_data.get('project_end') else ''),
                ('Time Elapsed', fmt_pct(evm_data['time_elapsed_fraction'] * 100)),
                ('Budget Data Available', 'Yes' if evm_data.get('has_budget_data') else 'No (using duration proxy)'),
            ]

            for row_data in evm_rows:
                ws_evm.append(list(row_data))

            for row_idx in range(2, len(evm_rows) + 2):
                metric_cell = ws_evm.cell(row=row_idx, column=1)
                value_cell = ws_evm.cell(row=row_idx, column=2)
                metric_name = metric_cell.value or ''

                if metric_name.startswith('CV') or metric_name.startswith('SV') or metric_name.startswith('VAC'):
                    try:
                        val = float(value_cell.value)
                        if val >= 0:
                            value_cell.fill = PatternFill(start_color="D3F9D8", end_color="D3F9D8", fill_type="solid")
                        else:
                            value_cell.fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
                    except (ValueError, TypeError):
                        pass
                elif metric_name.startswith('CPI') or metric_name.startswith('SPI'):
                    try:
                        val = float(value_cell.value)
                        if val >= 1:
                            value_cell.fill = PatternFill(start_color="D3F9D8", end_color="D3F9D8", fill_type="solid")
                        else:
                            value_cell.fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
                    except (ValueError, TypeError):
                        pass

            ws_evm.column_dimensions['A'].width = 35
            ws_evm.column_dimensions['B'].width = 25

    if original_text:
        try:
            from .format_converter import extract_comms_plan, parse_comms_markdown
            comms_text = extract_comms_plan(original_text)
            if comms_text:
                comms_items = parse_comms_markdown(comms_text)
                if comms_items:
                    ws_comms = wb.create_sheet("Comms Plan")
                    comms_headers = ['ID', 'Activity', 'Audience', 'Content', 'Frequency', 'Channel', 'Owner', 'Status']
                    comms_header_fill = PatternFill(start_color="4A90D9", end_color="4A90D9", fill_type="solid")
                    comms_header_font = Font(bold=True, color="FFFFFF", size=11)

                    for col, header in enumerate(comms_headers, 1):
                        cell = ws_comms.cell(row=1, column=col, value=header)
                        cell.fill = comms_header_fill
                        cell.font = comms_header_font
                        cell.alignment = Alignment(horizontal='center')

                    for row_idx, item in enumerate(comms_items, 2):
                        ws_comms.cell(row=row_idx, column=1, value=item.get('id', row_idx - 1))
                        ws_comms.cell(row=row_idx, column=2, value=item.get('activity', ''))
                        ws_comms.cell(row=row_idx, column=3, value=item.get('audience', ''))
                        ws_comms.cell(row=row_idx, column=4, value=item.get('content', ''))
                        ws_comms.cell(row=row_idx, column=5, value=item.get('frequency', ''))
                        ws_comms.cell(row=row_idx, column=6, value=item.get('channel', ''))
                        ws_comms.cell(row=row_idx, column=7, value=item.get('owner', ''))
                        ws_comms.cell(row=row_idx, column=8, value=item.get('status', ''))

                    comms_col_widths = {'A': 5, 'B': 25, 'C': 20, 'D': 30, 'E': 12, 'F': 12, 'G': 18, 'H': 10}
                    for col_letter, width in comms_col_widths.items():
                        ws_comms.column_dimensions[col_letter].width = width
        except Exception as e:
            logger.warning(f"Failed to add comms plan worksheet: {e}")

    if original_text:
        try:
            from .format_converter import extract_lessons, parse_lessons_markdown
            lessons_text = extract_lessons(original_text)
            if lessons_text:
                lessons_items = parse_lessons_markdown(lessons_text)
                if lessons_items:
                    ws_lessons = wb.create_sheet("Lessons Learned")
                    lessons_headers = [
                        'ID', 'Project Manager', 'Project Type', 'Technology',
                        'Project Phase', 'Area', 'Impact Type', 'Observation',
                        'Impact', 'Recommendations', 'Date',
                    ]
                    lessons_header_fill = PatternFill(
                        start_color="6F42C1", end_color="6F42C1", fill_type="solid"
                    )
                    lessons_header_font = Font(bold=True, color="FFFFFF", size=11)

                    for col, header in enumerate(lessons_headers, 1):
                        cell = ws_lessons.cell(row=1, column=col, value=header)
                        cell.fill = lessons_header_fill
                        cell.font = lessons_header_font
                        cell.alignment = Alignment(horizontal='center')

                    for row_idx, item in enumerate(lessons_items, 2):
                        ws_lessons.cell(row=row_idx, column=1, value=item.get('id', row_idx - 1))
                        ws_lessons.cell(row=row_idx, column=2, value=item.get('project_manager', ''))
                        ws_lessons.cell(row=row_idx, column=3, value=item.get('project_type', ''))
                        ws_lessons.cell(row=row_idx, column=4, value=item.get('technology', ''))
                        ws_lessons.cell(row=row_idx, column=5, value=item.get('project_phase', ''))
                        ws_lessons.cell(row=row_idx, column=6, value=item.get('area', ''))
                        ws_lessons.cell(row=row_idx, column=7, value=item.get('impact_type', ''))
                        ws_lessons.cell(row=row_idx, column=8, value=item.get('observation', ''))
                        ws_lessons.cell(row=row_idx, column=9, value=item.get('impact', ''))
                        ws_lessons.cell(row=row_idx, column=10, value=item.get('recommendations', ''))
                        ws_lessons.cell(row=row_idx, column=11, value=item.get('date', ''))

                    lessons_col_widths = {
                        'A': 5, 'B': 18, 'C': 16, 'D': 16, 'E': 14,
                        'F': 16, 'G': 16, 'H': 40, 'I': 30, 'J': 30, 'K': 12,
                    }
                    for col_letter, width in lessons_col_widths.items():
                        ws_lessons.column_dimensions[col_letter].width = width
        except Exception as e:
            logger.warning(f"Failed to add lessons learned worksheet: {e}")

    try:
        deliverables = [t for t in tasks
                        if t.get('deliverable') and t.get('product_type', 'internal') != 'group']

        if deliverables:
            people = {}
            for t in tasks:
                qr = t.get('quality_roles', {})
                if qr and isinstance(qr, dict):
                    for name in qr:
                        key = name.lower()
                        if key not in people:
                            people[key] = resource_map.get(key, name)
                res_str = t.get('resources', '')
                if res_str:
                    for r in res_str.split(','):
                        key = r.strip().lower()
                        if key and key not in people:
                            people[key] = resource_map.get(key, r.strip())
            for short, full in resource_map.items():
                if short.lower() not in people:
                    people[short.lower()] = full

            people_list = sorted(people.items(), key=lambda x: x[0])

            # Build a role lookup: shortname -> role title (from Resources and Stakeholders)
            # Strip out email addresses and full names — headers should show only the role.
            def _clean_role(role_str):
                """Remove email addresses and return only the role title."""
                if not role_str:
                    return ''
                parts = [p.strip() for p in role_str.split(',')]
                cleaned = [p for p in parts if p and not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', p)]
                return ', '.join(cleaned)

            dm_role_map = {}
            if original_text:
                dm_role_map = parse_resource_roles(original_text)
                for s in parse_stakeholders_from_frontmatter(original_text):
                    skey = (s.get('shortname') or s.get('name', '')).lower()
                    if skey and s.get('role') and skey not in dm_role_map:
                        dm_role_map[skey] = s['role']
                # Clean email addresses from all role values; drop empty entries
                dm_role_map = {k: c for k, v in dm_role_map.items()
                               if (c := _clean_role(v))}

            ws_dm = wb.create_sheet("Deliverables Matrix")

            fixed_headers = ['ID', 'Deliverable', 'Dates', 'Status']
            # Use role as the column header; fall back to shortname if no role is known
            person_headers = [dm_role_map.get(short, short) for short, _ in people_list]
            dm_headers = fixed_headers + person_headers + ['QA']

            dm_header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            dm_header_font = Font(bold=True, color="FFFFFF", size=10)
            ws_dm.row_dimensions[1].height = 80

            for col, header in enumerate(dm_headers, 1):
                cell = ws_dm.cell(row=1, column=col, value=header)
                cell.fill = dm_header_fill
                cell.font = dm_header_font
                if col > len(fixed_headers) and col <= len(fixed_headers) + len(person_headers):
                    cell.alignment = Alignment(text_rotation=90, horizontal='center', vertical='bottom')
                else:
                    cell.alignment = Alignment(horizontal='center', vertical='center')

            role_fills = {
                'P': PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid"),
                'R': PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid"),
                'A': PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid"),
            }
            role_font = Font(bold=True, color="FFFFFF")

            for d_idx, dtask in enumerate(deliverables, start=1):
                row_num = d_idx + 1

                merged_roles = {}
                qr = dtask.get('quality_roles', {})
                if qr and isinstance(qr, dict):
                    for name, role in qr.items():
                        merged_roles[name.lower()] = role
                res_str = dtask.get('resources', '')
                if res_str:
                    for r in res_str.split(','):
                        key = r.strip().lower()
                        if key and key not in merged_roles:
                            merged_roles[key] = 'P'
                dtask_name = (dtask.get('name') or '').lower()
                for t in tasks:
                    if (t.get('parent') or '').lower() == dtask_name:
                        child_qr = t.get('quality_roles', {})
                        if child_qr and isinstance(child_qr, dict):
                            for name, role in child_qr.items():
                                if name.lower() not in merged_roles:
                                    merged_roles[name.lower()] = role
                        child_res = t.get('resources', '')
                        if child_res:
                            for r in child_res.split(','):
                                key = r.strip().lower()
                                if key and key not in merged_roles:
                                    merged_roles[key] = 'P'

                ws_dm.cell(row=row_num, column=1, value=d_idx)
                deliverable_name = dtask.get('deliverable', dtask.get('name', ''))
                ws_dm.cell(row=row_num, column=2, value=deliverable_name.replace('_', ' '))

                # Dates column — show start – finish range
                d_start = dtask.get('start')
                d_finish = dtask.get('finish')
                dates_str = ''
                if d_start and d_finish:
                    fmt = '%Y-%m-%d'
                    s = d_start.strftime(fmt) if hasattr(d_start, 'strftime') else str(d_start)
                    f = d_finish.strftime(fmt) if hasattr(d_finish, 'strftime') else str(d_finish)
                    dates_str = f"{s} \u2013 {f}"
                elif d_start:
                    dates_str = str(d_start.strftime('%Y-%m-%d') if hasattr(d_start, 'strftime') else d_start)
                ws_dm.cell(row=row_num, column=3, value=dates_str)

                status = 'Complete' if dtask.get('percent', 0) == 100 else ('In Progress' if dtask.get('percent', 0) else '')
                status_cell = ws_dm.cell(row=row_num, column=4, value=status)
                if status == 'Complete':
                    status_cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
                    status_cell.font = Font(color="FFFFFF")
                elif status == 'In Progress':
                    status_cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")

                role_letters = []
                for p_idx, (shortname, _) in enumerate(people_list):
                    col_num = len(fixed_headers) + p_idx + 1
                    role = merged_roles.get(shortname)
                    if role:
                        cell = ws_dm.cell(row=row_num, column=col_num, value=role)
                        if role in role_fills:
                            cell.fill = role_fills[role]
                            cell.font = role_font
                        cell.alignment = Alignment(horizontal='center')
                        role_letters.append(role)

                has_p = 'P' in role_letters
                has_r = 'R' in role_letters
                has_a = 'A' in role_letters
                qa_col = len(dm_headers)
                qa_cell = ws_dm.cell(row=row_num, column=qa_col,
                                     value='\u2713' if (has_p and has_r and has_a) else '')
                if has_p and has_r and has_a:
                    qa_cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
                    qa_cell.font = Font(color="FFFFFF", bold=True)
                qa_cell.alignment = Alignment(horizontal='center')

            ws_dm.column_dimensions['A'].width = 5
            ws_dm.column_dimensions['B'].width = 25
            ws_dm.column_dimensions['C'].width = 24
            ws_dm.column_dimensions['D'].width = 12
            for p_idx in range(len(people_list)):
                col_letter = get_column_letter(len(fixed_headers) + p_idx + 1)
                ws_dm.column_dimensions[col_letter].width = 4
            if dm_headers:
                qa_letter = get_column_letter(len(dm_headers))
                ws_dm.column_dimensions[qa_letter].width = 5
    except Exception as e:
        logger.warning(f"Failed to add Deliverables Matrix worksheet: {e}")

    if original_text:
        try:
            from .front_matter_parser import FrontMatterParser
            fm = FrontMatterParser(original_text)
            kv = fm.parse_key_values()
            title = fm.parse_title() or kv.get('title', project_name)

            ws_summary = wb.create_sheet("Summary", 0)
            summary_label_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            summary_label_font = Font(bold=True, color="FFFFFF", size=11)
            summary_value_font = Font(size=11)

            # Calculate overall project completion from leaf (non-summary) tasks.
            # Matches the frontend's calculateProjectCompletionFromTasks (portfolio-status.js):
            # simple average of leaf task percent values.
            leaf_tasks = [t for t in tasks if not t.get('summary')]
            if leaf_tasks:
                total_percent = sum(float(t.get('percent') or 0) for t in leaf_tasks)
                percent_complete = round(total_percent / len(leaf_tasks))
            else:
                percent_complete = 0

            summary_rows = [
                ('Project Name', title),
                ('Version', kv.get('version', '')),
                ('Start Date', kv.get('start date', kv.get('start', ''))),
                ('Project Manager', kv.get('project manager', kv.get('pm', ''))),
                ('Budget', kv.get('budget', '')),
                ('Sponsor', kv.get('sponsor', '')),
                ('Percentage Complete', percent_complete / 100.0),
                ('Date Exported', datetime.now().strftime('%Y-%m-%d %H:%M')),
            ]

            for row_idx, (label, value) in enumerate(summary_rows, 1):
                label_cell = ws_summary.cell(row=row_idx, column=1, value=label)
                label_cell.fill = summary_label_fill
                label_cell.font = summary_label_font
                label_cell.alignment = Alignment(horizontal='right')
                value_cell = ws_summary.cell(row=row_idx, column=2, value=value)
                value_cell.font = summary_value_font
                if label == 'Percentage Complete':
                    value_cell.number_format = '0%'

            ws_summary.column_dimensions['A'].width = 20
            ws_summary.column_dimensions['B'].width = 35
        except Exception as e:
            logger.warning(f"Failed to add Summary worksheet: {e}")

    wb.save(output_path)
    logger.info(f"Exported project data to {output_path}")


def text_to_markdown_table(text, is_yaml=True, project_name="Project", terminal_width=80, original_text=None):
    """Convert text (YAML or natural language) to markdown table."""
    # Import here to avoid circular imports
    from .scheduling_engine import natural_language_to_yaml, schedule_tasks

    today = datetime.today()

    resource_map = {}
    resource_nwd = {}
    project_holidays = set()
    if original_text:
        resource_map, resource_nwd = parse_resource_mappings(original_text)
        from .front_matter_parser import FrontMatterParser
        _fm_parser = FrontMatterParser(original_text)
        project_holidays = _fm_parser.parse_non_working_days()

    if is_yaml:
        data = yaml.safe_load(text)
        project_name = list(data.keys())[0]
        phases_raw = data[project_name]
    else:
        data = natural_language_to_yaml(text, project_name)
        phases_raw = data[project_name]

    if isinstance(phases_raw, list):
        phases = phases_raw
    elif isinstance(phases_raw, dict):
        phases = [phases_raw]
    else:
        phases = []
    tasks = schedule_tasks(phases, holidays=project_holidays, resource_non_working_days=resource_nwd)

    phase_dates = {}
    for t in tasks:
        phase = t.get('phase')
        start = t.get('start') or today
        duration = t.get('duration') if 'duration' in t else timedelta(days=1)
        finish = t.get('finish') or (start + duration)
        if phase:
            if phase not in phase_dates:
                phase_dates[phase] = {'start': start, 'end': finish}
            else:
                if start < phase_dates[phase]['start']:
                    phase_dates[phase]['start'] = start
                if finish > phase_dates[phase]['end']:
                    phase_dates[phase]['end'] = finish

    col_widths = {
        'id': len('ID'),
        'task_name': len('Task Name'),
        'start': len('Start'),
        'finish': len('Finish'),
        'duration': len('Dur'),
        'resources': len('Res'),
        'percent': len('%'),
        'rag': len('RAG'),
        'comment': len('Comment')
    }

    date_format = '%d %b' if terminal_width < 100 else '%Y-%m-%d'

    task_rows = []
    for idx, t in enumerate(tasks, start=1):
        start = t.get('start') or today
        duration = t.get('duration') if 'duration' in t else timedelta(days=1)
        finish = t.get('finish') or (start + duration)
        resources = t.get('resources','')
        if resources:
            resources = ', '.join([r.lstrip('@').strip() for r in resources.split(',')])
            if resource_map:
                resource_list = [r.strip() for r in resources.split(',')]
                mapped_resources = [resource_map.get(r.lower(), r) for r in resource_list]
                resources = ', '.join(mapped_resources)
        percent = t.get('percent', '')
        comment = t.get('comment','')

        if t.get('summary'):
            rag_status = ''
        else:
            rag_status = calculate_rag_status(t)

        task_name = t.get('description') or t.get('name', '')

        level = t.get('level', 0)
        if level > 0:
            indent = '  ' * level
            task_name = f"{indent}{task_name}"

        row = {
            'id': str(idx),
            'task_name': task_name,
            'start': start.strftime(date_format),
            'finish': finish.strftime(date_format),
            'duration': f"{t.get('duration',timedelta(days=1)).days}d",
            'resources': str(resources),
            'percent': str(percent),
            'rag': rag_status,
            'comment': str(comment)
        }
        task_rows.append(row)

        for key, value in row.items():
            if key != 'comment':
                col_widths[key] = max(col_widths[key], len(value))

    fixed_width = (col_widths['id'] + col_widths['task_name'] + col_widths['start'] +
                   col_widths['finish'] + col_widths['duration'] + col_widths['resources'] +
                   col_widths['percent'] + col_widths['rag'])
    separators = 7 * 3
    available_for_comment = terminal_width - fixed_width - separators - 5

    max_comment_width = max(len('Comment'), min(50, available_for_comment))
    col_widths['comment'] = max_comment_width

    for row in task_rows:
        if len(row['comment']) > max_comment_width:
            row['comment'] = row['comment'][:max_comment_width - 3] + '...'

    md = f"# {project_name}\n\n"

    md += f"{'ID':<{col_widths['id']}} | {'Task Name':<{col_widths['task_name']}} | {'Start':<{col_widths['start']}} | {'Finish':<{col_widths['finish']}} | {'Dur':>{col_widths['duration']}} | {'Res':^{col_widths['resources']}} | {'%':^{col_widths['percent']}} | {'RAG':^{col_widths['rag']}} | {'Comment':<{col_widths['comment']}}\n"

    md += f"{'-' * (col_widths['id'] + 1)}|{'-' * (col_widths['task_name'] + 2)}|{'-' * (col_widths['start'] + 2)}|{'-' * (col_widths['finish'] + 2)}|{'-' * (col_widths['duration'] + 2)}|{'-' * (col_widths['resources'] + 2)}|{'-' * (col_widths['percent'] + 2)}|{'-' * (col_widths['rag'] + 2)}|{'-' * (col_widths['comment'])}\n"

    for row in task_rows:
        md += f"{row['id']:<{col_widths['id']}} | {row['task_name']:<{col_widths['task_name']}} | {row['start']:<{col_widths['start']}} | {row['finish']:<{col_widths['finish']}} | {row['duration']:>{col_widths['duration']}} | {row['resources']:^{col_widths['resources']}} | {row['percent']:^{col_widths['percent']}} | {row['rag']:^{col_widths['rag']}} | {row['comment']:<{col_widths['comment']}}\n"

    milestones = []
    if phase_dates or tasks:
        milestone_entries = []

        for phase, dates in phase_dates.items():
            milestone_entries.append({
                'name': phase, 'start': dates['start'], 'end': dates['end'], 'type': 'phase'
            })

        for t in tasks:
            if t.get('summary'):
                milestone_entries.append({
                    'name': t.get('description', t.get('name', '')),
                    'start': t.get('start'), 'end': t.get('finish'), 'type': 'summary'
                })

        for t in tasks:
            duration = t.get('duration', timedelta(days=1))
            if isinstance(duration, timedelta) and duration.days == 0:
                milestone_name = t.get('description', t.get('name', ''))
                display_name = milestone_name.replace('_', ' ')
                milestone_entries.append({
                    'name': display_name, 'start': t.get('start'),
                    'end': t.get('start'), 'type': 'milestone'
                })
                milestones.append({
                    'name': display_name, 'date': t.get('start'),
                    'percent': t.get('percent', 0)
                })

        if not phase_dates:
            for t in tasks:
                if t.get('summary'):
                    continue
                task_name = t.get('description', t.get('name', ''))
                display_name = task_name.replace('_', ' ')
                milestone_entries.append({
                    'name': display_name, 'start': t.get('start'),
                    'end': t.get('finish'), 'type': 'task'
                })
                milestones.append({'name': f"{display_name}", 'date': t.get('finish')})

        seen = set()
        unique_entries = []
        for entry in milestone_entries:
            key = (entry['name'], entry['start'], entry['end'])
            if key not in seen:
                seen.add(key)
                unique_entries.append(entry)

        milestone_entries = unique_entries
        milestone_entries.sort(key=lambda x: x['start'] if x['start'] else datetime.now())

        timeline_widths = {
            'name': len('Milestone'), 'start': len('Start'), 'end': len('End')
        }
        for entry in milestone_entries:
            timeline_widths['name'] = max(timeline_widths['name'], len(entry['name']))
            if entry['start']:
                timeline_widths['start'] = max(timeline_widths['start'], len(entry['start'].strftime(date_format)))
            if entry['end']:
                timeline_widths['end'] = max(timeline_widths['end'], len(entry['end'].strftime(date_format)))

        md += f"\n# Project Milestones\n\n{'Milestone':<{timeline_widths['name']}} | {'Start':<{timeline_widths['start']}} | {'End':<{timeline_widths['end']}}\n"
        md += f"{'-' * (timeline_widths['name'] + 1)}|{'-' * (timeline_widths['start'] + 2)}|{'-' * (timeline_widths['end'] + 1)}\n"

        for entry in milestone_entries:
            start_str = entry['start'].strftime(date_format) if entry['start'] else ''
            end_str = entry['end'].strftime(date_format) if entry['end'] else ''
            md += f"{entry['name']:<{timeline_widths['name']}} | {start_str:<{timeline_widths['start']}} | {end_str:<{timeline_widths['end']}}\n"

        for phase, dates in phase_dates.items():
            display_phase = phase.replace('_', ' ')
            milestones.append({'name': f"End of {display_phase}", 'date': dates['end']})
        phase_objs = [{'name': k, 'start': v['start']} for k,v in phase_dates.items()]
        tw = terminal_width
        if tasks:
            start_date = min([t.get('start', today) for t in tasks])
            finish_date = max([t.get('finish', today) for t in tasks])
        else:
            start_date = today
            finish_date = today
        timeline_row, milestone_labels, milestone_dates_str, milestone_full_dates, milestone_positions = render_custom_timeline(phase_objs, milestones, start_date, finish_date, tw)

        start_date_str = milestone_full_dates[0] if milestone_full_dates else ''
        finish_date_str = milestone_full_dates[-1] if milestone_full_dates else ''

        def create_connector_line():
            connector = [' '] * tw
            if not milestone_positions:
                return ''.join(connector)
            valid_positions = [p for p in milestone_positions if 0 <= p < tw]
            if not valid_positions:
                return ''.join(connector)
            for pos in valid_positions:
                connector[pos] = '\u2502'
            return ''.join(connector)

        md += f"Project: {project_name}\n"
        md += f"Start{' ' * (tw - 11)}Finish\n"
        md += f"{start_date_str}{' ' * (tw - len(start_date_str) - len(finish_date_str))}{finish_date_str}\n"

        for title_line in reversed(milestone_labels):
            md += f"{title_line}\n"

        if milestone_labels:
            md += f"{create_connector_line()}\n"

        combined_timeline = list(timeline_row)
        total_days = (finish_date - start_date).days or 1

        if phase_dates:
            for phase, dates in phase_dates.items():
                phase_tasks = [t for t in tasks if t.get('phase') == phase and not t.get('summary')]
                if phase_tasks:
                    total_progress = sum([t.get('percent', 0) for t in phase_tasks])
                    avg_progress = total_progress / len(phase_tasks) if phase_tasks else 0
                    phase_start_pos = int((dates['start'] - start_date).days / total_days * (tw - 1))
                    phase_end_pos = int((dates['end'] - start_date).days / total_days * (tw - 1))
                    phase_length = phase_end_pos - phase_start_pos + 1
                    progress_length = int(phase_length * avg_progress / 100)
                    for i in range(phase_start_pos, phase_end_pos + 1):
                        if i < tw and combined_timeline[i] != '\u25c6':
                            if i < phase_start_pos + progress_length:
                                combined_timeline[i] = '\u2550'
                            else:
                                combined_timeline[i] = '\u2500'
        else:
            for t in tasks:
                if t.get('summary'):
                    continue
                duration = t.get('duration', timedelta(days=1))
                if isinstance(duration, timedelta) and duration.days == 0:
                    continue
                task_start = t.get('start')
                task_finish = t.get('finish')
                task_percent = t.get('percent', 0)
                if task_start and task_finish:
                    task_start_pos = int((task_start - start_date).days / total_days * (tw - 1))
                    task_end_pos = int((task_finish - start_date).days / total_days * (tw - 1))
                    task_length = task_end_pos - task_start_pos + 1
                    progress_length = int(task_length * task_percent / 100)
                    for i in range(task_start_pos, task_end_pos + 1):
                        if i < tw and combined_timeline[i] != '\u25c6':
                            if i < task_start_pos + progress_length:
                                combined_timeline[i] = '\u2550'
                            else:
                                combined_timeline[i] = '\u2500'

        combined_timeline[0] = '\u251c'
        combined_timeline[-1] = '\u2524'

        md += f"{''.join(combined_timeline)}\n"
        md += f"{create_connector_line()}\n"
        md += f"{milestone_dates_str}\n"

    tw = terminal_width
    if tasks:
        start_date = min([t.get('start', today) for t in tasks])
        finish_date = max([t.get('finish', today) for t in tasks])
    else:
        start_date = today
        finish_date = today
    md += "\n"
    md += render_gantt_chart(tasks, start_date, finish_date, tw)
    md += "\n"
    md += render_resource_sheet(tasks, start_date, finish_date, holidays=project_holidays, terminal_width=tw, resource_map=resource_map)
    return md

def yaml_to_markdown_table(yaml_path, terminal_width=80):
    # Import here to avoid circular imports
    from .scheduling_engine import schedule_tasks

    today = datetime.today()

    with open(yaml_path, encoding="utf-8") as f:
        original_text = f.read()

    resource_map, resource_nwd = parse_resource_mappings(original_text)
    from .front_matter_parser import FrontMatterParser
    _fm_parser = FrontMatterParser(original_text)
    project_holidays = _fm_parser.parse_non_working_days()

    data = yaml.safe_load(original_text)
    project_name = list(data.keys())[0]
    phases_raw = data[project_name]
    if isinstance(phases_raw, list):
        phases = phases_raw
    elif isinstance(phases_raw, dict):
        phases = [phases_raw]
    else:
        phases = []
    tasks = schedule_tasks(phases, holidays=project_holidays, resource_non_working_days=resource_nwd)

    phase_dates = {}
    for t in tasks:
        phase = t.get('phase')
        start = t.get('start') or today
        duration = t.get('duration') if 'duration' in t else timedelta(days=1)
        finish = t.get('finish') or (start + duration)
        if phase:
            if phase not in phase_dates:
                phase_dates[phase] = {'start': start, 'end': finish}
            else:
                if start < phase_dates[phase]['start']:
                    phase_dates[phase]['start'] = start
                if finish > phase_dates[phase]['end']:
                    phase_dates[phase]['end'] = finish

    col_widths = {
        'id': len('ID'), 'task_name': len('Task Name'), 'start': len('Start'),
        'finish': len('Finish'), 'duration': len('Dur'), 'resources': len('Res'),
        'percent': len('%'), 'rag': len('RAG'), 'comment': len('Comment')
    }

    date_format = '%d %b' if terminal_width < 100 else '%Y-%m-%d'

    task_rows = []
    for idx, t in enumerate(tasks, start=1):
        start = t.get('start') or today
        duration = t.get('duration') if 'duration' in t else timedelta(days=1)
        finish = t.get('finish') or (start + duration)
        resources = t.get('resources','')
        if resources:
            resources = ', '.join([r.lstrip('@').strip() for r in resources.split(',')])
            if resource_map:
                resource_list = [r.strip() for r in resources.split(',')]
                mapped_resources = [resource_map.get(r.lower(), r) for r in resource_list]
                resources = ', '.join(mapped_resources)
        percent = t.get('percent', '')
        comment = t.get('comment','')

        if t.get('summary'):
            rag_status = ''
        else:
            rag_status = calculate_rag_status(t)

        task_name = t.get('description') or t.get('name', '')
        level = t.get('level', 0)
        if level > 0:
            indent = '  ' * level
            task_name = f"{indent}{task_name}"

        row = {
            'id': str(idx), 'task_name': task_name,
            'start': start.strftime(date_format), 'finish': finish.strftime(date_format),
            'duration': f"{t.get('duration',timedelta(days=1)).days}d",
            'resources': str(resources), 'percent': str(percent),
            'rag': rag_status, 'comment': str(comment)
        }
        task_rows.append(row)

        for key, value in row.items():
            if key != 'comment':
                col_widths[key] = max(col_widths[key], len(value))

    fixed_width = (col_widths['id'] + col_widths['task_name'] + col_widths['start'] +
                   col_widths['finish'] + col_widths['duration'] + col_widths['resources'] +
                   col_widths['percent'] + col_widths['rag'])
    separators = 7 * 3
    available_for_comment = terminal_width - fixed_width - separators - 5
    max_comment_width = max(len('Comment'), min(50, available_for_comment))
    col_widths['comment'] = max_comment_width

    for row in task_rows:
        if len(row['comment']) > max_comment_width:
            row['comment'] = row['comment'][:max_comment_width - 3] + '...'

    md = f"# {project_name}\n\n"
    md += f"{'ID':<{col_widths['id']}} | {'Task Name':<{col_widths['task_name']}} | {'Start':<{col_widths['start']}} | {'Finish':<{col_widths['finish']}} | {'Dur':>{col_widths['duration']}} | {'Res':^{col_widths['resources']}} | {'%':^{col_widths['percent']}} | {'RAG':^{col_widths['rag']}} | {'Comment':<{col_widths['comment']}}\n"
    md += f"{'-' * (col_widths['id'] + 1)}|{'-' * (col_widths['task_name'] + 2)}|{'-' * (col_widths['start'] + 2)}|{'-' * (col_widths['finish'] + 2)}|{'-' * (col_widths['duration'] + 2)}|{'-' * (col_widths['resources'] + 2)}|{'-' * (col_widths['percent'] + 2)}|{'-' * (col_widths['rag'] + 2)}|{'-' * (col_widths['comment'])}\n"

    for row in task_rows:
        md += f"{row['id']:<{col_widths['id']}} | {row['task_name']:<{col_widths['task_name']}} | {row['start']:<{col_widths['start']}} | {row['finish']:<{col_widths['finish']}} | {row['duration']:>{col_widths['duration']}} | {row['resources']:^{col_widths['resources']}} | {row['percent']:^{col_widths['percent']}} | {row['rag']:^{col_widths['rag']}} | {row['comment']:<{col_widths['comment']}}\n"

    milestones = []
    if phase_dates or tasks:
        milestone_entries = []
        for phase, dates in phase_dates.items():
            milestone_entries.append({'name': phase, 'start': dates['start'], 'end': dates['end'], 'type': 'phase'})
        for t in tasks:
            if t.get('summary'):
                milestone_entries.append({'name': t.get('description', t.get('name', '')), 'start': t.get('start'), 'end': t.get('finish'), 'type': 'summary'})
        for t in tasks:
            duration = t.get('duration', timedelta(days=1))
            if isinstance(duration, timedelta) and duration.days == 0:
                milestone_name = t.get('description', t.get('name', ''))
                display_name = milestone_name.replace('_', ' ')
                milestone_entries.append({'name': display_name, 'start': t.get('start'), 'end': t.get('start'), 'type': 'milestone'})
                milestones.append({'name': display_name, 'date': t.get('start'), 'percent': t.get('percent', 0)})
        if not phase_dates:
            for t in tasks:
                if t.get('summary'):
                    continue
                task_name = t.get('description', t.get('name', ''))
                display_name = task_name.replace('_', ' ')
                milestone_entries.append({'name': display_name, 'start': t.get('start'), 'end': t.get('finish'), 'type': 'task'})
                milestones.append({'name': f"{display_name}", 'date': t.get('finish')})
        seen = set()
        unique_entries = []
        for entry in milestone_entries:
            key = (entry['name'], entry['start'], entry['end'])
            if key not in seen:
                seen.add(key)
                unique_entries.append(entry)
        milestone_entries = unique_entries
        milestone_entries.sort(key=lambda x: x['start'] if x['start'] else datetime.now())
        timeline_widths = {'name': len('Milestone'), 'start': len('Start'), 'end': len('End')}
        for entry in milestone_entries:
            timeline_widths['name'] = max(timeline_widths['name'], len(entry['name']))
            if entry['start']:
                timeline_widths['start'] = max(timeline_widths['start'], len(entry['start'].strftime(date_format)))
            if entry['end']:
                timeline_widths['end'] = max(timeline_widths['end'], len(entry['end'].strftime(date_format)))
        md += f"\n# Project Milestones\n\n{'Milestone':<{timeline_widths['name']}} | {'Start':<{timeline_widths['start']}} | {'End':<{timeline_widths['end']}}\n"
        md += f"{'-' * (timeline_widths['name'] + 1)}|{'-' * (timeline_widths['start'] + 2)}|{'-' * (timeline_widths['end'] + 1)}\n"
        for entry in milestone_entries:
            start_str = entry['start'].strftime(date_format) if entry['start'] else ''
            end_str = entry['end'].strftime(date_format) if entry['end'] else ''
            md += f"{entry['name']:<{timeline_widths['name']}} | {start_str:<{timeline_widths['start']}} | {end_str:<{timeline_widths['end']}}\n"
        for phase, dates in phase_dates.items():
            display_phase = phase.replace('_', ' ')
            milestones.append({'name': f"End of {display_phase}", 'date': dates['end']})
        phase_objs = [{'name': k, 'start': v['start']} for k,v in phase_dates.items()]
        tw = terminal_width
        if tasks:
            start_date = min([t.get('start', today) for t in tasks])
            finish_date = max([t.get('finish', today) for t in tasks])
        else:
            start_date = today
            finish_date = today
        timeline_row, milestone_labels, milestone_dates_str, milestone_full_dates, milestone_positions = render_custom_timeline(phase_objs, milestones, start_date, finish_date, tw)
        start_date_str = milestone_full_dates[0] if milestone_full_dates else ''
        finish_date_str = milestone_full_dates[-1] if milestone_full_dates else ''
        def create_connector_line():
            connector = [' '] * tw
            if not milestone_positions:
                return ''.join(connector)
            valid_positions = [p for p in milestone_positions if 0 <= p < tw]
            if not valid_positions:
                return ''.join(connector)
            for pos in valid_positions:
                connector[pos] = '\u2502'
            return ''.join(connector)
        md += f"Project: {project_name}\n"
        md += f"Start{' ' * (tw - 11)}Finish\n"
        md += f"{start_date_str}{' ' * (tw - len(start_date_str) - len(finish_date_str))}{finish_date_str}\n"
        for title_line in reversed(milestone_labels):
            md += f"{title_line}\n"
        if milestone_labels:
            md += f"{create_connector_line()}\n"
        combined_timeline = list(timeline_row)
        total_days = (finish_date - start_date).days or 1
        if phase_dates:
            for phase, dates in phase_dates.items():
                phase_tasks = [t for t in tasks if t.get('phase') == phase and not t.get('summary')]
                if phase_tasks:
                    total_progress = sum([t.get('percent', 0) for t in phase_tasks])
                    avg_progress = total_progress / len(phase_tasks) if phase_tasks else 0
                    phase_start_pos = int((dates['start'] - start_date).days / total_days * (tw - 1))
                    phase_end_pos = int((dates['end'] - start_date).days / total_days * (tw - 1))
                    phase_length = phase_end_pos - phase_start_pos + 1
                    progress_length = int(phase_length * avg_progress / 100)
                    for i in range(phase_start_pos, phase_end_pos + 1):
                        if i < tw and combined_timeline[i] != '\u25c6':
                            if i < phase_start_pos + progress_length:
                                combined_timeline[i] = '\u2550'
                            else:
                                combined_timeline[i] = '\u2500'
        else:
            for t in tasks:
                if t.get('summary'):
                    continue
                duration = t.get('duration', timedelta(days=1))
                if isinstance(duration, timedelta) and duration.days == 0:
                    continue
                task_start = t.get('start')
                task_finish = t.get('finish')
                task_percent = t.get('percent', 0)
                if task_start and task_finish:
                    task_start_pos = int((task_start - start_date).days / total_days * (tw - 1))
                    task_end_pos = int((task_finish - start_date).days / total_days * (tw - 1))
                    task_length = task_end_pos - task_start_pos + 1
                    progress_length = int(task_length * task_percent / 100)
                    for i in range(task_start_pos, task_end_pos + 1):
                        if i < tw and combined_timeline[i] != '\u25c6':
                            if i < task_start_pos + progress_length:
                                combined_timeline[i] = '\u2550'
                            else:
                                combined_timeline[i] = '\u2500'
        combined_timeline[0] = '\u251c'
        combined_timeline[-1] = '\u2524'
        md += f"{''.join(combined_timeline)}\n"
        md += f"{create_connector_line()}\n"
        md += f"{milestone_dates_str}\n"
    tw = terminal_width
    if tasks:
        start_date = min([t.get('start', today) for t in tasks])
        finish_date = max([t.get('finish', today) for t in tasks])
    else:
        start_date = today
        finish_date = today
    md += "\n"
    md += render_gantt_chart(tasks, start_date, finish_date, tw)
    md += "\n"
    md += render_resource_sheet(tasks, start_date, finish_date, holidays=project_holidays, terminal_width=tw, resource_map=resource_map)
    return md


def export_to_pdf(text, output_path, is_yaml=True, project_name="Project", original_text=None):
    """Export the project plan to a PDF file."""
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.pdfgen import canvas

    ascii_output = text_to_markdown_table(
        text, is_yaml=is_yaml, project_name=project_name,
        terminal_width=120, original_text=original_text
    )

    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )

    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'CustomTitle', parent=styles['Heading1'],
        fontSize=24, textColor='#108BB9', spaceAfter=30, alignment=TA_CENTER
    )

    mono_style = ParagraphStyle(
        'Monospace', parent=styles['Code'],
        fontName='Courier', fontSize=8, leading=10,
        leftIndent=0, rightIndent=0, alignment=TA_LEFT
    )

    title = Paragraph(f"{project_name} - Project Plan", title_style)
    story.append(title)
    story.append(Spacer(1, 12))

    preformatted = Preformatted(ascii_output, mono_style)
    story.append(preformatted)

    doc.build(story)


def export_to_csv(text, output_path, is_yaml=True, project_name="Project", original_text=None):
    """Export project data to CSV format."""
    # Import here to avoid circular imports
    from .scheduling_engine import natural_language_to_yaml, schedule_tasks

    resource_map = {}
    if original_text:
        resource_map, _ = parse_resource_mappings(original_text)

    if is_yaml:
        data = yaml.safe_load(text)
        project_name = list(data.keys())[0]
        phases_raw = data[project_name]
    else:
        data = natural_language_to_yaml(text, project_name)
        phases_raw = data[project_name]

    if isinstance(phases_raw, list):
        phases = phases_raw
    elif isinstance(phases_raw, dict):
        phases = [phases_raw]
    else:
        phases = []

    tasks = schedule_tasks(phases)

    with open(output_path, 'w', newline='', encoding='utf-8') as csvfile:
        fieldnames = ['ID', 'Task Name', 'Start', 'Finish', 'Duration (days)',
                      'Resources', '% Complete', 'RAG', 'Priority', 'Bucket', 'Comment']
        writer = csv.DictWriter(csvfile, fieldnames=fieldnames)

        writer.writeheader()

        for idx, task in enumerate(tasks, start=1):
            duration = task.get('duration')
            if isinstance(duration, timedelta):
                duration_days = duration.days
            else:
                duration_days = 0

            start_date = task.get('start')
            finish_date = task.get('finish')

            start_str = start_date.strftime('%Y-%m-%d') if start_date else ''
            finish_str = finish_date.strftime('%Y-%m-%d') if finish_date else ''

            resources = task.get('resources', [])
            resource_str = ', '.join(resources) if resources else ''

            percentage = task.get('percentage', 0)
            rag_status = task.get('rag_status', 'N/A')
            comment = task.get('comment', '')

            writer.writerow({
                'ID': idx,
                'Task Name': task.get('name', ''),
                'Start': start_str,
                'Finish': finish_str,
                'Duration (days)': duration_days,
                'Resources': resource_str,
                '% Complete': percentage,
                'RAG': rag_status,
                'Priority': task.get('priority', 'Low'),
                'Bucket': task.get('bucket', ''),
                'Comment': comment
            })


# ---------------------------------------------------------------------------
# Tasks by Assignment and Slippage reports (#776)
#
# The browser computes both reports (static/plan-reports.js) and posts the
# rows; these functions only lay them out, so the file matches the view.
# ---------------------------------------------------------------------------

_REPORT_BLUE = RGBColor(33, 60, 114)
_REPORT_GREY = RGBColor(100, 100, 100)
_REPORT_RED = RGBColor(192, 57, 43)
_REPORT_GREEN = RGBColor(39, 124, 67)
_REPORT_HEADER_FILL = PatternFill(start_color="213C72", end_color="213C72", fill_type="solid")
_REPORT_ALERT_FILL = PatternFill(start_color="FDECEA", end_color="FDECEA", fill_type="solid")
_REPORT_ROWS_PER_SLIDE = 14


def _report_text(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", "" if value is None else str(value))


def _variance_text(days):
    days = int(days or 0)
    if days == 0:
        return "0"
    return f"{'+' if days > 0 else '-'}{abs(days)}d"


def _write_report_sheet(ws, headers, rows, alert=None):
    """Header row, data rows, widths; rows for which alert(row) holds are tinted."""
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _REPORT_HEADER_FILL
    for row in rows:
        ws.append([_report_text(v) if isinstance(v, str) else v for v in row])
        if alert and alert(row):
            for cell in ws[ws.max_row]:
                cell.fill = _REPORT_ALERT_FILL
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[index - 1])) for r in rows if index - 1 < len(r)])
        ws.column_dimensions[get_column_letter(index)].width = min(60, width + 2)
    ws.freeze_panes = "A2"


def _report_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def _report_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    p.text = _report_text(title)
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = _REPORT_BLUE
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.4))
        sp = sub.text_frame.paragraphs[0]
        sp.text = _report_text(subtitle)
        sp.font.size = Pt(12)
        sp.font.color.rgb = _REPORT_GREY
    return slide


def _report_table_slides(prs, title, subtitle, headers, rows, widths=None, highlight=None):
    """One or more slides holding a table, paginated; `highlight(row)` reddens a row."""
    chunks = [rows[i:i + _REPORT_ROWS_PER_SLIDE] for i in range(0, len(rows), _REPORT_ROWS_PER_SLIDE)] or [[]]
    for page, chunk in enumerate(chunks, start=1):
        suffix = f" ({page}/{len(chunks)})" if len(chunks) > 1 else ""
        slide = _report_slide(prs, title + suffix, subtitle)
        if not chunk:
            note = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(12), Inches(0.5))
            note.text_frame.paragraphs[0].text = "Nothing to show."
            continue
        shape = slide.shapes.add_table(len(chunk) + 1, len(headers), Inches(0.5), Inches(1.5),
                                       Inches(12.3), Inches(0.4) * (len(chunk) + 1))
        table = shape.table
        for c, header in enumerate(headers):
            cell = table.cell(0, c)
            cell.text = header
            cell.text_frame.paragraphs[0].font.size = Pt(11)
            cell.text_frame.paragraphs[0].font.bold = True
            if widths:
                table.columns[c].width = Inches(widths[c])
        for r, row in enumerate(chunk, start=1):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = _report_text(value)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(10)
                if highlight and highlight(row):
                    para.font.color.rgb = _REPORT_RED


def _assignment_rows(group):
    return [
        [t.get("name", ""), t.get("phase", ""), t.get("start", ""), t.get("finish", ""),
         f"{t.get('duration', 0)}d", f"{t.get('percent', 0):g}%", str(t.get("status", "")).title(),
         ", ".join(t.get("shared_with") or [])]
        for t in group.get("tasks", [])
    ]


_ASSIGNMENT_HEADERS = ["Task", "Phase", "Start", "Finish", "Duration", "%", "Status", "Shared with"]


def export_assignment_report_to_excel(output_path, data, project_name="Project"):
    """Tasks by Assignment as a workbook: a Summary sheet, then one sheet per person."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    groups = data.get("groups", [])
    _write_report_sheet(
        summary,
        ["Person", "Role", "Tasks", "Open", "Complete", "Work (days)", "Next due", "Overdue"],
        [[g.get("name", ""), g.get("role", ""), g["totals"].get("tasks", 0), g["totals"].get("open", 0),
          g["totals"].get("complete", 0), g["totals"].get("workDays", 0), g["totals"].get("nextDue") or "",
          g["totals"].get("overdue", 0)] for g in groups],
        alert=lambda row: row[0] == "Unassigned" or row[7],
    )
    used = {"Summary"}
    for group in groups:
        name = re.sub(r"[\[\]:*?/\\]", "", group.get("name") or "Person")[:31] or "Person"
        base, n = name, 2
        while name in used:
            name = f"{base[:28]} {n}"
            n += 1
        used.add(name)
        ws = wb.create_sheet(name)
        _write_report_sheet(ws, _ASSIGNMENT_HEADERS, _assignment_rows(group),
                            alert=lambda row: row[6] == "Overdue")
    wb.save(output_path)
    logger.info(f"Exported tasks by assignment to Excel: {output_path}")


def export_assignment_report_to_powerpoint(output_path, data, project_name="Project"):
    """Tasks by Assignment as a deck: an overview slide, then each person's tasks."""
    prs = _report_deck()
    groups = data.get("groups", [])
    _report_table_slides(
        prs, f"{project_name} — Tasks by assignment", "Who is carrying what",
        ["Person", "Role", "Tasks", "Open", "Complete", "Work", "Next due", "Overdue"],
        [[g.get("name", ""), g.get("role", ""), g["totals"].get("tasks", 0), g["totals"].get("open", 0),
          g["totals"].get("complete", 0), f"{g['totals'].get('workDays', 0)}d",
          g["totals"].get("nextDue") or "—", g["totals"].get("overdue", 0)] for g in groups],
        widths=[2.8, 2.5, 0.9, 0.9, 1.1, 1.0, 1.6, 1.5],
        highlight=lambda row: row[0] == "Unassigned" or row[7],
    )
    for group in groups:
        totals = group.get("totals", {})
        subtitle = (f"{totals.get('tasks', 0)} tasks · {totals.get('open', 0)} open · "
                    f"{totals.get('overdue', 0)} overdue · next due {totals.get('nextDue') or '—'}")
        _report_table_slides(
            prs, group.get("name", ""), subtitle, _ASSIGNMENT_HEADERS, _assignment_rows(group),
            widths=[3.2, 2.0, 1.2, 1.2, 0.9, 0.7, 1.1, 2.0],
            highlight=lambda row: row[6] == "Overdue",
        )
    prs.save(output_path)
    logger.info(f"Exported tasks by assignment to PowerPoint: {output_path}")


_SLIPPAGE_HEADERS = ["Task", "Phase", "Baselined start", "Baselined finish", "Start", "Finish",
                     "Start var.", "Finish var.", "Status", "Critical"]
_SLIPPAGE_STATUS = {"slipped": "Slipped", "on-track": "On track", "pulled-forward": "Pulled forward"}


def _slippage_row(t):
    return [t.get("name", ""), t.get("phase", ""), t.get("baseline_start", ""), t.get("baseline_finish", ""),
            t.get("start", ""), t.get("finish", ""), _variance_text(t.get("start_variance")),
            _variance_text(t.get("finish_variance")), _SLIPPAGE_STATUS.get(t.get("status"), t.get("status", "")),
            "Yes" if t.get("critical") else ""]


def _slippage_headline(data, project_name):
    project = data.get("project", {})
    counts = data.get("counts", {})
    variance = int(project.get("variance") or 0)
    state = ("late" if variance > 0 else "early" if variance < 0 else "on the baseline")
    lead = f"{abs(variance)} working days {state}" if variance else "On the baseline"
    return (f"{lead}: finishing {project.get('finish') or '—'} against "
            f"{project.get('baseline_finish') or '—'} ({data.get('baseline') or 'baseline'}). "
            f"{counts.get('slipped', 0)} slipped, {counts.get('onTrack', 0)} on track, "
            f"{counts.get('pulledForward', 0)} pulled forward, {counts.get('added', 0)} added, "
            f"{counts.get('removed', 0)} removed.")


def export_slippage_report_to_excel(output_path, data, project_name="Project"):
    """The Slippage report as a workbook: Summary, Critical path, Phases, Tasks, Scope."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    project = data.get("project", {})
    counts = data.get("counts", {})
    _write_report_sheet(summary, ["Measure", "Value"], [
        ["Project", project_name],
        ["Compared with", data.get("baseline", "")],
        ["Baselined finish", project.get("baseline_finish", "")],
        ["Current finish", project.get("finish", "")],
        ["Variance (working days)", int(project.get("variance") or 0)],
        ["Slipped tasks", counts.get("slipped", 0)],
        ["On track", counts.get("onTrack", 0)],
        ["Pulled forward", counts.get("pulledForward", 0)],
        ["Added since baseline", counts.get("added", 0)],
        ["Removed since baseline", counts.get("removed", 0)],
        ["Critical-path tasks slipped", len(data.get("critical", []))],
    ])
    _write_report_sheet(wb.create_sheet("Critical path"), _SLIPPAGE_HEADERS,
                        [_slippage_row(t) for t in data.get("critical", [])], alert=lambda row: True)
    _write_report_sheet(
        wb.create_sheet("Phases"),
        ["Phase", "Baselined finish", "Finish", "Finish var.", "Slipped tasks", "Worst slip"],
        [[p.get("name", ""), p.get("baseline_finish", ""), p.get("finish", ""),
          _variance_text(p.get("finish_variance")), p.get("slipped_tasks", 0),
          _variance_text(p.get("worst_slip")) if p.get("worst_slip") else ""] for p in data.get("phases", [])],
        alert=lambda row: str(row[3]).startswith("+"),
    )
    _write_report_sheet(wb.create_sheet("Tasks"), _SLIPPAGE_HEADERS,
                        [_slippage_row(t) for t in data.get("tasks", [])],
                        alert=lambda row: row[9] == "Yes" and str(row[7]).startswith("+"))
    _write_report_sheet(
        wb.create_sheet("Scope changes"), ["Task", "Phase", "Change"],
        [[t.get("name", ""), t.get("phase", ""), "Added"] for t in data.get("added", [])]
        + [[t.get("name", ""), "", "Removed"] for t in data.get("removed", [])],
    )
    wb.save(output_path)
    logger.info(f"Exported slippage report to Excel: {output_path}")


def export_slippage_report_to_powerpoint(output_path, data, project_name="Project"):
    """The Slippage report as a deck: headline, critical path, phases, every task, scope."""
    prs = _report_deck()
    slide = _report_slide(prs, f"{project_name} — Slippage", data.get("baseline") and f"Against {data['baseline']}")
    variance = int(data.get("project", {}).get("variance") or 0)
    big = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(12.3), Inches(1.4))
    bp = big.text_frame.paragraphs[0]
    bp.text = _variance_text(variance) + (" working days" if variance else "")
    bp.font.size = Pt(54)
    bp.font.bold = True
    bp.font.color.rgb = _REPORT_RED if variance > 0 else _REPORT_GREEN
    text = slide.shapes.add_textbox(Inches(0.5), Inches(3.4), Inches(12.3), Inches(2))
    text.text_frame.word_wrap = True
    tp = text.text_frame.paragraphs[0]
    tp.text = _report_text(_slippage_headline(data, project_name))
    tp.font.size = Pt(16)

    _report_table_slides(prs, "Slippage on the critical path", "These slips move the end date",
                         ["Task", "Baselined finish", "Finish", "Variance"],
                         [[t.get("name", ""), t.get("baseline_finish", ""), t.get("finish", ""),
                           _variance_text(t.get("finish_variance"))] for t in data.get("critical", [])],
                         widths=[6.3, 2.0, 2.0, 2.0], highlight=lambda row: True)
    _report_table_slides(prs, "Slippage by phase", None,
                         ["Phase", "Baselined finish", "Finish", "Variance", "Slipped", "Worst"],
                         [[p.get("name", ""), p.get("baseline_finish", ""), p.get("finish", ""),
                           _variance_text(p.get("finish_variance")), p.get("slipped_tasks", 0),
                           _variance_text(p.get("worst_slip")) if p.get("worst_slip") else "—"]
                          for p in data.get("phases", [])],
                         widths=[4.3, 1.8, 1.8, 1.4, 1.4, 1.6],
                         highlight=lambda row: str(row[3]).startswith("+"))
    _report_table_slides(prs, "Every task", "Largest slip first",
                         ["Task", "Baselined finish", "Finish", "Finish var.", "Status"],
                         [[t.get("name", ""), t.get("baseline_finish", ""), t.get("finish", ""),
                           _variance_text(t.get("finish_variance")),
                           _SLIPPAGE_STATUS.get(t.get("status"), "")] for t in data.get("tasks", [])],
                         widths=[5.3, 1.9, 1.9, 1.5, 1.7],
                         highlight=lambda row: str(row[3]).startswith("+"))
    if data.get("added") or data.get("removed"):
        _report_table_slides(prs, "Changed since the baseline", None, ["Task", "Change"],
                             [[t.get("name", ""), "Added"] for t in data.get("added", [])]
                             + [[t.get("name", ""), "Removed"] for t in data.get("removed", [])],
                             widths=[9.3, 3.0])
    prs.save(output_path)
    logger.info(f"Exported slippage report to PowerPoint: {output_path}")
